# -*- coding: utf-8 -*-
"""Planilha editável do Bloco 08: desenho MUS, amostra, avaliação (Stringer), ampliação, testes integrais,
distorções (PT 450) e comparação pós-congelamento. Base SINTÉTICA."""
import json, csv, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

T = '/root/cpfl/testes_b08'; S = f'{T}/saida_b08'
OUT = sys.argv[1] if len(sys.argv) > 1 else '/root/cpfl/producao/04_Testes/CPFL_B08_Amostragem_Testes_Distorcoes.xlsx'
P = json.load(open(f'{T}/parametros_b08.json')); R = json.load(open(f'{S}/resultados_b08.json'))
C = json.load(open(f'{S}/avaliacao_complementar_b08.json')); G = json.load(open(f'{S}/comparacao_gabarito_b08.json'))
K = json.load(open(f'{T}/B08_congelamento.json'))
F = 'Arial'; H = Font(name=F, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3864')
TT = Font(name=F, bold=True, size=12); N = Font(name=F); B = Font(name=F, bold=True)
BL = Font(name=F, color='0000FF'); GR = Font(name=F, color='008000'); INP = PatternFill('solid', fgColor='FFF2CC')
th = Side(style='thin', color='A6A6A6'); BD = Border(left=th, right=th, top=th, bottom=th)
RS = '#,##0.00;(#,##0.00);"-"'; PCT = '0.0%'; FAC = '0.0000'
wb = Workbook(); wb.remove(wb.active)
def sheet(name, title, widths):
    ws = wb.create_sheet(name); ws['A1'] = title; ws['A1'].font = TT
    ws['A2'] = 'Base SINTÉTICA para fins didáticos — resultados não são constatações sobre a CPFL Energia S.A. Preparado com apoio de IA em 23/10/2026.'
    ws['A2'].font = Font(name=F, italic=True, size=9)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    return ws
def head(ws, r, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=r, column=i, value=c); x.font = H; x.fill = HF; x.border = BD; x.alignment = Alignment(wrap_text=True)
def put(ws, r, vals, fmt=None, inp=(), bold=False, link=()):
    for i, v in enumerate(vals, 1):
        x = ws.cell(row=r, column=i, value=v); x.border = BD; x.font = B if bold else N
        x.alignment = Alignment(wrap_text=True, vertical='top')
        if fmt and fmt.get(i): x.number_format = fmt[i]
        if i in inp: x.font = BL; x.fill = INP
        if i in link: x.font = GR

# ---------------- Leia-me
ws = sheet('Leia_me', 'Bloco 08 — Amostragem (PT 530), testes integrais e distorções simuladas (PT 450)', [120])
for i, t in enumerate([
    'Sequência documentada: (1) desenho registrado no Drive antes da execução (B08_Desenho_pre_registro.md, 23/10/2026 19:54); (2) execução de testes_b08.py; (3) resposta ao resultado (avaliacao_complementar_b08.py); (4) congelamento por hash registrado no Drive (B08_congelamento.json, 19:55:53); (5) só então abertura dos arquivos reservados (aba Comparacao_pos).',
    'Cores: azul com fundo amarelo = dado de entrada ou valor calculado pelo código (fonte indicada); preto = fórmula; verde = vínculo entre abas.',
    'Abas: Desenho_MUS, Amostra_MUS, Avaliacao_MUS, Ampliacao, Testes_100, PT450_Distorcoes, Comparacao_pos, Congelamento.',
    'Limite: a mesma IA desenhou os erros e os testes; não é teste cego independente. Na simulação, o "valor auditado" de um item sorteado é obtido pelas mesmas regras T1–T3; a amostra não descobre tipos de erro fora dessas regras.'], 4):
    ws.cell(row=i, column=1, value=t).font = N; ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True)

# ---------------- Desenho
M = R['MUS']
ws = sheet('Desenho_MUS', 'Dimensionamento da amostra por unidade monetária (MUS)', [46, 22, 70])
head(ws, 4, ['Parâmetro', 'Valor', 'Fundamento / natureza'])
rows = [('Valor contábil da população (R$)', M['populacao_rs'], 'Base registrada conciliada (Bloco 07)', True),
        ('Distorção tolerável — DT (R$)', P['mus']['distorcao_toleravel_rs'], P['mus']['distorcao_toleravel_fundamento'], True),
        ('Risco de aceitação incorreta', P['mus']['risco_aceitacao_incorreta'], P['mus']['risco_fundamento'], True),
        ('Distorção esperada — DE (R$)', P['mus']['distorcao_esperada_rs'], P['mus']['distorcao_esperada_fundamento'], True),
        ('Fator de expansão — FE', P['mus']['fator_expansao'], P['mus']['fator_expansao_fonte'], True),
        ('Fator de confiança para 0 erros — FC', '=CHIINV(B7,2)/2', 'Limite superior de Poisson: qui-quadrado (95%; 2 g.l.) ÷ 2', False),
        ('Tamanho n = VC × FC ÷ (DT − DE × FE)', '=ROUNDUP(B5*B10/(B6-B8*B9),0)', 'Cálculo (literatura AICPA; NBC TA 530 item 7 e A10–A11)', False),
        ('Intervalo amostral J = VC ÷ n (R$)', '=B5/B11', 'Cálculo', False),
        ('Ponto inicial aleatório (R$)', M['ponto_inicial_rs'], f"random.Random({P['mus']['semente_ponto_inicial']}).uniform(0, J)", True),
        ('Itens ≥ J (seleção integral)', len(M['itens_selecao_integral_(>=J)']), 'Nenhum: maior item ≈ R$ 3,6 milhões', True),
        ('Itens selecionados mais de uma vez', M['itens_selecionados_mais_de_uma_vez'], 'Regra: testar uma vez', True),
        ('Conferência: n do código', M['tamanho_amostra_n'], 'Deve ser igual a B11', True),
        ('Diferença (B11 − B16)', '=B11-B16', 'Deve ser 0', False)]
for i, (a, b, c, inp) in enumerate(rows, 5):
    put(ws, i, [a, b, c], fmt={2: RS if isinstance(b, float) and b > 10 or (isinstance(b, str) and 'B5/' in b) else None}, inp=(2,) if inp else ())
for r in (7,): ws.cell(row=r, column=2).number_format = '0%'
for r in (9, 10): ws.cell(row=r, column=2).number_format = FAC
ws.cell(row=5, column=2).number_format = RS; ws.cell(row=6, column=2).number_format = RS; ws.cell(row=8, column=2).number_format = RS
ws.cell(row=12, column=2).number_format = RS; ws.cell(row=13, column=2).number_format = RS

# ---------------- Amostra
ws = sheet('Amostra_MUS', 'Amostra MUS (385 itens) — resultados por item', [7, 18, 18, 22, 16, 16, 16, 11, 16, 50])
head(ws, 4, ['Ordem', 'id_registro', 'Distribuidora', 'Categoria cadastrada', 'Valor registrado (R$)', 'Valor auditado simulado (R$)',
             'Distorção (R$)', 'Taint', 'Distorção projetada (R$)', 'Resultado'])
am = list(csv.DictReader(open(f'{S}/MUS_amostra_resultados.csv')))
for i, x in enumerate(am, 5):
    put(ws, i, [int(x['ordem']), x['id_registro'], x['distribuidora'], x['categoria_cadastrada'], float(x['valor_registrado_rs']),
                float(x['valor_auditado_simulado_rs']), f'=E{i}-F{i}', f'=G{i}/E{i}', f'=H{i}*Desenho_MUS!$B$12', x['resultado']],
        fmt={5: RS, 6: RS, 7: RS, 8: '0.0000', 9: RS}, inp=(5, 6))
L = 4 + len(am)
put(ws, L + 1, ['Total', '', '', '', f'=SUM(E5:E{L})', f'=SUM(F5:F{L})', f'=SUM(G5:G{L})', '', f'=SUM(I5:I{L})', ''], fmt={5: RS, 6: RS, 7: RS, 9: RS}, bold=True)
put(ws, L + 2, ['Itens com distorção', '', '', '', '', '', f'=COUNTIF(G5:G{L},"<>0")', '', '', ''], bold=True)
ws.freeze_panes = 'A5'

# ---------------- Avaliacao (Stringer)
ws = sheet('Avaliacao_MUS', 'Avaliação da amostra — projeção e limite superior de distorção (Stringer, Poisson 95%)', [36, 16, 16, 16, 20, 20])
head(ws, 4, ['Superavaliações ordenadas por taint', 'Taint', 'Fator UL(k)', 'Incremento do fator', 'Projeção (taint × J)', 'Componente do limite'])
put(ws, 5, ['Precisão básica (0 erros)', 1, '=CHIINV(Desenho_MUS!$B$7,2)/2', '=C5', '', '=Desenho_MUS!$B$12*C5'], fmt={3: FAC, 4: FAC, 6: RS})
nover = M['distorcoes_superavaliacao_na_amostra']
for k in range(1, nover + 1):
    r = 5 + k
    put(ws, r, [f'{k}ª maior', f'=LARGE(Amostra_MUS!$H$5:$H${L},{k})', f'=CHIINV(Desenho_MUS!$B$7,{2 * k + 2})/2', f'=C{r}-C{r-1}',
                f'=B{r}*Desenho_MUS!$B$12', f'=B{r}*D{r}*Desenho_MUS!$B$12'], fmt={2: '0.0000', 3: FAC, 4: FAC, 5: RS, 6: RS}, link=(2,))
e = 5 + nover
put(ws, e + 1, ['Distorção projetada (superavaliação)', '', '', '', f'=SUM(E6:E{e})', ''], fmt={5: RS}, bold=True)
put(ws, e + 2, ['Limite superior de distorção — UML', '', '', '', '', f'=SUM(F5:F{e})'], fmt={6: RS}, bold=True)
put(ws, e + 3, ['Distorção tolerável — DT', '', '', '', '', '=Desenho_MUS!B6'], fmt={6: RS}, link=(6,))
put(ws, e + 4, ['UML − DT (positivo = excede)', '', '', '', '', f'=F{e+2}-F{e+3}'], fmt={6: RS}, bold=True)
put(ws, e + 5, ['Conclusão pela regra de decisão', '', '', '', '', f'=IF(F{e+2}<=F{e+3},"UML <= DT","UML > DT: ampliar procedimentos")'], bold=True)
put(ws, e + 6, ['Conferência com o código (UML)', '', '', '', '', M['limite_superior_distorcao_UML_rs']], fmt={6: RS}, inp=(6,))
put(ws, e + 8, ['Subavaliações na amostra (projeção informativa)', '', '', '', f'=SUMIF(Amostra_MUS!I5:I{L},"<0")', 'MUS não é desenhado para subavaliação; coberta pelo T3 (100%)'], fmt={5: RS})

# ---------------- Testes 100%
ws = sheet('Testes_100', 'Procedimentos integrais (100% da população) — distorções conhecidas', [10, 52, 12, 20, 40])
head(ws, 4, ['Teste', 'Descrição', 'Itens', 'Efeito no saldo (R$)', 'Afirmação / natureza'])
t = [('T1', 'Bens duplicados (id_bem_tecnico repetido) — registro excedente', R['T1']['pares'], R['T1']['distorcao_rs'], 'Existência — factual (na simulação)'),
     ('T2', 'BAIXADO com valor registrado', R['T2']['itens'], R['T2']['distorcao_rs'], 'Existência / direitos — factual'),
     ('T3+', 'Recálculo: superavaliação (> 1%)', R['T3']['superavaliacao_itens'], R['T3']['superavaliacao_rs'], 'Avaliação — factual'),
     ('T3−', 'Recálculo: subavaliação (< −1%)', R['T3']['subavaliacao_itens'], R['T3']['subavaliacao_rs'], 'Avaliação — factual'),
     ('T4', 'Residual incoerente com a categoria cadastrada', R['T4']['itens'], 0, 'Classificação — sem efeito no saldo total')]
for i, x in enumerate(t, 5): put(ws, i, list(x), fmt={3: '#,##0', 4: RS}, inp=(3, 4))
put(ws, 10, ['', 'Superavaliação conhecida bruta (T1 + T2 + T3+)', '', '=D5+D6+D7', ''], fmt={4: RS}, bold=True)
put(ws, 11, ['', 'Subavaliação conhecida (T3−)', '', '=D8', ''], fmt={4: RS}, bold=True)
put(ws, 12, ['', 'Distorção conhecida líquida', '', '=D10+D11', ''], fmt={4: RS}, bold=True)
put(ws, 13, ['', 'Conferência com o código', '', R['distorcao_conhecida_liquida_rs'], ''], fmt={4: RS}, inp=(4,))

# ---------------- Ampliacao
ws = sheet('Ampliacao', 'Resposta ao resultado do MUS: população residual e amostra ampliada', [56, 22, 60])
head(ws, 4, ['Item', 'Valor', 'Fundamento'])
a = [('Superavaliação conhecida bruta (R$)', '=Testes_100!D10', 'Classes identificadas em 100% (NBC TA 530 item 12)'),
     ('Distorção tolerável — DT (R$)', '=Desenho_MUS!B6', ''),
     ('Folga para distorção não detectada (R$)', '=B6-B5', 'DT − conhecida'),
     ('Fator de confiança (0 erros)', '=Desenho_MUS!B10', ''),
     ('População residual (R$)', C['populacao_residual_rs'], f"Exclui {C['registros_sinalizados_excluidos']} registros sinalizados por T1/T2/T3"),
     ('Tamanho ampliado n2 = ROUNDUP(VC residual × FC ÷ folga)', '=ROUNDUP(B9*B8/B7,0)', 'Cálculo'),
     ('Intervalo J2 = VC residual ÷ n2 (R$)', '=B9/B10', 'Cálculo'),
     ('Distorções na amostra ampliada', C['ampliacao_distorcoes'], f"{C['ampliacao_itens_distintos']} itens; {C['ampliacao_itens_T4']} com incoerência de categoria (T4), sem efeito monetário; semente {C['ampliacao_semente']}"),
     ('UML residual = FC × J2 (R$)', '=B8*B11', 'Zero distorções'),
     ('Superavaliação conhecida + UML residual (R$)', '=B5+B13', ''),
     ('Comparação com DT', '=IF(B14<=B6,"<= DT: base razoável para concluir (NBC TA 530 item 15)","> DT")', ''),
     ('Conferência com o código (n2)', C['ampliacao_n'], '')]
for i, (x, y, z) in enumerate(a, 5):
    put(ws, i, [x, y, z], fmt={2: RS}, inp=(2,) if not isinstance(y, str) else ())
for r in (8,): ws.cell(row=r, column=2).number_format = FAC
for r in (10, 12, 16): ws.cell(row=r, column=2).number_format = '#,##0'

# ---------------- PT 450
ws = sheet('PT450_Distorcoes', 'PT 450 (parcela simulada) — acumulação e avaliação das distorções da base sintética', [52, 20, 16, 16, 16, 40])
head(ws, 4, ['Distorção (classe)', 'Valor (R$)', '% do trivial', '% da execução', '% da global', 'Tipo (NBC TA 450 A6) / tratamento'])
put(ws, 5, ['Materialidade global (R$)', P['materialidade_global_rs'], '', '', '', 'D-B04-01'], fmt={2: RS}, inp=(2,))
put(ws, 6, ['Materialidade de execução (R$)', P['materialidade_execucao_rs'], '', '', '', 'D-B04-01'], fmt={2: RS}, inp=(2,))
put(ws, 7, ['Claramente trivial (R$)', P['trivial_rs'], '', '', '', 'D-B04-01'], fmt={2: RS}, inp=(2,))
cl = [('Duplicidades (T1)', '=Testes_100!D5', 'Factual na simulação; acumular (A2: não trivial em conjunto)'),
      ('Bens baixados mantidos (T2)', '=Testes_100!D6', 'Factual; acumular'),
      ('Superavaliação de valor (T3+)', '=Testes_100!D7', 'Factual; acumular'),
      ('Subavaliação de valor (T3−)', '=Testes_100!D8', 'Factual; acumular (efeito oposto)'),
      ('Classificação de categoria (T4)', 0, 'Sem efeito no saldo; categoria não é divulgada nas DF — avaliar por natureza (A5)')]
for i, (x, y, z) in enumerate(cl, 9):
    put(ws, i, [x, y, f'=B{i}/$B$7', f'=B{i}/$B$6', f'=B{i}/$B$5', z], fmt={2: RS, 3: PCT, 4: PCT, 5: PCT}, link=(2,))
put(ws, 14, ['Total líquido acumulado', '=SUM(B9:B13)', '=B14/$B$7', '=B14/$B$6', '=B14/$B$5', 'Superavaliação líquida do ativo financeiro da concessão'], fmt={2: RS, 3: PCT, 4: PCT, 5: PCT}, bold=True)
put(ws, 15, ['Superavaliação bruta (T1 + T2 + T3+)', '=SUM(B9:B11)', '', '=B15/$B$6', '', 'Gatilho D-B04-01: reduzir execução se > 50% da execução'], fmt={2: RS, 4: PCT}, bold=True)
put(ws, 16, ['Gatilho D-B04-01 acionado?', '=IF(D15>0.5,"SIM — reduzir execução para 60%","NÃO")', '', '', '', ''], bold=True)
put(ws, 17, ['Margem para distorção não detectada (UML residual)', '=Ampliacao!B13', '', '=B17/$B$6', '', 'Risco de amostragem (não é distorção)'], fmt={2: RS, 4: PCT}, link=(2,))
put(ws, 18, ['Conhecida bruta + margem', '=B15+B17', '', '=B18/$B$6', '=B18/$B$5', 'Comparar com execução e global'], fmt={2: RS, 4: PCT, 5: PCT}, bold=True)
put(ws, 20, ['LAIR consolidado 2025 (R$)', 8091246000, '', '', '', 'DF25 DRE p. 36/4'], fmt={2: RS}, inp=(2,))
put(ws, 21, ['Efeito no LAIR se a contrapartida for resultado (R$)', '=B14', '', '', '=B21/B20', '% do LAIR (col. E)'], fmt={2: RS, 5: '0.00%'})
put(ws, 23, ['Conclusão (simulada)', '=IF(B14<B6,"Distorção líquida abaixo da execução e da global: não relevante isoladamente; comunicar à administração e solicitar correção (450 item 8)","Avaliar relevância")', '', '', '', ''], bold=True)

# ---------------- Comparacao
ws = sheet('Comparacao_pos', 'Comparação com o registro reservado de erros — feita APÓS o congelamento (19:55:53)', [40, 14, 14, 14, 14, 16, 16])
head(ws, 4, ['Teste × tipo de erro', 'Sinalizados', 'Verdadeiros', 'Detectados', 'Omitidos', 'Falsos positivos', 'Taxa de detecção'])
cmp = [('T1 duplicidade (nível de par)', G['T1_pares']), ('T2 baixado mantido', G['T2']), ('T3 superavaliação', G['T3_super']),
       ('T3 subavaliação', G['T3_sub']), ('T4 classificação incorreta', G['T4_classificacao'])]
for i, (nm, d) in enumerate(cmp, 5):
    put(ws, i, [nm, d['sinalizados'], d['verdadeiros'], d['detectados'], d['omitidos'], d['falsos_positivos'], f'=D{i}/C{i}'], fmt={7: PCT}, inp=(2, 3, 4, 5, 6))
put(ws, 11, ['T1: convenção "id maior = excedente" acertou o registro', G['T1_convencao_id_maior_acertou_o_excedente'], 1000, '', '', '', '=B11/C11'], fmt={7: PCT}, inp=(2,))
head(ws, 13, ['Efeito monetário por tipo', 'Estimado (R$)', 'Verdadeiro (R$)', 'Diferença (R$)', '', '', ''])
for i, (tp, d) in enumerate(G['efeito_monetario_por_tipo'].items(), 14):
    put(ws, i, [tp, d['estimado_testes_rs'], d['verdadeiro_rs'], f'=B{i}-C{i}', '', '', ''], fmt={2: RS, 3: RS, 4: RS}, inp=(2, 3))
put(ws, 19, ['Total', '=SUM(B14:B18)', '=SUM(C14:C18)', '=B19-C19', '', '', ''], fmt={2: RS, 3: RS, 4: RS}, bold=True)
mv = G['MUS_vs_verdade']
head(ws, 21, ['MUS × verdade', 'Valor (R$)', '', '', '', '', ''])
for i, (a_, b_) in enumerate([('Superavaliação verdadeira bruta', mv['superavaliacao_verdadeira_bruta_rs']),
                               ('Projeção do desenho original', mv['projecao_superavaliacao_desenho_original_rs']),
                               ('UML do desenho original', mv['UML_desenho_original_rs']),
                               ('Conhecida + UML residual', mv['conhecida_mais_UML_residual_rs'])], 22):
    put(ws, i, [a_, b_, '', '', '', '', ''], fmt={2: RS}, inp=(2,))
put(ws, 26, ['Erro da projeção (projeção − verdade)', '=B23-B22', '', '', '', '', ''], fmt={2: RS})
put(ws, 27, ['UML cobre a verdade?', '=IF(B24>=B22,"SIM","NÃO")', '', '', '', '', ''])
put(ws, 28, ['Saldo de referência sem erros (ponte)', G['ponte']['saldo_referencia_sem_erros_rs'], '', '', '', '', ''], fmt={2: RS}, inp=(2,))
put(ws, 29, ['Ponte: publicado − efeito total = referência', '=29623619000-C19-B28', '', '', '', '', 'Deve ser ~0 (centavos)'], fmt={2: RS})
put(ws, 31, ['Omissões do T4 (4): bens totalmente depreciados (residual no piso 0,01) — o residual não distingue a categoria. Limitação do teste; não foi ajustado após ver o resultado.', '', '', '', '', '', ''])

# ---------------- Congelamento
ws = sheet('Congelamento', 'Hashes congelados antes da abertura dos reservados', [48, 70])
head(ws, 4, ['Arquivo', 'SHA-256'])
put(ws, 5, ['Congelado em', K['congelado_em'] + ' (registro no Drive 23/10/2026 19:55:53 −03)'])
put(ws, 6, ['Desenho registrado no Drive', '23/10/2026 19:54:18 −03 (antes da execução às 19:54:28)'])
r = 7
for f_, h_ in K['arquivos_sha256'].items(): put(ws, r, [f_, h_]); r += 1
put(ws, r, ['Base registrada', K['base_sha256']]); r += 1
for f_, h_ in K['reservados_sha256_esperados'].items(): put(ws, r, ['Reservado: ' + f_ + ' (hash conferiu na abertura)', h_]); r += 1
os.makedirs(os.path.dirname(OUT), exist_ok=True); wb.save(OUT); print('OK', OUT)
