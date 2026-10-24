# -*- coding: utf-8 -*-
"""B13 — Conferência DF25 (PDF, transcrição do B02) × dados abertos da CVM (DFP 2025 v1, consolidado).
Uso: python b13_conferencia_cvm.py <pasta_csv_cvm> <planilha_B02> <saida_xlsx>"""
import csv, io, sys, os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
CVM = sys.argv[1] if len(sys.argv) > 1 else '/root/cpfl/fontes/cvm'
B02 = sys.argv[2] if len(sys.argv) > 2 else '/root/cpfl/producao/01_Dados_reais/CPFL_B02_Dados_Reais_2025.xlsx'
OUT = sys.argv[3] if len(sys.argv) > 3 else '/root/cpfl/producao/01_Dados_reais/CPFL_B13_Conferencia_PDF_CVM.xlsx'
ARQ = ['dfp_cia_aberta_BPA_con_2025_CPFL_FILTRADO.csv', 'dfp_cia_aberta_BPP_con_2025_CPFL_FILTRADO.csv', 'dfp_cia_aberta_DFC_MI_con_2025_CPFL_FILTRADO.csv']
rows = []
for a in ARQ:
    rows += [dict(r, ARQ=a) for r in csv.DictReader(io.open(os.path.join(CVM, a), encoding='latin-1'), delimiter=';')]
assert all(r['CNPJ_CIA'] == '02.429.144/0001-93' and r['ESCALA_MOEDA'] == 'MIL' and r['VERSAO'] == '1' for r in rows)
ano = lambda r: '2025' if r['ORDEM_EXERC'] == 'ÚLTIMO' else '2024'
idx = {}
for r in rows:
    idx.setdefault(ano(r), {}).setdefault(round(abs(float(r['VL_CONTA']))), []).append(r['CD_CONTA'] + ' ' + r['DS_CONTA'])
EXPL = {  # diferenças explicadas após análise (Bloco 13)
 ('AC12', '2024'): ('Arredondamento', '1.01.08.03.01 Outros Ativos = 1.508.407 (±1)'),
 ('Total do circulante (divulgado)', '2024'): ('Arredondamento', '1.01 = 16.551.548 (±2); soma das linhas do PDF = 16.551.547'),
 ('ANC13', '2025'): ('Arredondamento', '1.02.03 = 9.375.880 (±1)'),
 ('ANC14', '2025'): ('Arredondamento', '1.02.04 = 6.892.838 (±1)'),
 ('Total do não circulante (divulgado)', '2024'): ('Arredondamento', '1.02 = 60.577.646 (±2)'),
 ('Total do passivo circulante (divulgado)', '2025'): ('Arredondamento', '2.01 = 17.875.964 (±1)'),
 ('PL05', '2024'): ('Arredondamento', '2.03.04.04 = 3.075.192 (±1)'),
 ('Total do patrimônio líquido (divulgado)', '2025'): ('Arredondamento', '2.03 = 23.504.509 (±1)'),
 ('PC13', '2024'): ('Agregação', '2.01.01 221.712 + 2.01.05.02.06 23.253 + 2.01.05.02.07 2.638.554 + 2.01.05.02.08 100.722 = 2.984.241 (±1)'),
 ('PNC14', '2024'): ('Agregação', '2.02.02.02.07 175.914 + 2.02.02.02.08 653.079 = 828.993 (±1)'),
 ('PL atribuído aos controladores (divulgado)', '2025'): ('Agregação', '2.03 − 2.03.09 = 23.504.509 − 1.014.735 = 22.489.774 (±1)'),
 ('PL atribuído aos controladores (divulgado)', '2024'): ('Agregação', '2.03 − 2.03.09 = 21.795.029 − 1.063.267 = 20.731.762'),
 ('FC06', '2025'): ('Subtotal derivado', '6.01.01 13.055.906 + 6.01.02 (2.249.102) = 10.806.804'),
 ('FC06', '2024'): ('Subtotal derivado', '6.01.01 12.470.994 + 6.01.02 (620.383) = 11.850.611'),
}
wb_in = load_workbook(B02, data_only=True)
res = []
for sh in ['BP_Ativo', 'BP_Passivo_PL', 'DFC']:
    for r in wb_in[sh].iter_rows(min_row=5, values_only=True):
        if r[1] is None or not isinstance(r[3], (int, float)) or 'calculado' in str(r[1]) or 'Diferença' in str(r[1]): continue
        for y, v in (('2025', r[3]), ('2024', r[4])):
            if not isinstance(v, (int, float)): continue
            hit = idx[y].get(round(abs(v)))
            if hit: res.append((sh, r[0] or '', r[1], y, v, 'Idêntico', hit[0]))
            else:
                k = (r[0] or r[1], y); c, e = EXPL.get(k, ('SEM EXPLICAÇÃO', ''))
                res.append((sh, r[0] or '', r[1], y, v, c, e))
wb = Workbook(); ws = wb.active; ws.title = 'Conferencia'
H = Font(name='Arial', bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3864'); th = Side(style='thin', color='A6A6A6'); BD = Border(left=th, right=th, top=th, bottom=th)
ws['A1'] = 'Conferência DF25 (PDF) × CVM — DFP 2025 v1, documento 155108 (recebido em 05/03/2026), consolidado, R$ mil'; ws['A1'].font = Font(name='Arial', bold=True, size=12)
ws['A2'] = 'Fonte CVM: dados abertos (dfp_cia_aberta_2025.zip), arquivos BPA_con, BPP_con e DFC_MI_con filtrados pelo CNPJ 02.429.144/0001-93. PDF: planilha B02 (DF25 p. 34–35 e 39). Preparado em 24/10/2026 (assistido por IA).'
ws['A2'].font = Font(name='Arial', italic=True, size=9)
cols = ['Demonstração (B02)', 'ID', 'Linha', 'Exercício', 'Valor PDF', 'Classe', 'Conta CVM / explicação']
for i, c in enumerate(cols, 1):
    x = ws.cell(row=4, column=i, value=c); x.font = H; x.fill = HF; x.border = BD
for j, row in enumerate(res, 5):
    for i, v in enumerate(row, 1):
        x = ws.cell(row=j, column=i, value=v); x.border = BD; x.font = Font(name='Arial', size=9)
        if i == 5: x.number_format = '#,##0;(#,##0)'
n = 5 + len(res) + 1
ws.cell(row=n, column=3, value='Resumo (fórmulas)').font = Font(name='Arial', bold=True)
for k, cl in enumerate(['Idêntico', 'Arredondamento', 'Agregação', 'Subtotal derivado', 'SEM EXPLICAÇÃO'], 1):
    ws.cell(row=n + k, column=3, value=cl); ws.cell(row=n + k, column=4, value=f'=COUNTIF(F5:F{4 + len(res)},"{cl}")')
ws.cell(row=n + 6, column=3, value='Total comparado'); ws.cell(row=n + 6, column=4, value=f'=COUNTA(F5:F{4 + len(res)})')
for c, w in zip('ABCDEFG', [16, 8, 44, 10, 14, 18, 70]): ws.column_dimensions[c].width = w
wb.save(OUT)
from collections import Counter
print('total', len(res), Counter(x[5] for x in res)); print('OK', OUT)
