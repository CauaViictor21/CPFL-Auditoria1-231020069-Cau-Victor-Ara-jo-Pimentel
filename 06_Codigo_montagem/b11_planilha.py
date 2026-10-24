# -*- coding: utf-8 -*-
"""Planilha do Bloco 11: PT 265 (matriz de deficiências), PT 450 (consolidação final), PT 500 (matriz de evidência)
e PT 260 (checklist de comunicações). Separa fatos da companhia (evidência pública), hipóteses dependentes de acesso
e resultados exclusivamente sintéticos."""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
OUT = sys.argv[1] if len(sys.argv) > 1 else '/root/cpfl/producao/02_Planejamento/CPFL_B11_Comunicacoes_Distorcoes_Evidencia.xlsx'
F = 'Arial'; H = Font(name=F, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3864')
TT = Font(name=F, bold=True, size=12); N = Font(name=F); B = Font(name=F, bold=True); BL = Font(name=F, color='0000FF')
INP = PatternFill('solid', fgColor='FFF2CC'); th = Side(style='thin', color='A6A6A6'); BD = Border(left=th, right=th, top=th, bottom=th)
MIL = '#,##0;(#,##0);"-"'; PCT = '0.0%;(0.0%);"-"'
CAT = {'R': PatternFill('solid', fgColor='E2EFDA'), 'H': PatternFill('solid', fgColor='FFF2CC'), 'S': PatternFill('solid', fgColor='DDEBF7'), 'X': PatternFill('solid', fgColor='EDEDED')}
wb = Workbook(); wb.remove(wb.active)
def sheet(name, title, widths, nota):
    ws = wb.create_sheet(name); ws['A1'] = title; ws['A1'].font = TT
    ws['A2'] = nota; ws['A2'].font = Font(name=F, italic=True, size=9)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    return ws
def head(ws, r, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=r, column=i, value=c); x.font = H; x.fill = HF; x.border = BD; x.alignment = Alignment(wrap_text=True)
def put(ws, r, vals, fmt=None, inp=(), bold=False, fill=None):
    for i, v in enumerate(vals, 1):
        x = ws.cell(row=r, column=i, value=v); x.border = BD; x.font = B if bold else N
        x.alignment = Alignment(wrap_text=True, vertical='top')
        if fmt and i in fmt: x.number_format = fmt[i]
        if i in inp: x.font = BL; x.fill = INP
        if fill: x.fill = fill
LEG = 'Categorias: R = real, sustentada por evidência pública; H = hipótese dependente de acesso; S = exclusivamente sintética (base simulada); X = erro de extração/arredondamento do preparador (não é da companhia).'

# ============ PT 265
ws = sheet('PT265_Matriz', 'PT 265 — Matriz de deficiências de controle (condição, critério, causa, efeito, recomendação, evidência)', [7, 6, 34, 30, 26, 30, 32, 34, 26], LEG)
head(ws, 4, ['Ref.', 'Cat.', 'Condição', 'Critério', 'Causa (provável)', 'Efeito', 'Recomendação', 'Evidência / fonte', 'Avaliação de significância (item 8)'])
d = [
 ('DC-01', 'S', 'Base sintética: 1.000 bens duplicados, 1.000 baixados com valor, 2.000 valores divergentes do recálculo e 996 categorias incoerentes', 'Cadastro deve refletir bens existentes, ativos e corretamente classificados (CPC 47/ICPC 01; MCSE)', 'Simulada pelo desenho do exercício', 'Superavaliação líquida simulada de R$ 117,5 mi (38,7% da execução)', 'Na vida real: controles de unitização, baixa tempestiva e chave única por bem; conciliação cadastro × razão', 'PT 530 e PT 450; CPFL_B08_Amostragem_Testes_Distorcoes.xlsx', 'Seria significativa em ambiente real (volume e reincidência). NÃO é deficiência da CPFL: resultado só da simulação'),
 ('DC-02', 'H', 'Diferença de +R$ 150 mi no VNR 2025 da CPFL Paulista sem explicação pública (+12,2% no consolidado)', 'Estimativa deve seguir método e premissas consistentes (CPC 46; nota 4)', 'Desconhecida: possíveis ajustes de base, índices ou falha de controle sobre o cálculo', 'Possível superavaliação de ativo e receita; ou explicação legítima', 'Indagar e obter o cálculo; testar controles de revisão do VNR', 'PT 520 e PT 540; DFs individuais notas 11; DF25 nota 26.2', 'Não avaliável sem acesso; NÃO classificada como deficiência'),
 ('DC-03', 'H', 'Reconhecimento do débito do IRRF Jantus (139.291) em 31/12/2025 não demonstrado nas notas', 'Obrigação presente após decisão definitiva deve ser reconhecida (CPC 25; CPC 24)', 'Desconhecida', 'Se não reconhecido: passivo subavaliado (46% da execução)', 'Verificar rubrica e lançamento; controles de monitoramento jurídico-fiscal', 'DF25 nota 22 p. 96; nota 10 p. 68', 'Não avaliável sem acesso'),
 ('DC-04', 'R', 'Divulgação da renovação das concessões em 2025 (nota 1.3) sem quantificar o efeito possível sobre a classificação de ~99,9% do ativo financeiro da concessão', 'Divulgar fontes de incerteza de estimativa com risco de ajuste relevante (CPC 26 item 125)', 'Julgamento da administração (aditivos ainda não assinados)', 'Usuário não dimensiona o efeito que se materializou em 2026 (29.599.494 reclassificados)', 'Discutir com a governança a suficiência da divulgação (PT 260)', 'DF25 nota 1.3 p. 44; ITR 2T26 nota 11 p. 49', 'Ponto de divulgação, não de controle interno; NÃO caracterizado como deficiência de controle'),
 ('DC-05', 'H', 'Controles de TI sobre medição e faturamento (R15) e sobre a receita não faturada (R03) não testados', 'NBC TA 315 (R2); NBC TA 330 itens 8 e 18', 'Ausência de acesso do aluno', '—', 'Testar controles gerais de TI e de aplicação', 'PT 315; PT 330', 'Ausência de acesso NÃO é deficiência da companhia'),
 ('DC-06', 'R', 'Divergências de transcrição corrigidas (DV01, DV11–DV13, DV17, DV18)', 'Documentação de auditoria exata (NBC TA 230)', 'Erro do preparador (leitura de imagens)', 'Nenhum nas DF da companhia', 'Manter dupla conferência e conciliação por fórmula', 'Apêndice A do dossiê', 'Controle de qualidade do trabalho (PT 220), não da companhia'),
]
for i, row in enumerate(d, 5): put(ws, i, list(row), fill=CAT[row[1]])
put(ws, 12, ['', '', 'Conclusão: nenhuma deficiência significativa de controle interno da CPFL foi identificada com evidência pública; DC-02 e DC-03 dependem de acesso; DC-01 é sintética.', '', '', '', '', '', ''], bold=True)

# ============ PT 450 consolidado
ws = sheet('PT450_Consolidado', 'PT 450 — Consolidação final das distorções e divergências (sem duplicar exceções e projeções)', [7, 6, 40, 16, 22, 16, 16, 30, 34], LEG + ' Valores em R$ mil.')
head(ws, 4, ['Ref.', 'Cat.', 'Item', 'Afirmação', 'Tipo (A6) / natureza', 'Valor (R$ mil)', '% execução', 'Correção', 'Avaliação'])
x = [
 ('DS-01', 'S', 'Base sintética: duplicidades (T1)', 'Existência', 'Factual na simulação', 59975.20, 'Simulada (não houve correção real)', 'Acumulada só na simulação'),
 ('DS-02', 'S', 'Base sintética: bens baixados mantidos (T2)', 'Existência', 'Factual na simulação', 56070.25, 'Simulada', 'Idem'),
 ('DS-03', 'S', 'Base sintética: superavaliação de valor (T3+)', 'Avaliação', 'Factual na simulação', 8250.70, 'Simulada', 'Idem'),
 ('DS-04', 'S', 'Base sintética: subavaliação de valor (T3−)', 'Avaliação', 'Factual na simulação', -6811.81, 'Simulada', 'Não compensa a superavaliação na avaliação bruta'),
 ('DS-05', 'S', 'Base sintética: classificação de categoria (T4)', 'Classificação', 'Por natureza (A5)', 0, 'Simulada', 'Sem efeito no saldo'),
 ('DS-06', 'S', 'Margem para distorção não detectada (UML residual do MUS)', 'Existência/avaliação', 'Risco de amostragem (não é distorção)', 178890.80, '—', 'Não somar como distorção; usar só na avaliação de suficiência'),
 ('DH-01', 'H', 'IRRF Jantus: débito de 139.291 em 31/12/2025 (reconhecimento não demonstrado)', 'Integridade (passivo)', 'Potencial factual — pendente', 139291, 'Não aplicável (não verificado)', 'Não acumulado: hipótese. Se confirmado o não reconhecimento: 46% da execução'),
 ('DH-02', 'H', 'VNR 2025 acima da expectativa (+139.479; Paulista +150.451)', 'Avaliação', 'Diferença analítica não explicada (não é distorção)', 139479, '—', 'Não acumulado: não há base para afirmar distorção (PT 520/540)'),
 ('DQ-01', 'R', 'Renovação: divulgação sem quantificação do efeito possível (nota 1.3)', 'Divulgação', 'Qualitativa / julgamento', None, 'Discussão com a governança (planejada)', 'Não caracterizada como distorção relevante; ponto de comunicação'),
 ('DX-01', 'X', 'Erros de transcrição do preparador (DV01, DV11–DV13, DV17, DV18)', '—', 'Extração (não é da companhia)', None, 'Corrigidos no dossiê', 'Excluídos da acumulação'),
 ('DX-02', 'X', 'Diferenças de soma ±1 a ±3 (arredondamento, nota 2.4)', '—', 'Arredondamento', None, '—', 'Excluídas da acumulação'),
]
for i, (a, c, it, af, tp, v, cor, av) in enumerate(x, 5):
    put(ws, i, [a, c, it, af, tp, v, (f'=F{i}/303421.73' if isinstance(v, (int, float)) else ''), cor, av], fmt={6: MIL, 7: PCT}, inp=(6,) if isinstance(v, (int, float)) else (), fill=CAT[c])
r = 5 + len(x) + 1
put(ws, r, ['', '', 'Distorções REAIS da companhia acumuladas (evidência sustentada)', '', '', 0, '=F{0}/303421.73'.format(r), '', 'Nenhuma distorção da CPFL foi sustentada por evidência pública'], fmt={6: MIL, 7: PCT}, bold=True, fill=CAT['R']); r += 1
put(ws, r, ['', '', 'Distorção sintética líquida (DS-01 a DS-05)', '', '', '=SUM(F5:F9)', f'=F{r}/303421.73', '', 'Igual ao PT 530 (117.484); não é constatação sobre a CPFL'], fmt={6: MIL, 7: PCT}, bold=True, fill=CAT['S']); r += 1
put(ws, r, ['', '', 'Superavaliação sintética bruta (DS-01 + DS-02 + DS-03)', '', '', '=F5+F6+F7', f'=F{r}/303421.73', '', 'Gatilho D-B04-01 (50%): não acionado'], fmt={6: MIL, 7: PCT}, bold=True, fill=CAT['S']); r += 1
put(ws, r, ['', '', 'Hipóteses pendentes (DH-01 + DH-02) — apenas para dimensionar o risco remanescente', '', '', '=F11+F12', f'=F{r}/303421.73', '', 'Mesmo somadas, < materialidade global (404.562); não acumuladas'], fmt={6: MIL, 7: PCT}, bold=True, fill=CAT['H'])

# ============ PT 500 matriz
ws = sheet('PT500_Matriz', 'PT 500 — Matriz de evidência por risco/afirmação: relevância, confiabilidade, suficiência e lacunas', [7, 30, 22, 40, 14, 14, 14, 40, 14], 'Suficiência: S = suficiente para a conclusão limitada do exercício; P = parcial; I = insuficiente (procedimento permanece planejado).')
head(ws, 4, ['Risco', 'Assunto', 'Afirmações', 'Evidência obtida (fonte)', 'Relevância', 'Confiabilidade', 'Suficiência', 'Lacuna / procedimento planejado', 'PT'])
m = [
 ('R01', 'Ativo financeiro da concessão', 'Avaliação, classificação', 'Movimentação reexecutada; saldos por distribuidora; analítico VNR; sensibilidade; renovação (ITR 2T26); testes na base sintética', 'Alta', 'Alta (DF auditadas) / sintética: nenhuma sobre a CPFL', 'P', 'Base homologada, laudo, cálculo da Paulista, cadastro real', '500/520/530/540'),
 ('R02', 'Ativo contratual', 'Avaliação, exatidão', 'Saldos e PAA 1 da PwC; programa PG-ESTIM', 'Média', 'Alta', 'I', 'Modelo de remuneração das transmissoras; medições de obras', '315/330/540'),
 ('R03', 'Receita não faturada', 'Ocorrência, corte', 'Saldo 1.763.858 e PAA 2; analíticos de receita', 'Média', 'Alta', 'I', 'Controles de medição e faturamento; recálculo por classe', '315/330/520'),
 ('R04', 'Demais receitas', 'Ocorrência, exatidão', 'Variações por componente; RTA 2025 (nota 26.2)', 'Média', 'Alta', 'P', 'Testes de detalhe em faturas', '520'),
 ('R06', 'Dívidas e debêntures', 'Integridade, classificação', 'Notas 18–19; vencimentos (nota 34 g); captações de 2026', 'Alta', 'Alta', 'P', 'Confirmações CF-04 (não enviadas); escrituras', '505/570'),
 ('R07', 'Provisões e contingências', 'Integridade, avaliação', 'Movimentação conciliada; casos individuais; IRRF Jantus', 'Alta', 'Média (prognósticos da administração)', 'P', 'Carta jurídica CF-08 (não enviada)', '501/560'),
 ('R13', 'Partes relacionadas', 'Integridade, divulgação', 'Nota 31 conciliada com o BP; política; comunicados 2026', 'Alta', 'Média (lista da própria entidade)', 'P', 'Atas, contratos, FR íntegro, CF-02/03', '550/505'),
 ('R14', 'Consolidação', 'Integridade, apresentação', 'Mapa do grupo; segmentos conciliados; julgamento Chapecoense', 'Alta', 'Alta', 'P', 'Instruções aos componentes; eliminações', '600/501'),
 ('R16', 'Burla de controles', 'Todas', 'Indicador de viés no VNR (não conclusivo); análise de estimativas', 'Média', 'Média', 'I', 'Teste de lançamentos manuais (NBC TA 240 item 32)', '240/540'),
 ('R17', 'Continuidade', 'Divulgação', 'Liquidez, vencimentos, release (2,30x), captações', 'Alta', 'Média (release não auditado)', 'P', 'Previsão de caixa; recálculo contratual de covenants', '570'),
 ('—', 'Saldos de abertura', 'Todas', '61 linhas DF24 × DF25 sem diferença', 'Alta', 'Alta', 'S', 'Papéis do antecessor (não acessíveis)', '510'),
 ('—', 'Eventos subsequentes', 'Integridade, divulgação', 'Notas 37/22; ITRs 1T26 e 2T26; RI até 23/10/2026', 'Alta', 'Alta', 'P', 'Indagações, atas, representação formal', '560/580'),
]
for i, row in enumerate(m, 5): put(ws, i, list(row))
put(ws, 18, ['', 'Leitura do relatório da PwC', '', 'Usada como contexto (PAA, opinião), nunca como substituta de evidência própria', '', '', '', '', '200/700'], bold=True)

# ============ PT 260 checklist
ws = sheet('PT260_Checklist', 'PT 260 — Assuntos a comunicar à governança e onde estão documentados (comunicação NÃO enviada)', [10, 48, 40, 30], 'Destinatário pretendido: Comitê de Auditoria / Conselho de Administração. Nenhuma comunicação foi enviada.')
head(ws, 4, ['Item 260', 'Assunto', 'Conteúdo no exercício', 'Origem'])
c = [('14', 'Responsabilidades do auditor', 'Trabalho acadêmico; sem opinião; base NBC TA 200', 'PT 200/210'),
     ('15', 'Alcance e época planejados', 'Perímetro consolidado; materialidade; riscos significativos R01, R02, R03, R16; cronograma', 'PT 300/315/320'),
     ('16(a)', 'Aspectos qualitativos das práticas contábeis', 'VNR (nível 3); divulgação da renovação; julgamento de controle conjunto da Chapecoense', 'PT 540/600'),
     ('16(b)', 'Dificuldades significativas', 'Falta de acesso a sistemas, pessoas e terceiros; FR truncado; relatório 2024 ausente do PDF', 'PT 500/510'),
     ('16(c)', 'Assuntos significativos discutidos com a administração', 'Nenhuma discussão ocorreu (planejado: VNR Paulista, IRRF Jantus)', 'PT 520/560'),
     ('16(d)', 'Representações solicitadas', 'Minuta de carta de representação (não enviada)', 'PT 580'),
     ('16(e)', 'Outros assuntos relevantes', 'Partes relacionadas (mútuo SGEL); continuidade (CCL negativo)', 'PT 550/570'),
     ('17', 'Independência (entidade listada)', 'Declaração do aluno: sem vínculo com o Grupo CPFL', 'PT 200'),
     ('265 item 9', 'Deficiências significativas por escrito', 'Nenhuma deficiência significativa da CPFL identificada com evidência pública', 'PT 265'),
     ('450 item 12', 'Distorções não corrigidas', 'Nenhuma real; distorções sintéticas apresentadas apenas como resultado didático', 'PT 450')]
for i, row in enumerate(c, 5): put(ws, i, list(row))
os.makedirs(os.path.dirname(OUT), exist_ok=True); wb.save(OUT); print('OK', OUT)
