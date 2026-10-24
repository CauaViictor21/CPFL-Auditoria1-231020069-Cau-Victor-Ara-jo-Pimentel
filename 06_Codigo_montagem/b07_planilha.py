# -*- coding: utf-8 -*-
"""Planilha editável do Bloco 07: conciliação, verificações e prévia da base SINTÉTICA v2."""
import json, gzip, csv, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

B = '/root/cpfl/base_v2'
OUT = sys.argv[1] if len(sys.argv) > 1 else '/root/cpfl/producao/03_Base_sintetica/CPFL_B07_Conciliacao_Verificacoes_Base.xlsx'
cfg = json.load(open(f'{B}/config_base_cpfl_v2.json', encoding='utf-8'))
ver = json.load(open(f'{B}/saida/CPFL_SINTETICA_verificacoes.json', encoding='utf-8'))
resu = json.load(open(f'{B}/saida/CPFL_SINTETICA_resumo_execucao.json', encoding='utf-8'))
limpa = json.load(open('/tmp/claude-0/pasta_limpa_b07/saida/CPFL_SINTETICA_resumo_execucao.json', encoding='utf-8'))

F = 'Arial'
H = Font(name=F, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3864')
T = Font(name=F, bold=True, size=12); N = Font(name=F); BOLD = Font(name=F, bold=True)
BLUE = Font(name=F, color='0000FF'); GREEN = Font(name=F, color='008000')
INP = PatternFill('solid', fgColor='FFF2CC'); thin = Side(style='thin', color='A6A6A6')
BD = Border(left=thin, right=thin, top=thin, bottom=thin)
RS = '#,##0.00;(#,##0.00);"-"'; MIL = '#,##0;(#,##0);"-"'; PCT = '0.0%'

wb = Workbook()
def sheet(name, title, widths):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = T
    ws['A2'] = 'Base SINTÉTICA para fins didáticos — não são dados internos da CPFL Energia S.A. Preparado com apoio de IA em 23/10/2026.'
    ws['A2'].font = Font(name=F, italic=True, size=9)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    return ws
def head(ws, r, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=r, column=i, value=c); x.font = H; x.fill = HF; x.border = BD
        x.alignment = Alignment(wrap_text=True, vertical='center')
def put(ws, r, vals, fonts=None, fmts=None):
    for i, v in enumerate(vals, 1):
        x = ws.cell(row=r, column=i, value=v); x.border = BD; x.font = N
        x.alignment = Alignment(wrap_text=True, vertical='top')
        if fonts and fonts.get(i): x.font = fonts[i]
        if fmts and fmts.get(i): x.number_format = fmts[i]

wb.remove(wb.active)
# 0 Leia-me
ws = sheet('Leia_me', 'Bloco 07 — Conciliação e verificações da base sintética v2 (ativo financeiro da concessão)', [110])
txt = [
    'Conteúdo: Conciliacao (base × saldos publicados), Distrib_Categoria (abertura), Dicionario (definitivo), Verificacoes (integridade), Execucao (ambiente, comandos, hashes, reprodução em pasta limpa), Previa_1000 (primeiras 1.000 linhas).',
    'Legenda de cores: texto azul com fundo amarelo = dado de entrada (saldo publicado ou total calculado pelo verificador, com fonte); preto = fórmula; verde = vínculo a outra aba.',
    'A base integral (501.000 linhas) está em CPFL_SINTETICA_Base_AFC_2025_registrada.csv.gz; não é carregada nesta planilha.',
    'Os arquivos de referência sem erros, registro de erros e ponte estão em pasta reservada e NÃO foram lidos neste bloco; só os hashes estão registrados (aba Execucao). Serão abertos após o congelamento dos testes do Bloco 08.',
    'Modalidade: Simulado para fins didáticos (base). Os saldos por distribuidora são reais (fontes na aba Conciliacao).',
]
for i, t in enumerate(txt, 4): ws.cell(row=i, column=1, value=t).font = N; ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True)

# 1 Conciliacao
ws = sheet('Conciliacao', 'Conciliação da base registrada com os saldos publicados (31/12/2025)', [22, 18, 20, 20, 16, 12, 60])
head(ws, 4, ['Distribuidora', 'Saldo publicado (R$ mil)', 'Saldo publicado (R$) = mil × 1.000', 'Total da base (R$)', 'Diferença (R$)', 'Linhas na base', 'Fonte do saldo publicado'])
r = 5
for d, v in cfg['saldos_publicados_rs_mil'].items():
    pdv = ver['por_distribuidora'][d]
    put(ws, r, [d, v['saldo'], f'=B{r}*1000', pdv['total_rs'], f'=D{r}-C{r}', pdv['linhas'], v['fonte']],
        fonts={2: BLUE, 4: BLUE, 6: BLUE}, fmts={2: MIL, 3: RS, 4: RS, 5: RS, 6: MIL})
    ws.cell(row=r, column=2).fill = INP; ws.cell(row=r, column=4).fill = INP; ws.cell(row=r, column=6).fill = INP
    r += 1
put(ws, r, ['Total das distribuidoras', f'=SUM(B5:B{r-1})', f'=SUM(C5:C{r-1})', f'=SUM(D5:D{r-1})', f'=SUM(E5:E{r-1})', f'=SUM(F5:F{r-1})', ''],
    fonts={i: BOLD for i in range(1, 8)}, fmts={2: MIL, 3: RS, 4: RS, 5: RS, 6: MIL}); TOT = r; r += 1
put(ws, r, ['Saldo consolidado (BP)', cfg['saldo_consolidado_rs_mil']['saldo'], f'=B{r}*1000', '', '', '', cfg['saldo_consolidado_rs_mil']['fonte']],
    fonts={2: BLUE}, fmts={2: MIL, 3: RS}); ws.cell(row=r, column=2).fill = INP; CONS = r; r += 1
put(ws, r, ['Diferença distribuidoras − consolidado', f'=B{TOT}-B{CONS}', f'=C{TOT}-C{CONS}', '', '', '', 'Deve ser zero (Santa Cruz derivada por diferença)'], fmts={2: MIL, 3: RS}); r += 1
put(ws, r, ['Diferença base − BP consolidado (R$)', '', '', f'=D{TOT}-C{CONS}', '', '', 'Critério: R$ 0,00'], fonts={4: BOLD}, fmts={4: RS}); r += 1
put(ws, r, ['Ponte de escala', '', '', '', '', '', 'R$ mil × 1.000 = R$. Base em reais com 2 casas; somas feitas em centavos inteiros pelo verificador; arredondamento por maior resto (≤ 1 centavo por linha).']); r += 2
head(ws, r, ['Prova da Santa Cruz (R$ mil)', 'Valor', '', '', '', '', 'Fonte']); r += 1
sc0 = r
rows = [('Saldo 2024 consolidado', 25209768, 'DF25 nota 11 p. 68'),
        ('(−) Paulista 2024', -10688700, 'DF Paulista nota 11 p. 29'), ('(−) RGE 2024', -10927569, 'DF RGE nota 11 p. 31'),
        ('(−) Piratininga 2024', -3343565, 'DF Piratininga nota 11 p. 29')]
for a, b, c in rows:
    put(ws, r, [a, b, '', '', '', '', c], fonts={2: BLUE}, fmts={2: MIL}); ws.cell(row=r, column=2).fill = INP; r += 1
put(ws, r, ['Santa Cruz 2024 (por diferença)', f'=SUM(B{sc0}:B{r-1})', '', '', '', '', 'Cálculo'], fonts={1: BOLD, 2: BOLD}, fmts={2: MIL}); s24 = r; r += 1
mov = [('(+) Ajuste ao valor justo', 1285173, -(635251 + 469006 + 171723)), ('(+) Transferência do ativo contratual', 3368915, -(1383147 + 1458964 + 348021)),
       ('(−) Transferência para o intangível', -12576, (3072 + 6118 + 2397)), ('(−) Baixas', -227661, (86789 + 115187 + 24311))]
m0 = r
for a, c, p in mov:
    put(ws, r, [a + ' — consolidado menos Paulista, RGE e Piratininga', f'={c}+({p})', '', '', '', '', 'DF25 nota 11 p. 68; notas 11 das três DFs individuais'], fmts={2: MIL}); r += 1
put(ws, r, ['Santa Cruz 2025 reconstruída pela movimentação', f'=B{s24}+SUM(B{m0}:B{r-1})', '', '', '', '', 'Cálculo'], fonts={1: BOLD, 2: BOLD}, fmts={2: MIL}); rec = r; r += 1
put(ws, r, ['Diferença vs. saldo usado (435.548)', f'=B{rec}-B8', '', '', '', '', 'Esperado ±1 (arredondamento, DF25 nota 2.4 p. 46)'], fmts={2: MIL})

# 2 Distribuidora x Categoria
ws = sheet('Distrib_Categoria', 'Abertura da base registrada por distribuidora e categoria', [22, 24, 12, 22, 12])
head(ws, 4, ['Distribuidora', 'Categoria cadastrada', 'Linhas', 'Total (R$)', '% do total'])
r = 5
for it in ver['por_distribuidora_categoria']:
    put(ws, r, [it['distribuidora'], it['categoria'], it['linhas'], it['total_rs'], f'=D{r}/$D$26'], fonts={3: BLUE, 4: BLUE}, fmts={3: MIL, 4: RS, 5: PCT}); r += 1
put(ws, 25, ['', '', '', '', ''])
put(ws, 26, ['Total', '', '=SUM(C5:C24)', '=SUM(D5:D24)', '=SUM(E5:E24)'], fonts={i: BOLD for i in range(1, 6)}, fmts={3: MIL, 4: RS, 5: PCT})
put(ws, 27, ['Confere com Conciliacao (deve ser 0)', '', f"=C26-Conciliacao!F{TOT}", f"=D26-Conciliacao!D{TOT}", ''], fonts={3: GREEN, 4: GREEN}, fmts={3: MIL, 4: RS})
head(ws, 29, ['Resumo por categoria', '', 'Linhas', 'Total (R$)', '% do total'])
for i, c in enumerate(cfg['categorias'], 30):
    put(ws, i, [c, '', f'=SUMIF($B$5:$B$24,A{i},$C$5:$C$24)', f'=SUMIF($B$5:$B$24,A{i},$D$5:$D$24)', f'=D{i}/$D$26'], fmts={3: MIL, 4: RS, 5: PCT})

# 3 Dicionario
ws = sheet('Dicionario', 'Dicionário de dados definitivo — base sintética v2', [28, 18, 62, 24])
head(ws, 4, ['Campo', 'Tipo / formato', 'Definição e domínio', 'Natureza'])
dic = [
    ('id_registro', 'texto', 'Chave primária. AFC-SIM-0000001 a AFC-SIM-0501000, atribuída após embaralhamento (não revela origem nem erro).', 'Sintético'),
    ('id_bem_tecnico', 'texto', 'BEM-SIM-nnnnnnn. Identifica o bem físico; deveria ser único. Repetição = indício de duplicidade cadastral.', 'Sintético'),
    ('distribuidora', 'texto', 'CPFL Paulista | RGE | CPFL Piratininga | CPFL Santa Cruz. Linhas proporcionais ao saldo publicado.', 'Nome real; alocação sintética'),
    ('categoria_cadastrada', 'texto', 'Redes de distribuicao | Transformadores | Postes e estruturas | Medidores | Subestacoes (sem acento no arquivo).', 'Sintético'),
    ('data_incorporacao', 'data AAAA-MM-DD', 'De max(1985; 2025 − vida útil + 1) a 2025; dia 1–28.', 'Sintético'),
    ('vnr_base_simulado_rs', 'decimal (2 casas), R$', 'VNR unitário da categoria × U(0,35; 1,75). Não é VNR homologado pela ANEEL.', 'Sintético'),
    ('fator_atualizacao_simulado', 'decimal (4 casas)', 'U(1,00; 1,30). Substituto didático do IPCA acumulado desde a avaliação (DF25 nota 4 p. 59).', 'Sintético'),
    ('fator_residual_simulado', 'decimal (4 casas)', '1 − idade/vida útil da categoria verdadeira (mínimo 0,0100). Idade = (31/12/2025 − data de incorporação)/365,25.', 'Sintético, dependente'),
    ('status_cadastral', 'texto', 'ATIVO | BAIXADO.', 'Sintético'),
    ('data_baixa', 'data ou vazio', 'Vazio quando ATIVO (vazio justificado); data de 2025 quando BAIXADO.', 'Sintético'),
    ('valor_registrado_rs', 'decimal (2 casas), R$', 'SALDO em 31/12/2025 (estoque, não fluxo) = VNR × fator × residual × fator de escala da distribuidora, arredondado ao centavo por maior resto.', 'Sintético; total conciliado'),
    ('data_base', 'data', '2025-12-31 em todas as linhas.', 'Constante'),
    ('origem_dados', 'texto', 'SINTETICO em todas as linhas.', 'Rótulo obrigatório'),
]
for i, d in enumerate(dic, 5): put(ws, i, list(d))
r = 5 + len(dic) + 1
head(ws, r, ['Categoria', 'Peso nas linhas', 'VNR unitário base (R$) / vida útil (anos)', 'Origem']); r += 1
for c, v in cfg['categorias'].items():
    put(ws, r, [c, v['peso'], f"{v['vnr_base_unitario_rs']:,.0f}".replace(',', '.') + f" / {v['vida_util_anos']}", 'Arbitrado (peso e VNR da etapa 1; vida útil v2)'], fmts={2: '0%'}); r += 1

# 4 Verificacoes
ws = sheet('Verificacoes', 'Relatório de verificações de integridade (leitura do arquivo salvo)', [8, 52, 22, 20, 46])
head(ws, 4, ['Ref.', 'Verificação', 'Resultado', 'Esperado', 'Classificação'])
fx = ver['falhas_de_geracao']
chk = [
    ('V01', 'Cabeçalho = 13 campos do dicionário', 'OK' if ver['cabecalho_confere'] else 'FALHA', 'OK', 'Integridade de geração'),
    ('V02', 'Linhas de dados (sem cabeçalho)', ver['linhas_dados'], 501000, 'Integridade de geração'),
    ('V03', 'IDs de registro únicos', ver['ids_registro_unicos'], 501000, 'Integridade de geração'),
    ('V04', 'Falhas de formato, domínio, datas, precisão, intervalos (F01–F16)', sum(fx.values()), 0, 'Integridade de geração'),
    ('V05', 'Campos vazios fora de data_baixa', sum(v for k, v in ver['campos_vazios'].items() if k != 'data_baixa'), 0, 'Integridade de geração'),
    ('V06', 'data_baixa vazia (ATIVO)', ver['campos_vazios'].get('data_baixa', 0), ver['status'].get('ATIVO', 0), 'Vazio justificado'),
    ('V07', 'Total da base − saldo do BP (R$)', ver['diferenca_total_rs'], 0, 'Conciliação'),
    ('V08', 'Diferenças por distribuidora (soma dos módulos, R$)', sum(abs(v['diferenca_rs']) for v in ver['por_distribuidora'].values()), 0, 'Conciliação'),
    ('V09', 'Mediana de valor/(VNR×fator×residual) ≈ fator de escala (máx. desvio absoluto)', max(abs(ver['mediana_valor_sobre_vnr_fator_residual'][d] - ver['fator_escala_publicado_no_resumo'][d]) for d in ver['por_distribuidora']), '< 0,001', 'Consistência econômica agregada'),
    ('V10', 'id_bem_tecnico repetidos (pares)', ver['ids_bem_tecnico_repetidos'], 'Exceção prevista no desenho', 'EXCEÇÃO INTENCIONAL — investigar no PT 530 (não é falha de geração)'),
    ('V11', 'Registros BAIXADO com valor positivo', ver['status'].get('BAIXADO', 0), 'Exceção prevista no desenho', 'EXCEÇÃO INTENCIONAL — investigar no PT 530'),
    ('V12', 'Hash da base: execução original = execução em pasta limpa', 'IGUAL' if resu['sha256'] == limpa['sha256'] else 'DIFERENTE', 'IGUAL', 'Reprodutibilidade'),
]
for i, c in enumerate(chk, 5):
    put(ws, i, list(c), fmts={3: '#,##0.######', 4: '#,##0'})
r = 5 + len(chk) + 1
ws.cell(row=r, column=1, value='Nota: recálculo individual do valor e coerência residual × idade × vida útil (por registro) ficam para os testes do Bloco 08, para não antecipar a detecção. Super/subavaliação e classificação incorreta não são visíveis nas verificações estruturais.').font = Font(name=F, italic=True, size=9)

# 5 Execucao
ws = sheet('Execucao', 'Registro de execução, ambiente e integridade dos arquivos', [44, 80])
head(ws, 4, ['Item', 'Registro'])
ex = [
    ('Gerador', 'gerar_base_cpfl_v2.py (versão 2.0) + config_base_cpfl_v2.json; biblioteca padrão do Python'),
    ('Verificador', 'verificar_base_cpfl_v2.py (lê só a base registrada e o config)'),
    ('Ambiente', f"Python {resu['python']}; {resu['sistema']}"),
    ('Semente', str(resu['semente'])),
    ('Comando 1 (pasta de trabalho)', 'python3 gerar_base_cpfl_v2.py'),
    ('Execução 1', f"{resu['executado_em']} — {resu['duracao_s']} s"),
    ('Comando 2 (pasta limpa, só .py + .json copiados)', 'python3 gerar_base_cpfl_v2.py > execucao_limpa.log'),
    ('Execução 2', f"{limpa['executado_em']} — {limpa['duracao_s']} s"),
    ('Comando 3', 'python3 verificar_base_cpfl_v2.py'),
    ('Execução 3', f"{ver['executado_em']} — {ver['duracao_s']} s"),
    ('Linhas por distribuidora (referência)', json.dumps(resu['linhas_por_distribuidora_referencia'], ensure_ascii=False)),
    ('Fator de escala por distribuidora', json.dumps(resu['fator_escala_por_distribuidora'], ensure_ascii=False)),
]
r = 5
for a, b in ex: put(ws, r, [a, b]); r += 1
r += 1; head(ws, r, ['Arquivo (caminho relativo)', 'SHA-256 | bytes | igual na pasta limpa?']); r += 1
for k in resu['sha256']:
    put(ws, r, [k, f"{resu['sha256'][k]} | {resu['bytes'][k]:,} | {'sim' if resu['sha256'][k] == limpa['sha256'][k] else 'NÃO'}".replace(',', '.')]); r += 1
r += 1
ws.cell(row=r, column=1, value='Os três arquivos da pasta RESERVADO_NAO_ABRIR_ATE_CONGELAMENTO não foram lidos; os hashes acima congelam seu conteúdo e podem ser regenerados pelo gerador.').font = Font(name=F, italic=True, size=9)

# 6 Previa
ws = sheet('Previa_1000', 'Prévia: primeiras 1.000 linhas da base registrada (de 501.000)', [17, 17, 18, 22, 13, 14, 12, 12, 10, 12, 16, 11, 12])
with gzip.open(f"{B}/{cfg['saidas']['base_registrada']}", 'rt', encoding='utf-8') as f:
    rd = csv.reader(f); cab = next(rd); head(ws, 4, cab)
    for i, row in enumerate(rd):
        if i >= 1000: break
        vals = row[:5] + [float(row[5]), float(row[6]), float(row[7])] + row[8:10] + [float(row[10])] + row[11:]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=5 + i, column=j, value=v); c.font = N
            if j in (6, 11): c.number_format = RS
            if j in (7, 8): c.number_format = '0.0000'
ws.freeze_panes = 'A5'
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT); print('OK', OUT)
