# -*- coding: utf-8 -*-
"""Monta o dossiê consolidado (Markdown -> DOCX -> PDF) e os PTs individuais.
Uso: python build_dossie.py <versao> <pasta_saida>
Requer: pandoc, LibreOffice (soffice), python-docx."""
import sys, os, re, subprocess, importlib.util
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

VER = sys.argv[1] if len(sys.argv) > 1 else "0.1"
OUT = sys.argv[2] if len(sys.argv) > 2 else "out"
HERE = os.path.dirname(os.path.abspath(__file__))
os.makedirs(OUT, exist_ok=True)

def load(mod):
    spec = importlib.util.spec_from_file_location(mod, os.path.join(HERE, mod + ".py"))
    m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m); return m

blocos = [load("pts_b03"), load("pts_b04"), load("pts_b05"), load("pts_b07"), load("pts_b08"), load("pts_b09"), load("pts_b10"), load("pts_b11"), load("pts_b12"), load("pts_b13")]
PROG_MD = blocos[1].PROG_MD
PTS = [dict(p, data=p.get("data", b.DATA)) for b in blocos for p in b.PTS]
for _b in blocos:
    for p in PTS:
        u = getattr(_b, "UPDATES", {}).get(p["codigo"])
        if not u: continue
        for k in ("situacao", "titulo", "resultados", "conclusao"):
            if k in u: p[k] = u[k]
        p["procedimentos"] = [(a, bb, u.get("proc_fix", {}).get(a, c)) for a, bb, c in p["procedimentos"]]
        p["corpo"] = p["corpo"].rstrip() + "\n" + u.get("corpo_add", "")
        p["resultados"] += u.get("resultados_add", ""); p["evidencias"] += u.get("evidencias_add", "")
        p["data"] = _b.DATA
PTS.sort(key=lambda p: 0 if p["codigo"].endswith("EC") else int(p["codigo"].split()[-1]))
PREP, DATA = blocos[0].PREP, blocos[-1].DATA
B11, B13 = blocos[-3], blocos[-1]
FINAL = not VER.startswith("0.")

def pt_md(p, level=2):
    data_pt = p.get("data", DATA)
    h = "#" * level
    lines = [f"{h} {p['codigo']} — {p['titulo']}", "",
             f"*Versão {VER} · Estado: {'concluído pelo preparador; aguarda revisão humana' if FINAL else p.get('situacao','em elaboração')} · Perímetro: consolidado CPFL Energia S.A., exercício 2025 · Preparado em {data_pt}*", ""]
    lines += ["| Campo | Conteúdo |", "|------|----------------------|",
              f"| 1. Código, título, norma e itens | **{p['codigo']}** — {p['titulo']}. {p['norma']} |",
              f"| 2. Objetivo | {p['objetivo']} |",
              f"| 2. Conta / processo / divulgação | {p['conta']} |",
              f"| 3. Afirmações e riscos | {p['afirmacoes']} |",
              f"| 4. População, fontes e seleção | {p['populacao']} |", ""]
    lines += [f"**5. Procedimentos planejados ou executados**", "", "| Ref. | Procedimento | Modalidade |", "|---|--------------|--------|"]
    lines += [f"| {a} | {b} | {c} |" for a, b, c in p["procedimentos"]]
    lines += ["", f"**Desenvolvimento**", "", p["corpo"].strip(), ""]
    lines += ["| Campo | Conteúdo |", "|------|----------------------|",
              f"| 6. Evidências e referências cruzadas | {p['evidencias']} |",
              f"| 7. Resultados e exceções | {p['resultados']} |",
              f"| 7. Conclusão | {p['conclusao']} |",
              f"| 7. Limitações | {p['limitacoes']} |",
              f"| 7. Requisito × orientação × julgamento | **Requisito:** {B13.RJ[p['codigo']][0]} **Orientação (não obrigatória):** {B13.RJ[p['codigo']][1]} **Julgamento do preparador:** {B13.RJ[p['codigo']][2]} |",
              f"| 8. Preparador e data | {PREP} — {data_pt} |",
              f"| 9. Revisão | Revisão humana: ____  Data: ____ (em branco). Revisão automatizada: RV-B01 a RV-B12 (PT 220; CPFL_B12_Revisao_Tecnica_Cobertura.xlsx). Nenhuma revisão de terceiros simulada |", ""]
    return "\n".join(lines)

STATUS = [("PT EC","Estrutura Conceitual","B03"),("PT 200","Objetivos gerais","B03"),("PT 210","Termos do trabalho","B03"),("PT 220","Gestão da qualidade","B03"),("PT 230","Documentação","B03"),
 ("PT 240","Fraude","B04"),("PT 250","Leis e regulamentos","B04"),("PT 260","Comunicação com a governança","B11"),("PT 265","Deficiências de controle","B11"),("PT 300","Planejamento","B04"),
 ("PT 315","Identificação e avaliação dos riscos","B04"),("PT 320","Materialidade","B04"),("PT 330","Respostas aos riscos","B04"),("PT 402","Organização prestadora de serviços","B05"),
 ("PT 450","Distorções","B08/B11"),("PT 500","Evidência","B07/B11"),("PT 501","Itens específicos","B10"),("PT 505","Confirmações externas","B10"),("PT 510","Saldos iniciais","B09"),
 ("PT 520","Procedimentos analíticos","B09"),("PT 530","Amostragem","B08"),("PT 540","Estimativas contábeis","B09"),("PT 550","Partes relacionadas","B09"),("PT 560","Eventos subsequentes","B10"),
 ("PT 570","Continuidade operacional","B10"),("PT 580","Representações formais","B11"),("PT 600","Auditoria de grupos","B05")]
done = {p["codigo"]: ("Concluído pelo preparador (aguarda revisão humana)" if FINAL else p.get("situacao", "Redigido")) for p in PTS}

memo = open(os.path.join(HERE, "..", "producao", "01_Dados_reais", "CPFL_B02_Memorando_Entidade.md"), encoding="utf-8").read()
memo = re.sub(r"^---.*?---\s*", "", memo, flags=re.S)
memo = re.sub(r"(?m)^# ", "### ", memo)
B06 = open(os.path.join(HERE, "..", "producao", "03_Base_sintetica", "CPFL_B06_Projeto_Base_Sintetica.md"), encoding="utf-8").read()
B06 = re.sub(r"^---.*?---\s*", "", B06, flags=re.S)
B06 = re.sub(r"(?m)^## ", "#### ", B06); B06 = re.sub(r"(?m)^# ", "### ", B06)
B06 = "*Documento do Bloco 06 reproduzido sem alteração de conteúdo; a decisão nele pedida foi aprovada pelo aluno em 23/10/2026 (D-B06-01 e D-B06-02).*\n\n" + B06
MIN11 = open(os.path.join(HERE, "..", "producao", "02_Planejamento", "CPFL_B11_Minutas_Governanca_Representacoes_NAO_ENVIADAS.md"), encoding="utf-8").read()
MIN11 = re.sub(r"^---.*?---\s*", "", MIN11, flags=re.S)
MIN11 = re.sub(r"(?m)^# ", "## ", MIN11)

SUF = " (versão final)" if FINAL else " (em elaboração)"
md = f"""---
title: "Dossiê de Auditoria — CPFL Energia S.A."
subtitle: "Demonstrações contábeis consolidadas de 2025 · Trabalho acadêmico"
author: "Matrícula 231020069 · UnB · Auditoria 1 (CCA0024) · Turma 02 · 2026.2 · Prof. Lucas Oliveira Gomes Ferreira"
date: "Versão {VER} — {DATA}{SUF}"
lang: pt-BR
---

> **AVISO.** Trabalho acadêmico individual. Não é opinião profissional nem relatório de auditoria e não fornece asseguração (ver PT EC). Usa somente documentos públicos (CVM, B3, RI da companhia) e dados **sintéticos** identificados como tais. Não houve contato com a CPFL, seus auditores, clientes ou fornecedores. Procedimentos são classificados como **Realizado com evidência pública**, **Planejado por falta de acesso** ou **Simulado para fins didáticos**.

# 1. Introdução e escopo

- **Companhia atribuída:** CPFL Energia S.A. — CNPJ 02.429.144/0001-93 — CPFE3 (Novo Mercado). Nome anterior: "Draft II" (FR 2025 v6, item 1.1).
- **Objeto:** demonstrações contábeis individuais e consolidadas de 31/12/2025, comparativos de 2024, notas explicativas e relatório do auditor independente (PwC, 04/03/2026), publicadas em DF25 CPFL.pdf (115 p.).
- **Perímetro principal:** consolidado societário. A controladora individual é usada para a parte de *holding* e o PT 600.
- **Normas:** NBC TA Estrutura Conceitual e NBC TA 200 a 600, nas versões aplicáveis ao exercício de 2025 (seção 2.3).
- **Base sintética obrigatória:** ativo financeiro da concessão (R$ 29.623.619 mil; 36,5% do ativo total — DF25 p. 34/2), com base sintética de 501.000 registros (500.000 bens + 1.000 duplicidades previstas no desenho), conciliada no total e por distribuidora com os saldos publicados (PT 500; Bloco 07).
- **Entrega:** repositório GitHub até 25/10/2026, com este dossiê e os anexos editáveis (planilhas, base, códigos, README).

# 2. Metodologia, convenções e limitações

## 2.1 Modalidades e marcas

Cada procedimento recebe uma das três modalidades do enunciado. Um PT pode combinar modalidades, desde que cada procedimento tenha a sua.

| Sigla | Modalidade | Quando usar |
|---|---|---|
| [Pub] | Realizado com evidência pública | Executado com documentos efetivamente obtidos (CVM, B3, RI, DF) |
| [P] | Planejado por falta de acesso | Detalhado, mas não executado: depende de registros, pessoas, terceiros ou sistemas |
| [S] | Simulado para fins didáticos | Executado sobre dados sintéticos ou minuta não enviada, identificados como tais |

| Marca | Significado |
|---|---|
| ✓ | Conferido com o documento-fonte |
| Σ | Soma ou recálculo conferido |
| ⇄ | Cruzado com fonte independente (por exemplo, PDF × CVM) |
| ↔ | Referência cruzada a outro PT ou planilha |
| ⚠ | Exceção ou diferença explicada no próprio PT |
| ◐ | Limitação: evidência indisponível ou parcial |
| ◇ | Hipótese ou julgamento do preparador (não é fato documentado) |

## 2.2 Citação de fontes

Todas as citações informam o documento e a página. "DF25 p. 34/2" = DF25 CPFL.pdf, página física 34 / página impressa 2. "NORMAS p. n" = CONSOLIDAÇÃO DAS NORMAS AUD 1.pdf, página n do PDF. A relação completa das fontes, com URL e data de acesso, está no Apêndice B.

## 2.3 Versões normativas aplicadas ao exercício de 2025

| Norma | Versão aplicada | Observação |
|---|---|---|
| Estrutura Conceitual | DOU 25/11/2015 | Efeitos a partir de 01/01/2016 (NORMAS p. 24) |
| NBC TA 200, 210, 230, 300, 320, 330, 450, 500, 510, 560, 580 | (R1), DOU 05/09/2016 | — |
| NBC TA 220 | (R3), DOU 09/12/2021 | Períodos iniciados a partir de 01/01/2023 (NORMAS p. 93) |
| NBC TA 240 | **(R1)** | A (R2), DOU 25/09/2026, só se aplica a exercícios iniciados a partir de 01/01/2027 |
| NBC TA 250 | DOU 14/02/2019 | — |
| NBC TA 260 | (R2), DOU 04/07/2016 | — |
| NBC TA 265, 402, 501, 505, 520, 530, 550 | Versões de 2009 (Res. CFC 1.210 a 1.224/09) | — |
| NBC TA 315 | (R2), DOU 02/09/2021 | Períodos iniciados a partir de 01/01/2022 (NORMAS p. 256) |
| NBC TA 540 | (R2), DOU 23/10/2019 | Períodos iniciados a partir de 01/01/2020 (NORMAS p. 552) |
| NBC TA 570 | **Versão de 2016** | A (R1), DOU 25/09/2026, só se aplica a partir de 2027 (NORMAS p. 646) |
| NBC TA 600 | (R2), DOU 06/12/2023 | Períodos iniciados a partir de 01/01/2024 (NORMAS p. 688) |

Fonte: lista oficial do CFC (acesso em 22/10/2026) e textos do consolidado. A vigência da NBC TA 240 (R2) e da NBC TA 570 (R1) (DOU 25/09/2026; auditorias de períodos iniciados em ou após 01/01/2027) foi reconfirmada no Bloco 12 em notícias do CRCSP e do Portal Contnews (FC31); o DOU não foi aberto diretamente. Comentários didáticos do consolidado ("Comentado […]") não foram tratados como requisitos.

## 2.4 Limitações gerais

1. Não há acesso a registros, sistemas, pessoas ou terceiros. Procedimentos que dependem disso são apenas planejados.
2. As demonstrações primárias da DF25 são imagens e foram transcritas visualmente. Os totais foram recalculados por fórmula: 4 erros de uma transcrição anterior foram corrigidos (Apêndice A).
3. Os Formulários de Referência 2025 e 2026 disponíveis estão truncados (DV15).
4. A conferência com a base de dados da CVM foi feita no Bloco 13 para balanço e DFC consolidados (seção 3.2: 182 valores, 0 diferenças sem explicação). A DRE da CVM não pôde ser lida; o LAIR confere pela DFC.
5. Os resultados sobre a base sintética não são constatações sobre a CPFL. A mesma IA desenhou os erros e os testes (PT 530): não há teste cego independente.

## 2.5 Modelo de papel de trabalho

Todos os 27 PTs seguem o mesmo modelo, que atende aos 9 campos mínimos do enunciado e à NBC TA 230, itens 8 e 9 (NORMAS p. 133):

| Campo | Conteúdo |
|---|---|
| Cabeçalho | Versão, estado, perímetro e data de preparação |
| 1 | Código, título, norma (com revisão), itens e páginas no consolidado |
| 2 | Objetivo; conta, processo ou divulgação examinada |
| 3 | Afirmações e riscos (ligação à matriz R01–R19 do PT 315) |
| 4 | População, fontes e critério de seleção |
| 5 | Procedimentos, cada um com sua modalidade |
| 6 | Evidências e referências cruzadas |
| 7 | Resultados e exceções (fato documentado, hipótese e resultado sintético separados), conclusão, limitações e **requisito × orientação × julgamento** |
| 8 | Preparador e data (preparação assistida por IA) |
| 9 | Revisão humana em branco; revisões automatizadas identificadas |

## 2.6 Referências cruzadas

| Formato | Significado | Exemplo |
|---|---|---|
| DF25 p. 34/2 | Página física / página impressa da DF 2025 | DF25 p. 113/81 |
| NORMAS p. n | Página do consolidado de normas do curso | NBC TA 530 item 15 (NORMAS p. 536) |
| → PT 450 | Remissão a outro papel | → PT 540 seção 4 |
| D-B04-01, PB09-01, DV21, FC30, RV12-03 | Decisão, pendência, divergência, fonte e registro de revisão | — |
| R01–R19 | Riscos da matriz do PT 315 | R01 ativo financeiro da concessão |
| CVM 1.02.01.10.08 | Conta da DFP nos dados abertos da CVM | Ativo financeiro da concessão |

# 3. Entendimento da entidade e dados reais

""" + B13.PARTE_B_ID + """

## 3.3 Memorando factual da entidade (Bloco 02)

{memo}

# 4. Situação dos papéis de trabalho nesta versão

| PT | Norma / tema | Bloco | Situação na v{VER} |
|---|---|---|---|
""" + "\n".join(f"| {c} | {t} | {b} | {done.get(c, 'Pendente')} |" for c, t, b in STATUS) + """

Programas por ciclo (PG-REC, PG-DISP, PG-EST, PG-IMOB, PG-PASS, PG-ESTIM) e programa de consolidação (PG-CONS, Bloco 13): seção 6. Matriz final de cobertura das 33 obrigações do enunciado: CPFL_B12_Revisao_Tecnica_Cobertura.xlsx (aba Matriz_Cobertura_Final).

# 5. Papéis de trabalho

""" + "\n\n".join(pt_md(p) for p in PTS) + """

# 6. Programas de auditoria por ciclo

""" + PROG_MD + "\n\n" + B13.PG_CONS + """

""" + B11.CONCLUSAO_MD + "\n\n" + B13.LIGACAO_RISCOS + """

# 8. Base sintética: projeto (Bloco 06) e geração (Bloco 07)

""" + B13.PARTE_E_INTRO + "\n" + B06 + "\n\n" + B13.PARTE_E_GER + """

# Apêndice A — Divergências registradas

| ID | Divergência | Tratamento |
|---|---|---|
| DV01 | Controle antigo: 129.669 em "Investimentos em instrumentos patrimoniais" 2025; DF25 p. 34/2 mostra 129.659 | Corrigido; o ativo não circulante fecha sem diferença |
| DV02 | Guia indicava ativo total na "p. 23" | Página correta: p. física 34 / impressa 2 |
| DV03 | Base xlsx no Drive (41,3 MB) com tamanho diferente do hash registrado (35,8 MB) | Base reproduzida pelo gerador (D-B01-03); depois substituída pela base v2 (D-B06-01) |
| DV04 | Gerador com caminhos absolutos (/mnt/data) | Resolvido: gerador v2 com caminhos relativos, testado em pasta limpa (Bloco 07) |
| DV05 | Nomes de arquivos diferentes entre o registro e o Drive | Registrados os dois nomes |
| DV06 | PDF do enunciado não localizado | Usado o texto fornecido pelo aluno |
| DV07 | Pasta de produção inexistente | Criada em 22/10/2026 |
| DV08 | Cronograma do guia com datas incoerentes | Novo cronograma |
| DV09 | NBC TA 240 (R2) e 570 (R1) publicadas em 2026 | Não aplicadas a 2025 |
| DV10 | Ativo não circulante 2024: soma com diferença de −2 | Arredondamento (nota 2.4) |
| DV11 | IR/CS a compensar (circulante) 2024: 716.764 × 715.764 | Corrigido (DF25 p. 34/2) |
| DV12 | Despesas com vendas 2025: 986.664 × 985.664 | Corrigido (DF25 p. 36/4) |
| DV13 | Resultado do serviço 2025: 10.776.490 × 10.778.490 | Corrigido (DF25 p. 36/4) |
| DV14 | Nome anterior "Draftll Participações" × "Draft II" | Usado o FR 2025 v6 item 1.1; confirmar na CVM |
| DV15 | FR 2025 e FR 2026 truncados (16 MiB) | Limitação registrada; buscar versão íntegra |
| DV17 | Matriz preliminar (etapa 3), R08: 999.257 descrito como ativos, mas soma ativos (595.863) e passivos (403.394) de derivativos | Corrigido no PT 315 |
| DV18 | Nota 30 (p. 103): aquisições do segmento de geração lidas como 259.975 na primeira transcrição; a imagem ampliada mostra 269.975 | Corrigido; a soma dos segmentos passou a conciliar com o total (6.112.098) |
| DV19 | Base da etapa 1 distribuía bens por distribuidora em proporções hipotéticas (40/30/25/5%); notas 11 das DFs individuais mostram 42,6/43,0/13,0/1,5% | Base v2 conciliada por distribuidora (D-B06-01, aprovada pelo aluno em 23/10/2026) |
| DV20 | DF24 CPFL.pdf do Drive não contém o relatório do auditor independente (só o Relatório da Administração, as DF, o parecer do Conselho Fiscal e as declarações) | Opinião de 2024 inferida dos documentos das p. 138–139 (PT 510); obter o relatório integral na CVM/ENET (B12) |
| DV21 | Quadro de liquidez do PT 570 (planilha B10): total de debêntures registrado como 30.641.582; a nota 34 g (DF25 p. 113/81) mostra 30.641.585, igual à soma da linha | Corrigido no Bloco 12 (RV-B12); a falsa diferença de arredondamento de −3 deixou de existir |
| DV16 | "EXEMPLO 1/2/3 DE DOSSIE" são guias IFAC (vol. 2, 2ª ed. CFC 2010 e 4ª ed. em inglês) e a Orientação Prática da CGU (2022) | Usados só como referência de forma |

# Apêndice B — Fontes utilizadas

| ID | Documento | Páginas usadas | Local / URL | Acesso |
|---|---|---|---|---|
| FC01 | DF25 CPFL.pdf — DF individuais e consolidadas 2025 e relatório do auditor (SHA-256 0694f2b9…a578) | 26–36, 39, 41–42, 45–46, 54, 60–64, 68–85, 94–108, 113–115 | Drive (pasta 01). Publicação legal: https://publicidadelegal.valor.com.br/valor/2026/03/20/CPFLENERGIA1588267020032026.pdf | 22/10/2026 |
| FC09 | CONSOLIDAÇÃO DAS NORMAS AUD 1.pdf (756 p.; SHA-256 f2edf5fa…2775) | 3–145, 252–256, 548–552, 639–646, 684–702 | Drive (pasta 07) | 22/10/2026 |
| FC11 | CFC — NBC TA de Auditoria Independente | — | https://cfc.org.br/tecnica/normas-brasileiras-de-contabilidade/nbc-ta-de-auditoria-independente/ | 22/10/2026 |
| FC12 | Notícia sobre a vigência das NBC TA 240 (R2) e 570 (R1) (fonte secundária) | — | https://www.lrmconsultoria.com.br/noticias/tecnicas/2026/09/30/cfc-atualiza-normas-de-auditoria-sobre-fraude-e-continuidade-operacional.html | 22/10/2026 |
| FC13 | CVM Dados Abertos — DFP (diretório) | — | https://dados.cvm.gov.br/dados/CIA_ABERTA/DOC/DFP/DADOS/ | 22/10/2026 |
| FC14 | FR_CPFL Energia 2025_V6_06-04-26_PT.pdf (truncado) | Item 1.1, p. 1 | Drive (pasta 03); portal https://www.rad.cvm.gov.br/ENETWEB/frmConsultaExternaCVM.aspx | 22/10/2026 |
| FC15 | ITR 3TRI 2025.pdf — DF intermediárias e relatório de revisão (48 p.) | 14–15 | Drive (pasta 01) | 22/10/2026 |
| FC16 | Fato Relevante — Renovação das concessões (08/05/2026) | 1 | Drive (pasta 03) | 22/10/2026 |
| FC18 | Regimentos e políticas de governança (Comitê de Auditoria 11/11/2025 e 08/09/2026; Conselho Fiscal 10/05/2024; Auditoria Interna 24/01/2024; Política de Gerenciamento de Riscos; Política Anticorrupção) | Títulos e datas | Drive (pasta 03) | 22/10/2026 |
| FC19 | Planilha e relatório de materialidade (etapa 2) e matriz preliminar de riscos (etapa 3) | Integral | Drive (MANDAR PARA O CLAUDE PRO) | 22/10/2026 |
| FC20 | DF 2025 das controladas CPFL Paulista, RGE, Piratininga, Transmissão e Renováveis (relatórios da PwC) | Relatório do auditor e totais; notas 11 (Paulista p. 29, RGE p. 31, Piratininga p. 29, páginas impressas) | Drive (pasta 06) | 22/10/2026 |
| FC17 | Modelos de forma: guias IFAC vol. 2 (CFC 2ª ed.; 4ª ed. inglês) e CGU Orientação Prática (2022) | — | Drive (pasta MANDAR PARA O CLAUDE PRO) | 22/10/2026 |
| FC21 | Base SINTÉTICA v2 e códigos: gerar_base_cpfl_v2.py, config_base_cpfl_v2.json, verificar_base_cpfl_v2.py, CPFL_SINTETICA_Base_AFC_2025_registrada.csv.gz (SHA-256 877d4a02…4bb0), CPFL_B06_Projeto_Base_Sintetica, CPFL_B07_Conciliacao_Verificacoes_Base.xlsx | Integral | Drive (pasta 09/03_Base_sintetica) e repositório | 23/10/2026 |
| FC22 | Testes do Bloco 08: B08_Desenho_pre_registro.md (Drive 04_Testes, 23/10/2026 19:54), B08_congelamento.json (19:55:53), testes_b08.py, avaliacao_complementar_b08.py, comparar_gabarito_b08.py, saídas CSV/JSON, CPFL_B08_Amostragem_Testes_Distorcoes.xlsx | Integral | Drive (pasta 09/04_Testes) e repositório | 23/10/2026 |
| FC23 | DF24 CPFL.pdf (DF 2024: BP, DRE, nota 1.3, nota 2.1, parecer do Conselho Fiscal e declaração da diretoria) | p. 27, 34–36, 45–46, 138–139 | Drive (pasta 01) | 23/10/2026 |
| FC24 | ITR 2TRI 2026 (notas 1.2 e 11: assinatura dos aditivos em 06/05/2026 e reclassificação de 29.599.494 do ativo financeiro para o intangível) | p. 43 e 49 | Drive (pasta 01) | 23/10/2026 |
| FC25 | IBGE — IPCA acumulado 2025 (4,26%) e 2024 (4,83%), Agência de Notícias IBGE | — | agenciadenoticias.ibge.gov.br (notícias 45613 e 42416) | 23/10/2026 |
| FC26 | Política de Partes Relacionadas (CA de 16/12/2021) e Fato Relevante de 08/05/2026 (renovação das concessões) | Integral | Drive (pasta 03) | 23/10/2026 |
| FC27 | Release de Resultados 4T25 (alavancagem 2,30x no critério dos covenants; cronograma de amortização; dividendos propostos) | p. 3–4, 15 | Drive (pasta 01) | 23/10/2026 |
| FC28 | ITR 1TRI 2026 (autorização 11/05/2026; revisão PwC 13/05/2026; nota 23 de eventos subsequentes) | p. 42, 55, 58 | Drive (pasta 01) | 23/10/2026 |
| FC29 | RI CPFL — lista de comunicados ao mercado de 2026 (último: 14/10/2026, calendário do 3T26) e notícias de dividendos (02/10 e 23/10/2026) | — | ri.cpfl.com.br; suno.com.br; meusdividendos.com | 23/10/2026 (corte da pesquisa) |
| FC30 | Bloco 11: CPFL_B11_Comunicacoes_Distorcoes_Evidencia.xlsx (PT 265, PT 450, PT 500, PT 260) e minutas de governança, deficiências e representações NÃO ENVIADAS | Integral | Drive (pasta 09/02_Planejamento) e repositório | 23/10/2026 |
| FC31 | CRCSP, "NBC TA 240 (R2)" (notícia de 02/10/2026) e Portal Contnews, "CFC atualiza normas de auditoria sobre fraudes e continuidade operacional" | — | online.crcsp.org.br/portal/noticias/noticia.asp?c=11367; portalcontnews.com.br/?p=343355 | 23/10/2026 |
| FC32 | RI CPFL — lista de comunicados de 2026 relida no Bloco 12; Cenário Energia (09/03/2026) sobre a recomposição do Conselho | — | ri.cpfl.com.br/list.aspx?ano=2026; cenarioenergia.com.br | 23/10/2026 (corte final da pesquisa) |
| FC33 | Bloco 12: CPFL_B12_Revisao_Tecnica_Cobertura.xlsx (registro de revisão, matriz final de cobertura, conferências automatizadas, revisão analítica final) e scripts b12_*.py | Integral | Drive (pasta 09/00_Controle) e repositório | 23/10/2026 |
| FC34 | CVM — dados abertos, DFP 2025 (dfp_cia_aberta_2025.zip): índice de documentos, BPA, BPP e DFC-MI consolidados e pareceres/declarações; versão 1, documento 155108, recebido em 05/03/2026. Conferência: CPFL_B13_Conferencia_PDF_CVM.xlsx | Linhas da CPFL (CNPJ 02.429.144/0001-93) | dados.cvm.gov.br/dataset/cia_aberta-doc-dfp (arquivo obtido pelo aluno e salvo no Drive em 23/10/2026) | 24/10/2026 |

# Apêndice C — Anexos digitais que acompanham o dossiê (repositório)

""" + B13.PARTE_G_ANEXOS + """

# Apêndice D — Minutas didáticas NÃO ENVIADAS (PT 260, PT 265 e PT 580)

""" + MIN11 + """
"""

base = f"Dossie_CPFL_2025_231020069_v{VER}"
mdp = os.path.join(OUT, base + ".md"); open(mdp, "w", encoding="utf-8").write(md)
docx = os.path.join(OUT, base + ".docx")
subprocess.run(["pandoc", mdp, "-o", docx, "--toc", "--toc-depth=2", "-M", "toc-title=Sumário"], check=True)

# pós-processamento: fonte Arial, tabelas com bordas e fonte menor
d = Document(docx)
for st in d.styles:
    try:
        if st.type == 1 and st.font is not None:
            st.font.name = "Arial"
            rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
            if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), "Arial")
    except Exception: pass
for t in d.tables:
    tblPr = t._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x = OxmlElement(f"w:{e}"); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "4"); x.set(qn("w:color"), "A6A6A6"); b.append(x)
    tblPr.append(b)
    for row in t.rows:
        for c in row.cells:
            for par in c.paragraphs:
                for r in par.runs: r.font.size = Pt(8.5)
from docx.shared import Cm
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
def _field(par, instr):
    r = par.add_run(); r.font.size = Pt(8); f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin"); r._r.append(f1)
    r2 = par.add_run(); r2.font.size = Pt(8); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r2._r.append(it)
    r3 = par.add_run(); r3.font.size = Pt(8); f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end"); r3._r.append(f2)
for sec in d.sections:
    sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(2)
    fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    fp.text = ""; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Exercício acadêmico sem asseguração · CPFL Energia S.A. 2025 · Matrícula 231020069 · v" + VER + " · Página "); r.font.size = Pt(8)
    _field(fp, "PAGE"); fp.add_run(" de ").font.size = Pt(8); _field(fp, "NUMPAGES")
for par in d.paragraphs:
    if par.style.name in ("Heading 1",) and not par.text.startswith("Sumário"):
        par.paragraph_format.page_break_before = True
for par in d.paragraphs:
    if par.style.name == "Date":
        par.add_run().add_break(WD_BREAK.PAGE); break
d.save(docx)
subprocess.run([sys.executable, os.path.join(HERE, "lo_update_export.py"), docx], check=True)

# PTs individuais (Markdown) para a pasta 05_Papeis
ptdir = os.path.join(OUT, "PTs"); os.makedirs(ptdir, exist_ok=True)
for p in PTS:
    fn = re.sub(r"\W+", "_", p["codigo"]) + "_CPFL_2025.md"
    open(os.path.join(ptdir, fn), "w", encoding="utf-8").write(pt_md(p, 1))
print("OK", docx)
