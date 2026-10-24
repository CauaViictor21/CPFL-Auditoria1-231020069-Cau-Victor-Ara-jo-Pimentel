# -*- coding: utf-8 -*-
"""Bloco 05 - PT 600 (grupo) e PT 402 (prestadores). Gera CPFL_B05_Grupo_Prestadores.xlsx.
Fontes: DF25 nota 13.1 (p. 70/38: ativo total e resultado das controladas diretas, individual de cada uma),
nota 30 (p. 103/71: segmentos), nota 1 (p. 41/9), DFs das controladas (Drive pasta 06).
Uso: python b05_grupo.py <pasta_saida>"""
import sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else "."
F = "Arial"; HEAD = PatternFill("solid", fgColor="1F3864"); INP = PatternFill("solid", fgColor="FFF2CC"); CHK = PatternFill("solid", fgColor="E2EFDA")
th = Side(style="thin", color="BFBFBF"); BRD = Border(left=th, right=th, top=th, bottom=th)
wb = Workbook(); wb.remove(wb.active)

def sheet(name, title, cols, widths, note):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = Font(name=F, bold=True, size=12, color="1F3864")
    ws["A2"] = "Preparador: matrícula 231020069 (assistido por IA) | Data: 22/10/2026 | Revisão: ____ (sem revisão de terceiros simulada)"
    ws["A2"].font = Font(name=F, italic=True, size=8, color="595959")
    ws["A3"] = note; ws["A3"].font = Font(name=F, size=9, color="7F6000"); ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(cols)); ws.row_dimensions[3].height = 52
    for j, h in enumerate(cols, 1):
        c = ws.cell(row=4, column=j, value=h); c.font = Font(name=F, bold=True, color="FFFFFF", size=9); c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BRD
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A5"; return ws

def put(ws, r, vals, fills=None, fmt=None, bold=False):
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=j, value=v); c.font = Font(name=F, size=9, bold=bold); c.border = BRD
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if fills and j in fills: c.fill = fills[j]
        if fmt and j in fmt: c.number_format = fmt[j]

# ---------------------------------------------------------------- Parâmetros
ws = sheet("Parametros", "Parâmetros do grupo (R$ mil)", ["Item", "Valor", "Fonte"], [48, 18, 60],
           "Materialidades do grupo vêm do PT 320 (validadas pelo aluno em 22/10/2026).")
par = [("Ativo total consolidado", 81099930, "DF25 p. 34/2"), ("Lucro antes dos tributos consolidado", 8091246, "DF25 p. 36/4"),
       ("Materialidade global do grupo", 404562.30, "PT 320"), ("Materialidade para execução do grupo", 303421.725, "PT 320"),
       ("Claramente trivial do grupo", 20228.115, "PT 320"), ("% execução componente — componentes de maior porte", 0.70, "Julgamento (600 item 35(a): deve ser inferior à execução do grupo)"),
       ("% execução componente — demais componentes", 0.50, "Julgamento (maior risco de agregação em componentes menores e numerosos)")]
for i, (a, b, c) in enumerate(par, 5):
    put(ws, i, [a, b, c], fills={2: INP}, fmt={2: "0%" if b < 1 else "#,##0.00"})

# ---------------------------------------------------------------- Componentes
ws = sheet("Componentes_600", "PT 600 — Componentes, porte relativo e abordagem (controladas diretas — nota 13.1)",
           ["Componente", "Segmento", "Ativo total (individual)", "Resultado 2025", "Ativo / ativo consolidado", "|Resultado| / LAIR consolidado", "Classificação (julgamento)", "Abordagem planejada", "Mat. execução do componente", "Limite de comunicação", "Auditor do componente (fonte)"],
           [26, 16, 15, 14, 12, 12, 18, 40, 14, 13, 34],
           "Ativo e resultado: DF25 nota 13.1, p. 70/38 (informação individual de cada controlada; inclui investimentos em outras controladas — há sobreposição, por isso as proporções são indicativas e não somam 100%). Classificação = julgamento do auditor do grupo (NBC TA 600 R2 não usa mais a categoria 'componente significativo' da versão anterior; usa-se o risco e a agregação — itens 22, 33, 35).")
comps = [
 ("CPFL Paulista","Distribuição",22954592,1864350,"Maior porte","Procedimentos de auditoria nas informações financeiras completas (ativo financeiro da concessão, receita não faturada)","PwC — relatório 04/03/2026 (DF Paulista)"),
 ("CPFL RGE","Distribuição",21762223,1327456,"Maior porte","Idem","PwC — opinião sem modificação; data não localizada no texto extraído (DF RGE)"),
 ("CPFL Piratininga","Distribuição",8226299,548281,"Porte médio","Procedimentos sobre saldos de risco (ativo financeiro da concessão, receita, dívida)","PwC — 04/03/2026 (DF Piratininga)"),
 ("CPFL Santa Cruz","Distribuição",2647005,130563,"Porte menor","Procedimentos sobre ativo financeiro da concessão (base sintética inclui) + analíticos","Não verificado (DF não disponível no Drive)"),
 ("CPFL Transmissão","Transmissão",9425400,268600,"Porte médio","Procedimentos sobre ativo contratual (PAA 1) e receita de remuneração","PwC — 04/03/2026 (DF Transmissão)"),
 ("CPFL Renováveis","Geração",8335380,1062034,"Porte médio","Procedimentos sobre imobilizado/intangível e receita de geração","PwC — 04/03/2026 (DF Renováveis)"),
 ("CPFL Geração","Geração",3461790,1129886,"Porte menor (holding de geração)","Analíticos + equivalência de controladas em conjunto; incorporada em 02/02/2026 (nota 37.1)","Não verificado"),
 ("CPFL Brasil","Comercialização",4855998,142279,"Porte menor","Corte de receita de comercialização (CCEE) + analíticos","Não verificado"),
 ("CPFL Serviços e entidades de serviços compartilhados (Atende, Infra, Pessoas, Finanças, Supri, AUTHI, Telecom, Planalto, Jaguari Geração, Alesta, Clion)","Serviços/outros",741948+22742+22148+14934+24989+14961+10881+4433+4243+78804+213081+10155,46345+1317+12905+16032+29174+13691+857+1470+4181+2280+77788+367,"Agregado pouco relevante individualmente","Analíticos no nível do grupo; atenção a funções compartilhadas (finanças, pessoas, suprimentos) que afetam vários componentes","Não verificado"),
]
r = 5
for nm, sg, at, res, cl, ab, au in comps:
    exec_pct = "Parametros!$B$10" if cl == "Maior porte" else "Parametros!$B$11"
    put(ws, r, [nm, sg, at, res, f"=C{r}/Parametros!$B$5", f"=ABS(D{r})/Parametros!$B$6", cl, ab, f"=Parametros!$B$8*{exec_pct}", "=Parametros!$B$9", au],
        fills={3: INP, 4: INP, 5: CHK, 6: CHK, 9: CHK, 10: CHK}, fmt={3: "#,##0", 4: "#,##0", 5: "0.0%", 6: "0.0%", 9: "#,##0.00", 10: "#,##0.00"}); r += 1
put(ws, r + 1, ["Controladas em conjunto (equivalência patrimonial; nota 13.9.3, p. 74): BAESA (25,01%) e Chapecoense (51% — controle conjunto por acordo de acionistas)", "Geração/Transmissão", 380803, 275980, f"=C{r+1}/Parametros!$B$5", f"=ABS(D{r+1})/Parametros!$B$6", "Equivalência patrimonial", "Ler acordo de acionistas (julgamento de controle conjunto); obter DF auditadas; recalcular equivalência", "—", "—", "Não verificado"],
    fills={3: INP, 4: INP, 5: CHK, 6: CHK}, fmt={3: "#,##0", 4: "#,##0", 5: "0.0%", 6: "0.0%"})

# ---------------------------------------------------------------- Segmentos
ws = sheet("Segmentos_nota30", "Nota 30 — Informações por segmento 2025 (R$ mil)", ["Segmento", "Receita líquida", "Lucro antes dos impostos", "Aquisições de ativo contratual, imobilizado e intangível", "% do LAIR consolidado"], [24, 16, 16, 20, 14],
           "Fonte: DF25 nota 30, p. 103/71 (tabela em imagem; transcrição conferida pelos totais).")
seg = [("Distribuição", 36531151, 5331500, 4963868), ("Geração", 3272911, 2547712, 269975), ("Transmissão", 1635891, 350619, 804090), ("Comercialização", 2692087, -281094, 1500), ("Serviços", 234086, 259335, 71632), ("Outros", 1655, -116826, 1033)]
r = 5
for s in seg:
    put(ws, r, list(s) + [f"=C{r}/Parametros!$B$6"], fills={2: INP, 3: INP, 4: INP, 5: CHK}, fmt={2: "#,##0", 3: "#,##0", 4: "#,##0", 5: "0.0%"}); r += 1
put(ws, r, ["Soma (calculada)", f"=SUM(B5:B{r-1})", f"=SUM(C5:C{r-1})", f"=SUM(D5:D{r-1})", ""], bold=True, fmt={2: "#,##0", 3: "#,##0", 4: "#,##0"})
put(ws, r + 1, ["Divulgado (total consolidado)", 44367780, 8091246, 6112098, ""], fills={2: INP, 3: INP, 4: INP}, bold=True, fmt={2: "#,##0", 3: "#,##0", 4: "#,##0"})
put(ws, r + 2, ["Diferença", f"=B{r}-B{r+1}", f"=C{r}-C{r+1}", f"=D{r}-D{r+1}", "esperado 0/±1 (eliminações entre segmentos já líquidas na receita)"], fills={2: CHK, 3: CHK, 4: CHK}, fmt={2: "#,##0", 3: "#,##0", 4: "#,##0"})

# ---------------------------------------------------------------- Prestadores 402
ws = sheet("Prestadores_402", "PT 402 — Organizações prestadoras de serviços relevantes",
           ["Prestador / serviço", "Situação documental", "Fonte", "Processo e afirmações afetadas", "Relevante para as DF?", "Relatório tipo 1/2 disponível?", "Controles complementares da entidade usuária (exemplos)", "Procedimento planejado / alternativo", "Modalidade"],
           [28, 18, 22, 30, 14, 18, 40, 44, 18],
           "Identificação a partir das notas. 'Documentado' = citado na DF25; 'Hipótese' = serviço típico do setor, não comprovado. Nenhum relatório de asseguração de prestador foi obtido.")
prest = [
 ("Vivest e Fundação Família Previdência — administração dos planos de previdência e dos ativos garantidores","Documentado","DF25 nota 20, p. 86 e 92","Obrigação atuarial (606.924 no BP) e ativos do plano; avaliação, existência","Sim","Não obtido","Envio de cadastro de participantes; conciliação de contribuições; aprovação de premissas atuariais","Obter relatório do auditor do plano e confirmação dos ativos; testar dados censitários; usar especialista atuário (NBC TA 620)","Planejado por falta de acesso"),
 ("CCEE — contabilização e liquidação do mercado de curto prazo","Documentado","DF25 nota 3.11, p. 54; nota 12","Receita e custo de energia (comercialização, liquidação); ocorrência, exatidão","Sim","Não obtido (não verificado se existe)","Conferência dos relatórios de pré-liquidação; registro de contratos bilaterais","Conciliar relatórios de liquidação da CCEE com a contabilidade; inspecionar registros de contratos (planejado); usar dados abertos da CCEE apenas após validação da fonte (P07)","Planejado por falta de acesso"),
 ("Empreiteiras de construção da infraestrutura (alto grau de terceirização)","Documentado","DF25 nota 3, p. 47","Ativo contratual e receita/custo de construção (5.879.424 / 5.566.883); existência, exatidão, corte","Sim (afeta ativos da concessão)","Não aplicável (prestador operacional, não processa informação contábil)","Medição de obras e aceite; controle de ordens de serviço","Tratar como fornecedor (NBC TA 500): inspeção de medições e contratos; a NBC TA 402 só se aplica se o prestador processar informação que integra o sistema de informação","Planejado por falta de acesso"),
 ("Agentes arrecadadores (bancos) — cobrança das faturas e convênios de arrecadação","Documentado em parte (convênios de arrecadação na nota 23, p. 97); bancos = hipótese","DF25 nota 23, p. 97","Contas a receber e caixa; existência, corte","Sim","Não obtido","Conciliação diária de arquivos de retorno bancário","Conciliar arquivos de retorno × baixas de contas a receber; confirmação bancária (PT 505)","Planejado por falta de acesso"),
 ("Provedores de TI/nuvem (ERP, faturamento, medição)","Hipótese (não identificado na DF25)","—","Receita não faturada e faturamento (R03, R15)","A confirmar","Não obtido","Gestão de acessos e mudanças do lado da entidade","Indagar à administração sobre terceirização de TI; se houver, obter relatório tipo 2 com período coberto e avaliar exceções e subcontratadas (402 itens 13–18)","Planejado por falta de acesso"),
 ("Centros de serviços compartilhados do próprio grupo (CPFL Serviços, Finanças, Pessoas, Supri, Atende)","Documentado (entidades do grupo)","DF25 nota 13.1, p. 70/38","Folha, finanças, suprimentos, atendimento","Sim — mas são componentes do grupo, não prestadores externos","Não aplicável","—","Tratar no PT 600 (funções compartilhadas), não na NBC TA 402","Realizado com evidência pública (classificação)"),
]
r = 5
for p in prest: put(ws, r, list(p)); r += 1

os.makedirs(OUT, exist_ok=True); p = os.path.join(OUT, "CPFL_B05_Grupo_Prestadores.xlsx"); wb.save(p); print(p)
