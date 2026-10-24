# -*- coding: utf-8 -*-
"""Bloco 12 — revisão técnica automatizada: registro de revisão, matriz final de cobertura, conferências e revisão analítica final.
Revisão feita pela mesma IA que preparou o trabalho: NÃO é revisão independente."""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
OUT = sys.argv[1] if len(sys.argv) > 1 else '/root/cpfl/producao/00_Controle/CPFL_B12_Revisao_Tecnica_Cobertura.xlsx'
RV = '/root/cpfl/revisao_b12'
F = 'Arial'; H = Font(name=F, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3864')
TT = Font(name=F, bold=True, size=12); N = Font(name=F); B = Font(name=F, bold=True); BL = Font(name=F, color='0000FF')
INP = PatternFill('solid', fgColor='FFF2CC'); th = Side(style='thin', color='A6A6A6'); BD = Border(left=th, right=th, top=th, bottom=th)
MIL = '#,##0;(#,##0);"-"'; PCT = '0.0%;(0.0%);"-"'
G = {'Alta': PatternFill('solid', fgColor='F8CBAD'), 'Média': PatternFill('solid', fgColor='FFE699'), 'Baixa': PatternFill('solid', fgColor='E2EFDA'), 'Info': PatternFill('solid', fgColor='DDEBF7')}
wb = Workbook(); wb.remove(wb.active)
def sheet(name, title, widths, nota):
    ws = wb.create_sheet(name); ws['A1'] = title; ws['A1'].font = TT
    ws['A2'] = nota; ws['A2'].font = Font(name=F, italic=True, size=9)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A5'; return ws
def head(ws, r, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=r, column=i, value=c); x.font = H; x.fill = HF; x.border = BD; x.alignment = Alignment(wrap_text=True)
def put(ws, r, vals, fmt=None, inp=(), bold=False, fill=None):
    for i, v in enumerate(vals, 1):
        x = ws.cell(row=r, column=i, value=v); x.border = BD; x.font = B if bold else N
        x.alignment = Alignment(wrap_text=True, vertical='top')
        if fmt and i in fmt: x.number_format = fmt[i]
        if i in inp: x.font = BL; x.fill = INP
        if fill: x.fill = fill
AUTO = 'Revisão automatizada (segunda leitura crítica pela mesma IA que preparou o trabalho, em 23/10/2026). Não é revisão de auditor independente nem revisão humana.'

# ---------------- LEIA_ME
ws = wb.create_sheet('LEIA_ME'); ws.column_dimensions['A'].width = 120
for i, t in enumerate(['CPFL Energia S.A. — DF consolidadas 2025 — Bloco 12: revisão técnica e correção de inconsistências', 'UnB · Auditoria 1 (CCA0024) · Turma 02 · 2026.2 · Matrícula 231020069', AUTO,
    'Abas: Registro_Revisao (achado, gravidade, evidência, correção, estado) · Matriz_Cobertura_Final (33 obrigações do enunciado) · Conf_Normas · Conf_Numeros · Conf_Fontes_DF25 · Reproducao · PT520_Revisao_Final · Eventos_Corte · Pontos_Validacao',
    'Scripts: b12_check_normas.py, b12_check_numeros.py, b12_check_fontes.py (saídas JSON em revisao_b12/). Reexecução: python <script> Dossie_CPFL_2025_231020069_v0.8.md',
    'Gravidade: Alta = descumpre obrigação do enunciado ou altera conclusão; Média = afeta PT ou evidência sem mudar a conclusão; Baixa = número ou texto sem efeito na conclusão; Info = conferência sem achado.'], 1):
    ws.cell(row=i, column=1, value=t).font = TT if i == 1 else N; ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True)

# ---------------- Registro de revisão
ws = sheet('Registro_Revisao', 'Registro de revisão do Bloco 12', [9, 8, 40, 46, 46, 22], AUTO)
head(ws, 4, ['Ref.', 'Grav.', 'Achado', 'Evidência', 'Correção / tratamento', 'Estado'])
R = [
 ('RV12-01', 'Alta', '33 de 143 procedimentos com rótulo de modalidade fora das três categorias do enunciado ("Realizado", "Parcial", "Planejado (B12)", "Pendente — Bloco 11" etc.), incluindo rótulos desatualizados', 'Verificação automatizada sobre os 27 PTs (script em linha no Bloco 12); obrigação E-07', 'Rótulos normalizados nas fontes dos PTs; 0 fora das três categorias na v0.9 (reconferido)', 'Corrigido'),
 ('RV12-02', 'Média', 'Procedimento 520-05 (revisão analítica final, NBC TA 520 item 6) marcado "Pendente — Bloco 11" e não executado', 'PT 520, quadro de procedimentos; NORMAS p. 525', 'Executado no Bloco 12 (aba PT520_Revisao_Final; PT 520 seção 6). Sem variação nova inesperada; ativo setorial +228,6% explicado qualitativamente (nota 9, p. 64–65)', 'Corrigido'),
 ('RV12-03', 'Baixa', 'Quadro de liquidez (PT 570, planilha B10): debêntures 30.641.582; a fonte mostra 30.641.585. A "correção" do RV-B10 estava errada e criou falsa diferença de −3', 'DF25 nota 34 g, p. 113/81 (imagem ampliada a 300 dpi); soma da linha = 30.641.585', 'Valor corrigido; planilha B10 recalculada (0 erros; diferença da coluna Total = 0); DV21; PT 220 atualizado', 'Corrigido'),
 ('RV12-04', 'Média', 'Pesquisa de eventos subsequentes incompleta: a lista do RI de 2026 tinha itens não citados (renúncia de conselheiro 06/03; 21ª e 8ª emissões em março; 19ª emissão em setembro; atas de 15/10/2026)', 'RI CPFL, lista de 2026 relida em 23/10/2026 (FC32)', 'PT 560 seção 5 acrescentada; sem efeito em riscos, continuidade ou conclusão. Atas de 15/10/2026 e valores das emissões não lidos (limitação)', 'Corrigido com limitação'),
 ('RV12-05', 'Média', 'PB01-06: vigência da NBC TA 240 (R2) e 570 (R1) apoiada em fonte secundária', 'CRCSP (notícia de 02/10/2026) e Portal Contnews: DOU 25/09/2026, períodos iniciados a partir de 01/01/2027 (FC31)', 'Aplicação das versões 240 (R1) e 570 (2016) mantida; seção 2.3 atualizada. DOU não aberto diretamente', 'Resolvido (fonte secundária confiável)'),
 ('RV12-06', 'Info', 'Citações normativas item × página', '47 citações reais verificadas no texto do consolidado (NORMAS); 9 alertas eram falsos positivos (notas da DF ou itens A); itens A2, A6, A10, A13, A18, A23, A1, A5 e A7 localizados nas páginas citadas', 'Nenhuma correção necessária', 'Conferido'),
 ('RV12-07', 'Info', 'Números-chave consistentes entre PTs e planilhas', '20 valores-chave sem variante divergente (a única variante, 435.547, é a prova por movimentação ±1); 9 planilhas sem erro de fórmula', 'Nenhuma', 'Conferido'),
 ('RV12-08', 'Info', 'Números citados × páginas da DF25', '106 pares número-página: 24 localizados no texto, 58 em páginas-imagem já transcritas (B02), 24 não localizados por texto: são de outros documentos (ITR, DFs individuais, release), valores derivados ou tabelas-imagem. Conferência visual: p. 68, 111, 113, 114 (confirmados)', 'Achado RV12-03 decorre desta conferência', 'Conferido (amostra visual)'),
 ('RV12-09', 'Info', 'Reprodutibilidade da base e dos testes', 'Pasta limpa: SHA-256 da base 877d4a02…4bb0 idêntico; 501.000 linhas lidas; verificador sem falhas; total e 4 distribuidoras com diferença R$ 0,00; hashes dos reservados idênticos; T1–T4, MUS e ampliação com saídas idênticas; JSON difere só em "executado_em" e "duracao_s" (metadados)', 'Nenhuma', 'Conferido'),
 ('RV12-10', 'Baixa', 'Cópia da base no Drive não pôde ser lida pelo conector (limite de 10 MB)', 'Drive 03_Base_sintetica: 14.678.974 bytes, igual ao arquivo local lido integralmente', 'Conferência por tamanho exato; leitura integral da cópia publicada fica para o repositório (B13)', 'Limitação'),
 ('RV12-11', 'Info', 'Afirmações indevidas de execução, contato, assinatura, opinião ou revisão humana', '17 padrões pesquisados no dossiê; as ocorrências são negações ("nada foi enviado") ou descrição do relatório da PwC', 'Nenhuma', 'Conferido'),
 ('RV12-12', 'Info', 'Independência entre detecção e gabarito', 'Pré-registro (19:54) e congelamento (19:55:53) anteriores à abertura dos reservados (horários do Drive); "teste cego" só aparece negado', 'Nenhuma', 'Conferido'),
 ('RV12-13', 'Alta', 'PB01-01: conferência DF25 × base de dados da CVM não executada', 'dados.cvm.gov.br bloqueado no ambiente (proxy 403, 23/10/2026)', 'Mantida como limitação; ação opcional do aluno (Pontos_Validacao, ponto 1)', 'Pendente (depende do aluno)'),
 ('RV12-14', 'Média', 'Limitações remanescentes: hipóteses DH-01 (IRRF Jantus) e DH-02 (VNR Paulista); relatório do auditor 2024 (PB09-02); FR truncado (PB02-03); comunicados de partes relacionadas (PB10-02); código CVM e nome anterior (PB02-02)', 'PENDENCIAS.md v11', 'Sem acesso, não resolvíveis; mantidas nas conclusões limitadas do dossiê (seção 7)', 'Limitação final'),
]
for i, row in enumerate(R, 5): put(ws, i, list(row), fill=G[row[1]])

# ---------------- Matriz de cobertura final
ws = sheet('Matriz_Cobertura_Final', 'Matriz final de cobertura das obrigações do enunciado (FC10) — prova de aceite', [7, 44, 46, 42, 22], 'Estado em 23/10/2026 (fim do Bloco 12). "Dossiê" = Dossie_CPFL_2025_231020069_v0.9.')
head(ws, 4, ['ID', 'Obrigação', 'Onde é atendida (arquivo, PT, seção)', 'Prova de aceite', 'Estado'])
M = [
 ('E-01', 'Dossiê cobrindo EC e NBC TA 200 a 600', 'Dossiê, seção 5 (27 PTs) e seção 6 (programas)', '27 PTs × 13 campos sem vazio (RV12-01)', 'Atendida'),
 ('E-02', 'Usar DFs 2025, comparativos, notas e relatório do auditor', 'Dossiê seção 3; PTs 315, 320, 500–580; planilha B02', 'Saldos com documento e página; RV12-08', 'Atendida'),
 ('E-03', 'Julgamento, ceticismo, risco, planejamento, evidência', 'PTs 200, 300, 315, 320, 330, 500', 'Julgamentos com fundamento (DECISOES.md)', 'Atendida'),
 ('E-04', 'Fontes com URL, data, documento, página, exercício', 'Dossiê, Apêndice B (FC01–FC33)', 'Datas de acesso reais', 'Atendida'),
 ('E-05', 'Alterações de nome e código de negociação', 'Dossiê seção 3; PT EC', 'Nome anterior pelo FR (truncado); conferência no cadastro CVM pendente (PB02-02)', 'Atendida com limitação'),
 ('E-06', 'Não contatar empresa ou terceiros', 'Aviso do dossiê; PTs 260, 505, 580', 'Nenhum envio; minutas marcadas NÃO ENVIADAS', 'Atendida'),
 ('E-07', 'Classificar cada procedimento nas três modalidades', 'Quadro 5 de cada PT', '143 de 143 após RV12-01', 'Atendida (corrigida)'),
 ('E-08', 'Não apresentar confirmações, entrevistas, inspeções como realizadas', 'Todos os PTs', 'Varredura RV12-11 sem achado', 'Atendida'),
 ('E-09', 'Todos os PTs; "não se aplica" só com análise', '27 PTs; PT 402 e 501 (estoques) com análise de aplicabilidade', 'Índice da seção 4', 'Atendida'),
 ('E-10', 'Base simulada ≥ 500.000 linhas da conta de maior materialidade do ativo', 'CPFL_SINTETICA_Base_AFC_2025_registrada.csv.gz; PT 500', '501.000 linhas lidas; AFC = 36,5% do ativo', 'Atendida'),
 ('E-11', 'Registros individuais compatíveis com a conta', 'Projeto B06; PT 500', 'Unidade = bem cadastral; chaves únicas exceto 1.000 duplicidades do desenho', 'Atendida'),
 ('E-12', 'Justificar conta e materialidade', 'PT 320; PT 500; D-ETAPA-01; D-B04-01', 'Validado pelo aluno', 'Atendida'),
 ('E-13', 'Premissas e método de geração', 'CPFL_B06_Projeto_Base_Sintetica; config_base_cpfl_v2.json', 'Documento aprovado (D-B06-01/02)', 'Atendida'),
 ('E-14', 'Dicionário de variáveis', 'Planilha B07 (aba de dicionário); PT 500', '13 campos documentados', 'Atendida'),
 ('E-15', 'Conciliar base com saldo contábil', 'Planilha B07; PT 500', 'Diferença R$ 0,00 no total e por distribuidora (reconferida em RV12-09)', 'Atendida'),
 ('E-16', 'Usar a base: amostras, exceções, avaliação', 'PT 530, PT 450; planilha B08; testes_b08.py', 'Reexecução idêntica (RV12-09)', 'Atendida'),
 ('E-17', 'GitHub com código, instruções e PTs', 'Repositório (Bloco 13)', 'Link acessível', 'Pendente (B13)'),
 ('E-18', 'Erros inseridos documentados e comparáveis', 'Reservados (zip pós-congelamento); comparar_gabarito_b08.py; PT 530', 'Comparação após congelamento', 'Atendida'),
 ('E-19', 'Dados expressamente sintéticos', 'Campo origem_dados = SINTETICO; rótulos nos arquivos', 'Leitura do arquivo (cabeçalho e 501.000 linhas)', 'Atendida'),
 ('E-20', '9 campos mínimos em cada PT', 'Modelo de PT (dossiê seção 5)', 'Verificação automatizada sem campo vazio', 'Atendida'),
 ('E-21', 'Campo de revisão sem assinatura simulada', 'Campo 9 de cada PT', 'Em branco nos 27 PTs', 'Atendida'),
 ('E-22', 'Dossiê em PDF com referências cruzadas', 'Dossiê v0.9 PDF/DOCX', 'PDF gerado do mesmo DOCX', 'Atendida (versão final no B13)'),
 ('E-23', 'Base + dicionário no repositório ou link', 'Repositório (B13)', 'Link testado', 'Pendente (B13)'),
 ('E-24', 'Código de geração e análise', 'gerar/verificar_base_cpfl_v2.py; testes_b08.py; avaliacao_complementar_b08.py', 'Reexecução limpa (RV12-09)', 'Atendida'),
 ('E-25', 'Planilhas editáveis', 'Planilhas B02, B04, B05, B07–B12', '0 erros de fórmula (RV12-07)', 'Atendida'),
 ('E-26', 'Matriz de cobertura das normas', 'Controles B01 (Mapa_Normativo); esta aba; dossiê seção 2.3', 'Norma → PT; versões conferidas', 'Atendida'),
 ('E-27', 'Relação das fontes e links', 'Apêndice B', 'URL e data', 'Atendida'),
 ('E-28', 'README com matrícula, organização e reprodução', 'README.md (B13)', 'Passos testados', 'Pendente (B13)'),
 ('E-29', 'Link de download se a base exceder o limite do GitHub', 'Não necessário: base de 14,7 MB < 25 MB do upload pela web', 'Tamanho do arquivo', 'Atendida (verificar no B13)'),
 ('E-30', 'Seis programas por ciclo', 'Dossiê seção 6', 'PG-REC, PG-DISP, PG-EST, PG-IMOB, PG-PASS, PG-ESTIM', 'Atendida'),
 ('E-31', 'Conclusão integrada sem opinião profissional', 'Dossiê seção 7', 'Varredura RV12-11', 'Atendida'),
 ('E-32', 'Versões das normas aplicáveis a 2025', 'Dossiê seção 2.3', 'RV12-05', 'Atendida'),
 ('E-33', 'Postar link até 25/10/2026 e verificar acesso do professor', 'Ação do aluno (B13)', 'Teste em janela anônima', 'Pendente (aluno)'),
]
for i, row in enumerate(M, 5): put(ws, i, list(row), fill=(G['Média'] if row[4].startswith('Pendente') else None))
r = 5 + len(M) + 1
put(ws, r, ['', 'Atendidas (fórmula)', f'=COUNTIF(E5:E{4+len(M)},"Atendida*")', 'Pendentes', f'=COUNTIF(E5:E{4+len(M)},"Pendente*")'], bold=True)

# ---------------- conferências automatizadas
cn = json.load(open(os.path.join(RV, 'check_normas.json'), encoding='utf-8'))
ws = sheet('Conf_Normas', 'Citações normativas item × página (NORMAS = CONSOLIDAÇÃO DAS NORMAS AUD 1.pdf)', [10, 8, 12, 10, 10, 90], 'Gerado por b12_check_normas.py. "ok" = linha iniciada pelo número do item na página citada (ou na anterior, quando o item começa antes). Falsos positivos (notas da DF, itens A) conferidos à parte (RV12-06).')
head(ws, 4, ['Tipo', 'Item', 'Páginas', 'Localizado', 'Na página exata', 'Contexto'])
for i, x in enumerate(cn, 5): put(ws, i, [x['tipo'], x['item'], x['paginas'], 'sim' if x['ok'] else 'NÃO', 'sim' if x['exato'] else 'não', x['ctx'][-200:]])
cnum = json.load(open(os.path.join(RV, 'check_numeros.json'), encoding='utf-8'))
ws = sheet('Conf_Numeros', 'Números-chave: ocorrências e variantes próximas no dossiê', [24, 14, 12, 40], 'Gerado por b12_check_numeros.py.')
head(ws, 4, ['Número-chave', 'Valor (R$ mil)', 'Ocorrências', 'Variantes próximas (possível erro)'])
for i, (k, v) in enumerate(cnum.items(), 5): put(ws, i, [k, v['valor'], v['ocorrencias'], ', '.join(v['variantes']) or '—'])
cf = json.load(open(os.path.join(RV, 'check_fontes_df25.json'), encoding='utf-8'))
ws = sheet('Conf_Fontes_DF25', 'Números citados com página da DF25 × texto da página', [14, 14, 26, 90], 'Gerado por b12_check_fontes.py. Páginas 33–40 (demonstrações) são imagens transcritas no B02. Conferência visual adicional: p. 68 (1.285.173), p. 111 (4.958.307; 411.659), p. 113 (48.423.044; 12.916.273; 79.413.507; debêntures 30.641.585), p. 114 (1.700.000; 100.000).')
head(ws, 4, ['Número', 'Páginas', 'Estado', 'Contexto'])
for i, x in enumerate(cf, 5): put(ws, i, [x['numero'], ', '.join(map(str, x['paginas'])), x['estado'], x['ctx'][:200]])

# ---------------- Reprodução
ws = sheet('Reproducao', 'Reprodução em pasta limpa (23/10/2026, 21h05)', [44, 70, 18], 'Comandos: python gerar_base_cpfl_v2.py; python verificar_base_cpfl_v2.py; python testes_b08.py; python avaliacao_complementar_b08.py (estrutura de pastas: base_v2/ e testes_b08/ lado a lado).')
head(ws, 4, ['Arquivo', 'SHA-256 (16 primeiros) original = reprodução', 'Resultado'])
rp = [('CPFL_SINTETICA_Base_AFC_2025_registrada.csv.gz', '877d4a0256af842b (completo: 877d4a0256af842b859617f22c71678f30749e75b724c20cd205895ef85e4bb0)', 'Idêntico'),
      ('Reservado: base de referência', '48652ae2da56…', 'Idêntico'), ('Reservado: ponte', '7eabcc92f3d2…', 'Idêntico'), ('Reservado: registro de erros', '45e3ec3eba8d…', 'Idêntico'),
      ('T1_duplicidades.csv', '5b44a25b9239a200', 'Idêntico'), ('T2_baixados_com_valor.csv', 'f7c2c123f16a6499', 'Idêntico'), ('T3_recalculo_divergente.csv', 'e63489d299450916', 'Idêntico'),
      ('T4_residual_incoerente.csv', 'a8e7dfda972250e6', 'Idêntico'), ('MUS_amostra_resultados.csv', 'f681a63763e1afb5', 'Idêntico'), ('MUS_ampliacao_residual.csv', 'aa4fb53d279d77ac', 'Idêntico'),
      ('resultados_b08.json / avaliacao_complementar_b08.json', 'hashes diferentes', 'Diferem só em executado_em e duracao_s (metadados)'),
      ('Contagem de linhas (zcat | wc -l)', '501.001 = cabeçalho + 501.000 registros', 'OK (≥ 500.000)'), ('Verificador', 'falhas_de_geracao = {}; diferença total R$ 0,00', 'OK'), ('Tempo de geração', '16,8 s', '—')]
for i, row in enumerate(rp, 5): put(ws, i, list(row))

# ---------------- PT 520 revisão final
ws = sheet('PT520_Revisao_Final', 'PT 520 — Revisão analítica final das DF (item 6) — BP e DRE consolidados (R$ mil)', [40, 14, 14, 14, 10, 60], 'Fonte: DF25 p. 34–36/2–4 (valores da planilha B02, conferidos na imagem). Entradas em azul; variações por fórmula.')
head(ws, 4, ['Linha / indicador', '2025', '2024', 'Variação', 'Var. %', 'Leitura'])
L = [('Receita operacional líquida', 44367780, 42628210, 'Reajustes tarifários (nota 26)'), ('Custo com energia elétrica', 20376889, 19519468, 'Parcela A repassada'),
     ('Lucro operacional bruto', 14124196, 13568773, ''), ('Resultado do serviço', 10778490, 10493767, ''), ('Despesas financeiras', 4942969, 4343203, 'CDI/SELIC 14,90% (nota 34 e.2, p. 111)'),
     ('Lucro antes dos tributos', 8091246, 8088893, ''), ('IR + CS', 2348694, 2327339, ''), ('Ativo total', 81099930, 77129194, ''), ('Ativo financeiro da concessão', 29623619, 25209768, 'VNR acima da expectativa (PB09-01)'),
     ('Ativo contratual (circ. + não circ.)', 12859523, 11138706, ''), ('Ativo financeiro setorial (circ. + não circ.)', 2402307, 730965, 'Nota 9: CVA e transferência de 1.844.610'), ('Passivo financeiro setorial', 1990647, 1726197, ''),
     ('PIS/COFINS devolução consumidores', 3402844, 4936756, 'Transferência para o setorial'), ('Empréstimos + debêntures', 29055727, 28788338, ''), ('Mútuo com partes relacionadas', 4412650, 2485694, 'SGBP → SGEL (PT 550)'),
     ('Caixa + TVM', 2905082, 3547055, 'Dividendos e amortizações (não decomposto)'), ('Patrimônio líquido', 23504510, 21795029, '')]
for i, (n, a, b, t) in enumerate(L, 5):
    put(ws, i, [n, a, b, f'=B{i}-C{i}', f'=IF(C{i}=0,0,D{i}/C{i})', t], fmt={2: MIL, 3: MIL, 4: MIL, 5: PCT}, inp=(2, 3))
r = 5 + len(L) + 1
put(ws, r, ['Margem bruta', '=B7/B5', '=C7/C5', '=B{0}-C{0}'.format(r), '', 'Estável'], fmt={2: PCT, 3: PCT, 4: PCT}, bold=True); r += 1
put(ws, r, ['Resultado do serviço ÷ receita', '=B8/B5', '=C8/C5', '=B{0}-C{0}'.format(r), '', ''], fmt={2: PCT, 3: PCT, 4: PCT}, bold=True); r += 1
put(ws, r, ['Taxa efetiva (IR + CS ÷ LAIR)', '=B11/B10', '=C11/C10', '=B{0}-C{0}'.format(r), '', ''], fmt={2: PCT, 3: PCT, 4: PCT}, bold=True); r += 1
put(ws, r, ['Posição setorial líquida (ativo − passivo)', '=B15-B16', '=C15-C16', '=B{0}-C{0}'.format(r), '', '2025 = 411.660; nota 34 e.2 (p. 111): 411.659 (±1)'], fmt={2: MIL, 3: MIL, 4: MIL}, bold=True)

# ---------------- Eventos
ws = sheet('Eventos_Corte', 'Eventos subsequentes — corte final em 23/10/2026 (RI CPFL, lista de 2026 relida)', [16, 60, 30, 30], 'Itens novos em relação à versão do Bloco 10. Fonte: FC32.')
head(ws, 4, ['Data', 'Documento', 'Classificação', 'Efeito'])
ev = [('23/01 e 02/02/2026', 'Fatos relevantes — reorganização societária', 'Já tratado (nota 37.1)', 'Nenhum'), ('06/03/2026', 'Fato relevante — renúncia de membro do CA', 'Após o relatório', 'Nenhum nas DF de 2025'),
      ('10–17/03/2026', '21ª emissão (CPFL Energia) e 8ª emissão (CPFL Santa Cruz) de debêntures', 'Captação', 'Atenuante de liquidez (PT 570)'), ('01–10/09/2026', '19ª emissão de debêntures: escritura e anúncio de início', 'Captação', 'Idem'),
      ('29/04–02/10/2026', 'Avisos de dividendos (6)', 'Pagamento da proposta de 2025', 'Já considerado'), ('15/10/2026', 'Atas da 585ª e 586ª RCA', 'Não lidas', 'Limitação'),
      ('23/10/2026', 'Corte: nenhum fato relevante após 08/05/2026; 3T26 em 16/11/2026', 'Corte final', 'Conclusão inalterada')]
for i, row in enumerate(ev, 5): put(ws, i, list(row))

# ---------------- Pontos de validação
ws = sheet('Pontos_Validacao', 'Pontos para validação do aluno antes da montagem final (Bloco 13)', [6, 50, 50, 40], 'A validação não transforma hipótese em fato nem procedimento planejado em executado.')
head(ws, 4, ['Nº', 'Ponto', 'Recomendação', 'Alternativa'])
pv = [(1, 'Conferência DF25 × base da CVM (PB01-01) bloqueada no ambiente', 'Manter como limitação declarada (o PDF da DF25 é a fonte primária e conciliou internamente)', 'Aluno baixa dfp_cia_aberta_2025.zip (12 MB) de dados.cvm.gov.br e salva no Drive; executo a conferência'),
      (2, 'Atas de 15/10/2026 e valores das emissões de 2026 não lidos', 'Manter como limitação: não alteram as DF de 2025 nem a conclusão', 'Aluno abre as atas no RI e salva os PDFs no Drive para leitura'),
      (3, 'Entregar o dossiê v0.9 (sem pendência obrigatória de conteúdo) como base da versão final do Bloco 13', 'Aprovar; no B13 mudar o rótulo "em elaboração" para versão final 1.0 e montar o repositório', 'Pedir nova rodada de revisão de algum PT específico')]
for i, row in enumerate(pv, 5): put(ws, i, list(row))
os.makedirs(os.path.dirname(OUT), exist_ok=True); wb.save(OUT); print('OK', OUT)
