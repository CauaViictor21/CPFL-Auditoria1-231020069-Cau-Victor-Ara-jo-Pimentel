# -*- coding: utf-8 -*-
"""Planilha editável do Bloco 10: PT 501 (litígios, estoques, segmentos), PT 505 (controle de confirmações —
MODELOS NÃO ENVIADOS), PT 560 (linha do tempo de eventos subsequentes) e PT 570 (liquidez e covenants)."""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
OUT = sys.argv[1] if len(sys.argv) > 1 else '/root/cpfl/producao/02_Planejamento/CPFL_B10_Litigios_Confirmacoes_Eventos_Continuidade.xlsx'
F = 'Arial'; H = Font(name=F, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3864')
TT = Font(name=F, bold=True, size=12); N = Font(name=F); B = Font(name=F, bold=True)
BL = Font(name=F, color='0000FF'); INP = PatternFill('solid', fgColor='FFF2CC'); RED = PatternFill('solid', fgColor='FDE9E7')
th = Side(style='thin', color='A6A6A6'); BD = Border(left=th, right=th, top=th, bottom=th)
MIL = '#,##0;(#,##0);"-"'; PCT = '0.0%;(0.0%);"-"'; X2 = '0.00"x"'
wb = Workbook(); wb.remove(wb.active)
def sheet(name, title, widths, nota='Valores em R$ mil, salvo indicação. Azul/amarelo = dado de entrada com fonte; preto = fórmula.'):
    ws = wb.create_sheet(name); ws['A1'] = title; ws['A1'].font = TT
    ws['A2'] = nota; ws['A2'].font = Font(name=F, italic=True, size=9)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    return ws
def head(ws, r, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=r, column=i, value=c); x.font = H; x.fill = HF; x.border = BD; x.alignment = Alignment(wrap_text=True)
def put(ws, r, vals, fmt=None, inp=(), bold=False, fill=None):
    for i, v in enumerate(vals, 1):
        x = ws.cell(row=r, column=i, value=v); x.border = BD; x.font = B if bold else N
        x.alignment = Alignment(wrap_text=True, vertical='top')
        if fmt and i in fmt: x.number_format = fmt[i]
        if i in inp: x.font = BL; x.fill = INP
        if fill: x.fill = fill

GLOBAL, EXEC, TRIV = 404562.30, 303421.73, 20228.12

# ================= PT 501 — litígios
ws = sheet('PT501_Litigios', 'PT 501 — Litígios e reclamações (DF25 nota 22, p. 94–96/62–64)', [34, 14, 14, 14, 14, 14, 14, 40])
head(ws, 4, ['Natureza (provável — provisionado)', 'Saldo 31/12/2024', 'Adições', 'Reversões', 'Pagamentos', 'Atualização', 'Saldo 31/12/2025', 'Recálculo do saldo final'])
mov = [('Trabalhistas', 632533, 148117, -18954, -85938, 85382, 761139), ('Cíveis', 347977, 225660, -20052, -157871, 50463, 446178),
       ('Fiscais', 387442, 18356, -14203, -42161, 40601, 390034), ('Outros', 194028, 30978, -54266, -49097, 21909, 143552)]
for i, row in enumerate(mov, 5):
    put(ws, i, list(row) + [f'=SUM(B{i}:F{i})-G{i}'], fmt={k: MIL for k in range(2, 9)}, inp=tuple(range(2, 8)))
put(ws, 9, ['Total', '=SUM(B5:B8)', '=SUM(C5:C8)', '=SUM(D5:D8)', '=SUM(E5:E8)', '=SUM(F5:F8)', '=SUM(G5:G8)', '=SUM(H5:H8)'], fmt={k: MIL for k in range(2, 9)}, bold=True)
put(ws, 10, ['Total divulgado / BP (p. 35/3)', 1561980, 423111, -107475, -335067, 198354, 1740903, ''], fmt={k: MIL for k in range(2, 8)}, inp=tuple(range(2, 8)))
put(ws, 11, ['Diferença (calculado − divulgado)'] + [f'={c}9-{c}10' for c in 'BCDEFG'] + [''], fmt={k: MIL for k in range(2, 8)})
put(ws, 12, ['Depósitos judiciais (ativo)', 749936, '', '', '', '', 805776, 'Trabalhistas 274.015; cíveis 39.361; fiscais 449.343; outros 43.057 (2025)'], fmt={2: MIL, 7: MIL}, inp=(2, 7))
head(ws, 14, ['Perdas possíveis (não provisionadas; divulgadas)', '31/12/2025', '31/12/2024', 'Variação', '% materialidade global', '', '', 'Principais causas (p. 95/63)'])
pos = [('Trabalhistas', 824356, 684129, 'Acidentes de trabalho, periculosidade, horas extras'), ('Cíveis', 2801625, 2699688, 'Indenizações, danos elétricos, majoração tarifária, faixa de domínio'),
       ('Fiscais', 3414924, 3356569, 'IR e CS'), ('Fiscais — Outros', 3665609, 3246062, 'INSS, ICMS, FINSOCIAL, PIS e COFINS'), ('Regulatório', 41967, 195648, 'Fiscalização técnica, comercial e econômico-financeira')]
for i, (a, b, c, d) in enumerate(pos, 15):
    put(ws, i, [a, b, c, f'=B{i}-C{i}', f'=B{i}/{GLOBAL}', '', '', d], fmt={2: MIL, 3: MIL, 4: MIL, 5: PCT}, inp=(2, 3))
put(ws, 20, ['Total (quadro)', '=SUM(B15:B19)', '=SUM(C15:C19)', '=B20-C20', f'=B20/{GLOBAL}', '', '', 'Divulgado: 10.748.481 / 10.182.096'], fmt={2: MIL, 3: MIL, 4: MIL, 5: PCT}, bold=True)
put(ws, 21, ['Conferência com o total divulgado', '=B20-10748481', '=C20-10182096', '', '', '', '', 'Esperado 0'], fmt={2: MIL, 3: MIL})
head(ws, 23, ['Casos individuais divulgados', 'Valor', 'Prognóstico', '% materialidade global', '', '', '', 'Fonte / ponto de auditoria'])
cas = [('CEEE Prev — contribuições não paritárias (CPFL Transmissão)', 519456, 'Possível', 'p. 95/63'),
       ('ABRAGEL × União — garantia física (Renováveis)', 380500, 'Possível', 'p. 95/63'),
       ('ACP Estação Ecológica de Jataí — parcela possível', 51712, 'Possível', 'p. 95/63; restante 927.018 remoto (pedido 970.854)'),
       ('IRRF Jantus (Renováveis) — decisão definitiva desfavorável nov/2025', 139291, 'Obrigação em 31/12/2025', 'p. 96/64; quitado em jan/2026 por 60.950 (Lei 14.689/2023) → PT 560'),
       ('IRPJ/CSLL 2011 — ERSA/SMITA (CPFL Geração → Renováveis)', 504722, 'Possível', 'p. 96/64; CARF favorável nov/2024, cabe recurso'),
       ('IRPJ/CSLL 1999–2003 — ágio RGE (execução fiscal)', 650519, 'Possível (inferido do contexto)', 'p. 96/64; TRF4 desfavorável; REsp parcialmente provido em 26/11/2025'),
       ('CPFL Brasil × ANEEL — Notas Técnicas 23 e 81/2003 (ativo contingente)', 4678104, 'Ganho contingente (não reconhecido)', 'p. 94–95/62–63; cálculo da AGU até mar/2025')]
for i, (a, b, c, d) in enumerate(cas, 24):
    put(ws, i, [a, b, c, f'=B{i}/{GLOBAL}', '', '', '', d], fmt={2: MIL, 4: PCT}, inp=(2,))

# ================= PT 501 — estoques e segmentos
ws = sheet('PT501_Estoques_Segm', 'PT 501 — Estoques (aplicabilidade) e informações por segmento (nota 30)', [46, 16, 16, 16, 40])
head(ws, 4, ['Estoques', '2025', '2024', 'Variação', 'Fonte / leitura'])
put(ws, 5, ['Estoques (BP)', 236566, 221573, '=B5-C5', 'BP p. 34/2 (sem nota específica)'], fmt={2: MIL, 3: MIL, 4: MIL}, inp=(2, 3))
put(ws, 6, ['% do ativo total', '=B5/81099930', '=C5/77129194', '', ''], fmt={2: '0.00%', 3: '0.00%'})
put(ws, 7, ['% da materialidade global', f'=B5/{GLOBAL}', '', '', 'Abaixo da materialidade global'], fmt={2: PCT})
put(ws, 8, ['% da materialidade de execução', f'=B5/{EXEC}', '', '', 'Saldo < execução: inventário físico não obrigatório por relevância isolada; procedimento planejado por amostra'], fmt={2: PCT})
head(ws, 10, ['Segmento (nota 30, p. 103/71) — LAIR 2025', 'Valor', '', '', 'Fonte'])
seg = [('Distribuição', 5331500), ('Geração', 2547712), ('Transmissão', 350619), ('Comercialização', -281094), ('Serviços', 259335), ('Outros / eliminações', -116826)]
for i, (a, b) in enumerate(seg, 11): put(ws, i, [a, b, '', '', 'Nota 30 (transcrição do B05; DV18)'], fmt={2: MIL}, inp=(2,))
put(ws, 17, ['Soma dos segmentos', '=SUM(B11:B16)', '', '', ''], fmt={2: MIL}, bold=True)
put(ws, 18, ['LAIR consolidado (DRE p. 36/4)', 8091246, '', '', ''], fmt={2: MIL}, inp=(2,))
put(ws, 19, ['Diferença', '=B17-B18', '', '', 'Esperado 0'], fmt={2: MIL}, bold=True)
put(ws, 20, ['Participação da Distribuição no LAIR', '=B11/B18', '', '', 'Segmento que concentra o ativo financeiro da concessão (R01)'], fmt={2: PCT})

# ================= PT 505 — controle
ws = sheet('PT505_Controle', 'PT 505 — Controle de confirmações externas (MODELOS DIDÁTICOS — NENHUMA SOLICITAÇÃO FOI ENVIADA)', [8, 34, 30, 18, 14, 16, 14, 14, 14, 40],
           'Nenhum destinatário pessoal foi identificado; destinatários são categorias institucionais. Colunas de envio e resposta propositalmente vazias. Valores em R$ mil.')
head(ws, 4, ['Ref.', 'Saldo / assunto', 'Destinatário (categoria institucional)', 'Tipo', 'Valor no BP', 'Seleção', 'Data de envio', 'Resposta', 'Diferença', 'Procedimento alternativo se não houver resposta'])
cf = [('CF-01', 'Caixa e equivalentes + TVM; garantias; contas vinculadas', 'Instituições financeiras com saldo (lista da administração)', 'Positiva em branco', '=2229320+675762', 'Todas com saldo ≥ trivial + amostra das demais', 'Extratos de 31/12 e de janeiro; conciliações'),
      ('CF-02', 'Mútuo com a SGEL (PC 19.416 + PNC 4.393.234)', 'State Grid Europe Limited (área de tesouraria)', 'Positiva com saldo', 4412650, 'Integral (parte relacionada; R13)', 'Contratos, extratos de swap, recebimento dos recursos, atas'),
      ('CF-03', 'Debêntures detidas pela SGBP (nota 31 d)', 'SGBP (área financeira) e agente fiduciário da emissão', 'Positiva com saldo', 325563, 'Integral (parte relacionada)', 'Escritura, extrato do escriturador/B3, PU de 31/12'),
      ('CF-04', 'Empréstimos e debêntures (notas 18 e 19; o valor inclui as debêntures da CF-03)', 'Agentes fiduciários, BNDES e bancos credores', 'Positiva com saldo', '=3160115+6417553+1735590+17742469', 'Emissões ≥ execução + amostra estratificada', 'Escrituras, extratos B3, pagamentos subsequentes'),
      ('CF-05', 'Derivativos (swaps)', 'Contrapartes bancárias', 'Positiva com saldo', '=4513+591350', 'Integral por contraparte', 'Laudos de MTM e confirmações de negociação (term sheets)'),
      ('CF-06', 'Consumidores livres e concessionárias (recebíveis)', 'Clientes de maior saldo (pessoas jurídicas)', 'Positiva (maiores) e negativa (pulverizados)', '=6244460+133102', 'Estratificada; residenciais pulverizados não confirmados', 'Recebimentos subsequentes; faturas; medição'),
      ('CF-07', 'Liquidação CCEE', 'CCEE', 'Positiva / relatório de liquidação', None, 'Saldos de liquidação de dez/2025', 'Relatórios de contabilização CCEE (PT 402)'),
      ('CF-08', 'Litígios e reclamações (carta jurídica — NBC TA 501 item 10)', 'Escritórios de advocacia externos que patrocinam as causas', 'Carta de circularização jurídica', 1740903, 'Todos os escritórios com causas prováveis ou possíveis relevantes', 'Reunião com o jurídico interno; leitura de processos públicos')]
for i, row in enumerate(cf, 5):
    a, b, c, d, e, f_, g = row
    put(ws, i, [a, b, c, d, e, f_, 'NÃO ENVIADO', '—', '—', g], fmt={5: MIL}, inp=(5,) if isinstance(e, (int, float)) else ())
put(ws, 14, ['', 'Situação', 'Todos os modelos: NÃO ENVIADOS. Não há respostas, exceções ou confirmações recebidas.', '', '', '', '', '', '', ''], bold=True, fill=RED)

# ================= PT 560 — linha do tempo
ws = sheet('PT560_Linha_Tempo', 'PT 560 — Eventos subsequentes: linha do tempo com datas reais (corte da pesquisa: 23/10/2026)', [14, 58, 30, 30, 44],
           'Classificação conforme NBC TA 560 (R1) e CPC 24: "ajuste" = condição existente em 31/12/2025; "divulgação" = condição posterior relevante.')
head(ws, 4, ['Data', 'Evento', 'Janela da NBC TA 560', 'Classificação', 'Fonte'])
ev = [('31/12/2025', 'Data das demonstrações contábeis', 'Marco', '—', 'DF25'),
      ('jan/2026', 'Quitação do IRRF Jantus por 60.950 (decisão definitiva de nov/2025; débito de 139.291 em 31/12/2025)', 'Até o relatório (itens 6–9)', 'Condição existente em 31/12 (ajuste): verificar reconhecimento', 'DF25 nota 22 p. 96/64; nota 10 p. 68'),
      ('jan/2026', 'Emissão de debêntures: Paulista 1.700.000; Santa Cruz 100.000; Transmissão 100.000 (CDI + 0,30%)', 'Até o relatório', 'Não ajusta; divulgado', 'DF25 nota 37.3 p. 114/82'),
      ('02/02/2026', 'AGE aprova incorporação da CPFL Geração pela CPFL Brasil', 'Até o relatório', 'Não ajusta; divulgado (sem impacto no consolidado)', 'DF25 nota 37.1 p. 113–114/81–82'),
      ('até 23/02/2026', 'Curtailment: Lei 15.269/2025; termo de compromisso ainda incerto; nada reconhecido', 'Até o relatório', 'Incerteza divulgada (ativo contingente)', 'DF25 nota 37.2 p. 114/82'),
      ('23/02/2026', 'Autorização da emissão das DF pela administração', 'Marco', '—', 'DF25 nota 2.1 p. 45'),
      ('04/03/2026', 'Aprovação pelo Conselho de Administração e data do relatório do auditor (PwC, sem modificação)', 'Marco: fim da obrigação ativa (item 6)', '—', 'DF25 nota 2.1 p. 45; relatório p. 32'),
      ('20/03/2026', 'Publicação legal das DF (Valor Econômico)', 'Marco: divulgação das DF', '—', 'Ata metodológica da base (URL publicidade legal)'),
      ('06/04/2026', 'Despacho do MME (DOU) convoca a assinatura dos aditivos de concessão', 'Após a divulgação (itens 14–17)', 'Evento de 2026; não afeta 2025', 'Fato Relevante 08/05/2026'),
      ('22–24/04/2026', 'ANEEL: RTA da Paulista (+12,13%) e RTP da Santa Cruz (+15,12%)', 'Após a divulgação', 'Evento de 2026', 'ITR 1T26 nota 23.1 p. 55/22'),
      ('23/04/2026', 'Mútuo SGEL × CPFL Transmissão (200.000; EURIBOR + 1,0%)', 'Após a divulgação', 'Evento de 2026 (parte relacionada)', 'ITR 1T26 nota 23.2; comunicado RI 27/04/2026'),
      ('29/04/2026', 'AGO aprova dividendos de 2025 (proposta de 4.300 milhões)', 'Após a divulgação', 'Proposta já divulgada em 2025 (não ajusta)', 'Release 4T25 p. 3; dividendos pagos em parcelas até out/2026'),
      ('06/05/2026', 'Assinatura dos termos aditivos (prorrogação por 30 anos)', 'Após a divulgação', 'Evento de 2026: reclassificação prospectiva de 29.599.494', 'ITR 1T26 nota 23.3; ITR 2T26 notas 1.2 e 11'),
      ('11 e 13/05/2026', 'ITR 1T26 autorizado e revisado (PwC)', 'Após a divulgação', '—', 'ITR 1T26 p. 42 e 58'),
      ('06/07/2026', 'Mútuo SGEL × Transmissão, Renováveis e Brasil (1.805.000; CNH + 2,47%)', 'Após a divulgação', 'Evento de 2026 (parte relacionada)', 'ITR 2T26 nota 24.1; comunicado RI 07/07/2026'),
      ('jul/2026', 'Portaria MME 140/2026 (curtailment); manifestação de interesse em 07/08/2026', 'Após a divulgação', 'Evento de 2026; nada reconhecido no 2T26', 'ITR 2T26 nota 24.3 p. 60/26'),
      ('24/07/2026', 'Temporal em Ribeirão Preto (Paulista): baixas e provisões', 'Após a divulgação', 'Evento de 2026', 'ITR 2T26 nota 24.4 p. 60/26'),
      ('10 e 13/08/2026', 'ITR 2T26 autorizado e revisado (PwC)', 'Após a divulgação', '—', 'ITR 2T26 p. 44 e 65'),
      ('14/10/2026', 'Comunicado: conferência de resultados do 3T26 remarcada de 13/11 para 16/11/2026', 'Após a divulgação', 'Sem efeito', 'RI CPFL, lista de comunicados 2026'),
      ('23/10/2026', 'Corte da pesquisa: RI CPFL (lista de comunicados 2026), notícias de mercado; nenhum fato relevante novo após o ITR 2T26 identificado', 'Corte efetivo', '—', 'Pesquisa web em 23/10/2026')]
for i, row in enumerate(ev, 5): put(ws, i, list(row))

# ================= PT 570 — liquidez
ws = sheet('PT570_Liquidez', 'PT 570 — Liquidez, vencimentos e covenants (horizonte: 12 meses a partir de 31/12/2025, NBC TA 570 item 13)', [46, 14, 14, 14, 14, 14, 14, 16])
head(ws, 4, ['Passivos financeiros (fluxo não descontado; nota 34 g, p. 113/81)', '< 1 mês', '1–3 meses', '3–12 meses', '1–3 anos', '4–5 anos', '> 5 anos', 'Total'])
lq = [('Fornecedores', 3946688, 138471, 13240, 8027, 0, 231672, 4338098), ('Empréstimos e financiamentos — principal e encargos', 1149081, 478797, 2384105, 6065951, 651210, 691741, 11420885),
      ('Derivativos', 0, 419, 232121, 13274, 73836, 83744, 403394), ('Debêntures — principal e encargos', 0, 569096, 2636721, 5477838, 8120876, 13837054, 30641585),
      ('Outros', 276800, 996171, 94562, 78469, 92931, 80149, 1619082)]
for i, row in enumerate(lq, 5):
    put(ws, i, list(row) + [], fmt={k: MIL for k in range(2, 9)}, inp=tuple(range(2, 9)))
    ws.cell(row=i, column=9, value=f'=SUM(B{i}:G{i})-H{i}').number_format = MIL
ws.cell(row=4, column=9, value='Soma − total').font = B
put(ws, 10, ['Total calculado'] + [f'=SUM({c}5:{c}9)' for c in 'BCDEFGH'], fmt={k: MIL for k in range(2, 9)}, bold=True)
put(ws, 11, ['Total divulgado', 5372569, 2182955, 5360749, 11643558, 8938853, 14924359, 48423044], fmt={k: MIL for k in range(2, 9)}, inp=tuple(range(2, 9)))
put(ws, 12, ['Diferença (arredondamento)'] + [f'={c}10-{c}11' for c in 'BCDEFGH'], fmt={k: MIL for k in range(2, 9)})
head(ws, 14, ['Cobertura em 12 meses', 'Valor', '', '', '', '', '', 'Fonte / natureza'])
cob = [('Saídas contratuais em até 12 meses (passivos financeiros)', '=B11+C11+D11', 'Nota 34 g'),
       ('Caixa e equivalentes + TVM em 31/12/2025', '=2229320+675762', 'BP p. 34/2 (release: 2.905 milhões)'),
       ('Caixa líquido das atividades operacionais 2025 (referência histórica)', 7162797, 'DFC p. 39/7'),
       ('Dividendos de 2025 a pagar em 2026 (proposta)', -4300000, 'Release 4T25 p. 3; AGO 29/04/2026'),
       ('Captações de jan/2026 (debêntures)', 1900000, 'DF25 nota 37.3 p. 114/82'),
       ('Cobertura simples = (caixa + FCO − dividendos + captações) ÷ saídas', '=(B16+B17+B18+B19)/B15', 'Cenário didático; ignora investimentos (capex) e refinanciamentos'),
       ('Capital circulante líquido (AC − PC)', '=16616314-17875963', 'Nota 1.4 p. 44/12: −1.259.649'),
       ('Cobertura de amortizações de 12 meses divulgada pela companhia', 0.61, 'Release 4T25 p. 15 (caixa 2.905 ÷ amortizações 2026 de 4.732 milhões)'),
       ('Prazo médio da dívida (anos)', 4.13, 'Release 4T25 p. 15')]
for i, (a, b, c) in enumerate(cob, 15):
    put(ws, i, [a, b, '', '', '', '', '', c], fmt={2: MIL}, inp=(2,) if isinstance(b, (int, float)) else ())
ws.cell(row=20, column=2).number_format = X2; ws.cell(row=22, column=2).number_format = X2; ws.cell(row=23, column=2).number_format = '0.00'
head(ws, 25, ['Covenant consolidado (notas 18 e 19, p. 82–83/50–51)', 'Limite', 'Divulgado (release)', 'Folga', '', '', '', 'Natureza da evidência'])
put(ws, 26, ['Dívida líquida / EBITDA ajustado (máx.)', 3.75, 2.30, '=B26-C26', '', '', '', 'Informação da administração (release 4T25 p. 4 e 15); não auditada; definição pró-forma'], fmt={2: X2, 3: X2, 4: X2}, inp=(2, 3))
put(ws, 27, ['Dívida líquida no critério dos covenants (R$ milhões)', '', 30530, '', '', '', '', 'Release p. 15'], fmt={3: MIL}, inp=(3,))
put(ws, 28, ['EBITDA pró-forma (R$ milhões)', '', 13280, '', '', '', '', 'Release p. 15'], fmt={3: MIL}, inp=(3,))
put(ws, 29, ['Recálculo DL/EBITDA com os dados do release', '', '=C27/C28', '', '', '', '', 'Conferência aritmética'], fmt={3: X2})
put(ws, 30, ['Queda de EBITDA que levaria ao limite (dívida constante)', '', '=1-C27/(B26*C28)', '', '', '', '', 'Cenário didático'], fmt={3: PCT})
put(ws, 31, ['Aumento de dívida líquida que levaria ao limite (EBITDA constante; R$ milhões)', '', '=B26*C28-C27', '', '', '', '', 'Cenário didático'], fmt={3: MIL})
put(ws, 32, ['EBITDA ajustado / resultado financeiro (mín. 2,25) e PL / (PL + DL) (mín. 0,28)', '', 'não divulgado', '', '', '', '', 'Não recalculável com definições públicas'], )
put(ws, 33, ['Indicador didático com dados auditados: (dívida bruta − caixa − TVM) ÷ (resultado do serviço + D&A)', '', '=(29055727-2229320-675762)/(10778490+2397217)', '', '', '', '', 'NÃO é o covenant (sem consolidação proporcional nem mútuo SGEL)'], fmt={3: X2})
os.makedirs(os.path.dirname(OUT), exist_ok=True); wb.save(OUT); print('OK', OUT)
