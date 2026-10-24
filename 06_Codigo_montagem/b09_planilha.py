# -*- coding: utf-8 -*-
"""Planilha editável do Bloco 09: PT 510 (saldos iniciais), PT 520 (analíticos), PT 540 (estimativa do ativo
financeiro da concessão: retrospectiva e sensibilidade) e PT 550 (partes relacionadas). Valores em R$ mil."""
import re, ast, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else '/root/cpfl/producao/02_Planejamento/CPFL_B09_Abertura_Analiticos_Estimativas_PartesRel.xlsx'
src = open('/root/cpfl/scripts/b02_dados_reais.py', encoding='utf-8').read()
def get(name):
    return ast.literal_eval(re.search(r'^' + name + r' = (\[.*?\])\n', src, re.M | re.S).group(1))
F = 'Arial'; H = Font(name=F, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3864')
TT = Font(name=F, bold=True, size=12); N = Font(name=F); B = Font(name=F, bold=True)
BL = Font(name=F, color='0000FF'); GR = Font(name=F, color='008000'); INP = PatternFill('solid', fgColor='FFF2CC')
th = Side(style='thin', color='A6A6A6'); BD = Border(left=th, right=th, top=th, bottom=th)
MIL = '#,##0;(#,##0);"-"'; PCT = '0.0%;(0.0%);"-"'; PCT2 = '0.00%;(0.00%);"-"'
wb = Workbook(); wb.remove(wb.active)
def sheet(name, title, widths, nota='Valores em R$ mil, salvo indicação. Azul/amarelo = dado de entrada com fonte; preto = fórmula; verde = vínculo.'):
    ws = wb.create_sheet(name); ws['A1'] = title; ws['A1'].font = TT
    ws['A2'] = nota; ws['A2'].font = Font(name=F, italic=True, size=9)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    return ws
def head(ws, r, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=r, column=i, value=c); x.font = H; x.fill = HF; x.border = BD; x.alignment = Alignment(wrap_text=True)
def put(ws, r, vals, fmt=None, inp=(), bold=False, link=()):
    for i, v in enumerate(vals, 1):
        x = ws.cell(row=r, column=i, value=v); x.border = BD; x.font = B if bold else N
        x.alignment = Alignment(wrap_text=True, vertical='top')
        if fmt and i in fmt: x.number_format = fmt[i]
        if i in inp: x.font = BL; x.fill = INP
        if i in link: x.font = GR

# ======================= PT 510
ws = sheet('PT510_Abertura', 'PT 510 — Saldos de abertura: DF24 publicada × comparativo 2024 apresentado na DF25 (consolidado)', [52, 16, 16, 14, 40])
head(ws, 4, ['Linha (consolidado, 31/12/2024)', 'DF24 publicada', 'Comparativo na DF25', 'Diferença', 'Observação'])
df24 = {  # transcrito de DF24 CPFL.pdf p. 34 (ativo) e p. 35 (passivo/PL), conferência visual
 'Caixa e equivalentes de caixa': 1973401, 'Títulos e valores mobiliários': 1573654, 'Consumidores, concessionárias e permissionárias|C': 5883894,
 'Estoques': 221573, 'Dividendo e juros sobre o capital próprio': 19115, 'Imposto de renda e contribuição social a compensar|C': 715764,
 'Outros tributos a compensar|C': 653181, 'PIS/COFINS a compensar sobre ICMS|C': 1919127, 'Derivativos|C': 915621, 'Ativo financeiro setorial|C': 393443,
 'Ativo contratual|C': 774368, 'Outros ativos|C': 1508406,
 'Consumidores, concessionárias e permissionárias|N': 186371, 'Depósitos judiciais': 749936, 'Imposto de renda e contribuição social a compensar|N': 503271,
 'Outros tributos a compensar|N': 479460, 'PIS/COFINS a compensar sobre ICMS|N': 3652513, 'Ativo financeiro setorial|N': 337522, 'Derivativos|N': 193542,
 'Créditos fiscais diferidos': 266798, 'Ativo financeiro da concessão': 25209768, 'Investimentos em instrumentos patrimoniais': 144961, 'Outros ativos|N': 331087,
 'Investimentos': 581364, 'Imobilizado': 9752282, 'Intangível': 7824433, 'Ativo contratual|N': 10364338,
 'Fornecedores|PC': 3890827, 'Empréstimos e financiamentos|PC': 4587739, 'Debêntures|PC': 815233, 'Entidade de previdência privada|PC': 336398,
 'Imposto de renda e contribuição social a recolher|PC': 440904, 'Outros impostos, taxas e contribuições a recolher|PC': 784498, 'Mútuo com partes relacionadas|PC': 0,
 'Dividendo': 1510207, 'Derivativos|PC': 8178, 'Passivo financeiro setorial|PC': 927285, 'Provisões para desmobilização e gastos ambientais|PC': 12699,
 'PIS/COFINS devolução consumidores|PC': 1072326, 'Outras contas a pagar|PC': 2984242,
 'Fornecedores|PNC': 254364, 'Empréstimos e financiamentos|PNC': 7690254, 'Debêntures|PNC': 15695112, 'Entidade de previdência privada|PNC': 451514,
 'Imposto de renda e contribuição social a recolher|PNC': 245142, 'Outros impostos, taxas e contribuições a recolher|PNC': 997778, 'Débitos fiscais diferidos': 2665302,
 'Provisões para riscos fiscais, cíveis e trabalhistas': 1561980, 'Mútuo com partes relacionadas|PNC': 2485694, 'Derivativos|PNC': 272377,
 'Passivo financeiro setorial|PNC': 798912, 'Provisões para desmobilização e gastos ambientais|PNC': 151778, 'PIS/COFINS devolução consumidores|PNC': 3864430,
 'Outras contas a pagar|PNC': 828992,
 'Capital social': 9388071, 'Reserva de capital': -1394956, 'Reserva legal': 1877614, 'Reserva estatutária - reforço de capital de giro': 7033014,
 'Reserva de lucro a realizar': 3075193, 'Dividendo|PL': 1855190, 'Resultado abrangente acumulado': -1102363}
grupos = [('C', get('circ'), 'Ativo circulante', 16551546), ('N', get('nc'), 'Ativo não circulante', 60577648),
          ('PC', get('pc'), 'Passivo circulante', 17370535), ('PNC', get('pnc'), 'Passivo não circulante', 37963630),
          ('PL', get('pl'), 'PL dos controladores', 20731762)]
r = 5; totrows = {}; linhas = []
for g, lst, nome, tot24 in grupos:
    put(ws, r, [nome, '', '', '', ''], bold=True); r += 1; a = r
    for (lin, nota, v25, v24) in lst:
        k = lin + '|' + g
        d = df24.get(k, df24.get(lin))
        if g == 'PL' and lin == 'Dividendo': d = df24['Dividendo|PL']
        obs = 'Rubrica renomeada: na DF24 "Mútuo entre coligadas, controladas e controladora"' if 'Mútuo' in lin else ''
        put(ws, r, [lin, d, v24, f'=C{r}-B{r}', obs], fmt={2: MIL, 3: MIL, 4: MIL}, inp=(2, 3)); linhas.append(r); r += 1
    put(ws, r, [f'Soma das linhas — {nome}', f'=SUM(B{a}:B{r-1})', f'=SUM(C{a}:C{r-1})', f'=C{r}-B{r}', ''], fmt={2: MIL, 3: MIL, 4: MIL}, bold=True); s = r; r += 1
    put(ws, r, [f'Total publicado na DF24 — {nome}', tot24, '', f'=B{s}-B{r}', 'Diferença de soma = arredondamento (DF25 nota 2.4, p. 46)'], fmt={2: MIL, 4: MIL}, inp=(2,)); r += 2
    totrows[g] = s
put(ws, r, ['PL de não controladores', 1063267, 1063267, f'=C{r}-B{r}', 'DF24 p. 35; DF25 p. 35'], fmt={2: MIL, 3: MIL, 4: MIL}, inp=(2, 3)); r += 1
put(ws, r, ['Total do ativo (DF24 p. 34)', 77129194, 77129194, f'=C{r}-B{r}', ''], fmt={2: MIL, 3: MIL, 4: MIL}, inp=(2, 3)); r += 1
put(ws, r, ['Total do passivo e PL (DF24 p. 35)', 77129194, 77129194, f'=C{r}-B{r}', ''], fmt={2: MIL, 3: MIL, 4: MIL}, inp=(2, 3)); r += 2
head(ws, r, ['DRE consolidada 2024', 'DF24 (p. 36/4)', 'Comparativo DF25 (p. 36/4)', 'Diferença', '']); r += 1
for n_, v in [('Receita operacional líquida', 42628210), ('Lucro operacional bruto', 13568773), ('Resultado do serviço', 10493767),
              ('Lucro antes dos tributos', 8088893), ('Lucro líquido do exercício', 5761554), ('Lucro atribuído aos controladores', 5457652)]:
    put(ws, r, [n_, v, v, f'=C{r}-B{r}', ''], fmt={2: MIL, 3: MIL, 4: MIL}, inp=(2, 3)); r += 1
r += 1
put(ws, r, ['Linhas individuais comparadas', len(linhas), '', '', ''], bold=True); r += 1
put(ws, r, ['Linhas individuais com diferença ≠ 0', '=' + '+'.join(f'(D{x}<>0)*1' for x in linhas), '', '', 'Esperado 0. Totais: diferenças ±1/2 só de soma (arredondamento)'], bold=True)

# ======================= PT 520
ws = sheet('PT520_Planejamento', 'PT 520 / 315 — Analíticos de avaliação de risco (variações e indicadores, consolidado)', [44, 15, 15, 14, 12, 50],
           'Uso: avaliação de risco (NBC TA 315). Não são procedimentos substantivos. Valores em R$ mil.')
head(ws, 4, ['Item', '2025', '2024', 'Variação', 'Var. %', 'Fonte / leitura'])
it = [('Receita operacional líquida', 44367780, 42628210, 'DF25 DRE p. 36/4'), ('Custo com energia elétrica', -20376889, -19519468, 'p. 36/4'),
      ('Lucro operacional bruto', 14124196, 13568773, 'p. 36/4'), ('Resultado do serviço', 10778490, 10493767, 'p. 36/4'),
      ('Despesas financeiras', -4942969, -4343203, 'p. 36/4; nota 29'), ('Lucro antes dos tributos', 8091246, 8088893, 'p. 36/4'),
      ('Lucro líquido', 5742552, 5761554, 'p. 36/4'), ('Ativo circulante', 16616314, 16551546, 'BP p. 34/2'), ('Passivo circulante', 17875963, 17370535, 'BP p. 35/3'),
      ('Ativo total', 81099930, 77129194, 'BP p. 34/2'), ('Ativo financeiro da concessão', 29623619, 25209768, 'BP p. 34/2; nota 11'),
      ('Dívida bruta (empréstimos + debêntures, circ. + não circ.)', '=3160115+6417553+1735590+17742469', '=4587739+7690254+815233+15695112', 'BP p. 35/3'),
      ('Mútuo com partes relacionadas (total)', '=19416+4393234', 2485694, 'BP p. 35/3; nota 31'), ('Patrimônio líquido total', 23504510, 21795029, 'BP p. 35/3'),
      ('Caixa líquido das atividades operacionais', 7162797, 6788942, 'DFC p. 39/7')]
for i, (a, b, c, d) in enumerate(it, 5):
    put(ws, i, [a, b, c, f'=B{i}-C{i}', f'=IF(C{i}=0,"",D{i}/ABS(C{i}))', d], fmt={2: MIL, 3: MIL, 4: MIL, 5: PCT}, inp=(2, 3))
r = 5 + len(it) + 1
head(ws, r, ['Indicador', '2025', '2024', 'Variação (p.p./x)', '', 'Leitura para o risco']); r += 1
ind = [('Margem bruta', '=B7/B5', '=C7/C5', 'Estável'), ('Margem do LAIR', '=B10/B5', '=C10/C5', 'Leve queda'),
       ('Liquidez corrente (AC/PC)', '=B12/B13', '=C12/C13', 'Abaixo de 1: CCL negativo de 1.259.649 (nota 1.4, p. 44/12) → R17 (PT 570)'),
       ('Dívida bruta / PL', '=B16/B18', '=C16/C18', 'Alavancagem'), ('AFC / ativo total', '=B15/B14', '=C15/C14', 'Concentração crescente → R01'),
       ('Despesa financeira / (dívida + mútuo) média', '=-B9/((B16+B17+C16+C17)/2)', '', 'Custo médio aproximado (aproximação: dívida + mútuo)'),
       ('FCO / lucro líquido', '=B19/B11', '=C19/C11', 'Qualidade do lucro')]
for a, b, c, d in ind:
    put(ws, r, [a, b, c, f'=IF(C{r}="","",B{r}-C{r})', '', d], fmt={2: '0.00%' if 'Liquidez' not in a and 'FCO' not in a and 'PL' not in a else '0.00', 3: '0.00%' if 'Liquidez' not in a and 'FCO' not in a and 'PL' not in a else '0.00', 4: '0.00%' if 'Liquidez' not in a and 'FCO' not in a and 'PL' not in a else '0.00'}); r += 1

ws = sheet('PT520_Substantivo_VNR', 'PT 520 — Procedimento analítico substantivo: ajuste a valor justo (VNR) do ativo financeiro da concessão', [46, 16, 16, 16, 16, 16, 40],
           'Expectativa = saldo inicial × IPCA do ano + transferências do ano × IPCA/2 (meio período). Aproximação didática; NÃO é o modelo da administração (BRR homologada item a item).')
head(ws, 4, ['Parâmetro', 'Valor', '', '', '', '', 'Fonte / natureza'])
put(ws, 5, ['IPCA 2025', 0.0426, '', '', '', '', 'IBGE (acumulado jan–dez/2025); DF25 nota 34 e.2, p. 111/79 ("taxa no período")'], fmt={2: PCT2}, inp=(2,))
put(ws, 6, ['IPCA 2024', 0.0483, '', '', '', '', 'IBGE (acumulado jan–dez/2024)'], fmt={2: PCT2}, inp=(2,))
put(ws, 7, ['Materialidade de execução', 303421.73, '', '', '', '', 'D-B04-01'], fmt={2: MIL}, inp=(2,))
put(ws, 8, ['Limite de investigação (30% da execução)', '=0.3*B7', '', '', '', '', 'Julgamento: precisão exigida por risco significativo R01 (NBC TA 520 item 5(c)–(d))'], fmt={2: MIL})
head(ws, 10, ['Entidade / ano', 'Saldo inicial', 'Transferências do ativo contratual', 'VNR registrado', 'Expectativa', 'Diferença (reg. − esp.)', 'Fonte'])
rows = [('Consolidado 2025', 25209768, 3368915, 1285173, 5, 'DF25 nota 11 p. 68'), ('CPFL Paulista 2025', 10688700, 1383147, 635251, 5, 'DF Paulista nota 11 p. 29'),
        ('RGE 2025', 10927569, 1458964, 469006, 5, 'DF RGE nota 11 p. 31'), ('CPFL Piratininga 2025', 3343565, 348021, 171723, 5, 'DF Piratininga nota 11 p. 29'),
        ('CPFL Santa Cruz 2025 (por diferença)', 249934, 178783, 9193, 5, 'Consolidado − três DFs (PT 500)'),
        ('Consolidado 2024', 21617521, 2690645, 1104267, 6, 'DF25 nota 11 p. 68 (comparativo)'), ('CPFL Paulista 2024', 9148138, 1139005, 469588, 6, 'DF Paulista nota 11 p. 29'),
        ('RGE 2024', 9352489, 1207432, 476065, 6, 'DF RGE nota 11 p. 31'), ('CPFL Piratininga 2024', 2933783, 283428, 148482, 6, 'DF Piratininga nota 11 p. 29')]
for i, (a, s0, t, v, ipr, f_) in enumerate(rows, 11):
    put(ws, i, [a, s0, t, v, f'=B{i}*$B${ipr}+C{i}*$B${ipr}/2', f'=D{i}-E{i}', f_], fmt={2: MIL, 3: MIL, 4: MIL, 5: MIL, 6: MIL}, inp=(2, 3, 4))
r = 11 + len(rows)
head(ws, r + 1, ['Avaliação', 'Valor', '', '', '', '', '']); r += 2
put(ws, r, ['Diferença consolidada 2025 (R$ mil)', '=F11', '', '', '', '', ''], fmt={2: MIL}, link=(2,)); r += 1
put(ws, r, ['Diferença em % da expectativa', '=F11/E11', '', '', '', '', ''], fmt={2: PCT}); r += 1
put(ws, r, ['Excede o limite de investigação?', '=IF(ABS(F11)>B8,"SIM — investigar (520 item 7)","NÃO")', '', '', '', '', ''], bold=True); r += 1
put(ws, r, ['Parcela da diferença concentrada na CPFL Paulista', '=F12/F11', '', '', '', '', 'Distribuidora sem RTP em 2025 (RTA em abril; DF25 nota 26.2 p. 101/69)'], fmt={2: PCT}); r += 1
put(ws, r, ['Diferença consolidada 2024 (retrospectiva)', '=F16', '', '', '', '', 'Modelo de expectativa aderente em 2024'], fmt={2: MIL}, link=(2,)); r += 1
put(ws, r, ['Diferença 2024 em % da expectativa', '=F16/E16', '', '', '', '', ''], fmt={2: PCT}); r += 2
head(ws, r, ['Conciliação com a nota 34 (ganho reconhecido no resultado)', 'Valor', '', '', '', '', 'Fonte']); r += 1
put(ws, r, ['Ajuste a valor justo (VNR) 2025', '=D11', '', '', '', '', 'Nota 11 p. 68'], fmt={2: MIL}, link=(2,)); a = r; r += 1
put(ws, r, ['(−) Baixa da atualização relacionada ao ativo 2025', -113091, '', '', '', '', 'Nota 11 p. 68 (texto)'], fmt={2: MIL}, inp=(2,)); r += 1
put(ws, r, ['Ganho líquido calculado', f'=B{a}+B{r-1}', '', '', '', '', ''], fmt={2: MIL}, bold=True); c_ = r; r += 1
put(ws, r, ['Ganho divulgado na nota 34 c)', 1172082, '', '', '', '', 'DF25 p. 110/78'], fmt={2: MIL}, inp=(2,)); r += 1
put(ws, r, ['Diferença', f'=B{c_}-B{r-1}', '', '', '', '', 'Esperado 0'], fmt={2: MIL}, bold=True); r += 1
put(ws, r, ['Mesma conciliação 2024: 1.104.267 − 96.326 − 1.007.941', '=1104267-96326-1007941', '', '', '', '', 'Notas 11 e 34 (comparativos)'], fmt={2: MIL})

# ======================= PT 540
ws = sheet('PT540_Sensibilidade', 'PT 540 — Ativo financeiro da concessão: sensibilidade reproduzível (aproximação didática, não recálculo do modelo da administração)', [52, 16, 16, 16, 16, 40])
head(ws, 4, ['Entrada', 'Valor', '', '', '', 'Fonte'])
ent = [('Saldo do ativo financeiro da concessão 31/12/2025', 29623619, 'BP p. 34; nota 11 p. 68'), ('Saldo inicial 2025', 25209768, 'Nota 11 p. 68'),
       ('Transferências do ativo contratual 2025', 3368915, 'Nota 11 p. 68'), ('Materialidade global', 404562.30, 'D-B04-01'),
       ('Materialidade de execução', 303421.73, 'D-B04-01'), ('Lucro antes dos tributos 2025', 8091246, 'DRE p. 36/4')]
for i, (a, v, f_) in enumerate(ent, 5): put(ws, i, [a, v, '', '', '', f_], fmt={2: MIL}, inp=(2,))
put(ws, 11, ['Base sujeita à atualização (inicial + transferências/2)', '=B6+B7/2', '', '', '', 'Aproximação: transferências no meio do ano'], fmt={2: MIL})
head(ws, 13, ['Cenário', 'Variação do parâmetro', 'Efeito no saldo e no resultado', '% da materialidade global', '% do LAIR', 'Fundamento'])
sc = [('S1 IPCA +0,5 p.p.', 0.005, '=B11*B14', 'Erro de índice/defasagem na atualização'), ('S1 IPCA −0,5 p.p.', -0.005, '=B11*B15', ''),
      ('S1 IPCA +1,0 p.p.', 0.01, '=B11*B16', ''), ('S1 IPCA −1,0 p.p.', -0.01, '=B11*B17', ''),
      ('S2 Base homologada (BRR/VNR) −2%', -0.02, '=B5*B18', 'Ajuste na próxima revisão tarifária (nota 4, p. 59)'),
      ('S2 Base homologada (BRR/VNR) +2%', 0.02, '=B5*B19', ''), ('S2 Base homologada −5%', -0.05, '=B5*B20', ''), ('S2 Base homologada +5%', 0.05, '=B5*B21', '')]
for i, (a, v, f_, fu) in enumerate(sc, 14):
    put(ws, i, [a, v, f_, f'=C{i}/$B$8', f'=C{i}/$B$10', fu], fmt={2: PCT2, 3: MIL, 4: PCT, 5: PCT2}, inp=(2,))
head(ws, 23, ['Reprodução da sensibilidade divulgada (nota 34 e.2, p. 111/79) — instrumentos financeiros ativos indexados ao IPCA', 'Taxa', 'Efeito calculado', 'Efeito divulgado', 'Diferença', 'Fonte']); r = 24
put(ws, r, ['Exposição ativa ao IPCA (inclui o ativo financeiro da concessão)', 29653297, '', '', '', 'p. 111/79'], fmt={2: MIL}, inp=(2,)); r += 1
for nm, tx, dv in [('Cenário provável', 0.0963, 2855613), ('Redução de 25% do índice', '=0.0963*0.75', 2141709), ('Redução de 50% do índice', '=0.0963*0.5', 1427806)]:
    put(ws, r, [nm, tx, f'=$B$24*B{r}', dv, f'=C{r}-D{r}', 'Efeito em 12 meses (projeção da companhia)'], fmt={2: PCT2, 3: MIL, 4: MIL, 5: MIL}, inp=(4,) if not isinstance(tx, float) else (2, 4)); r += 1
put(ws, r, ['Perda máxima divulgada (cenário provável − 50%)', '', '=C25-C27', '', '', 'Comparar com a materialidade'], fmt={3: MIL}, bold=True); r += 2
head(ws, r, ['Renovação das concessões — evento posterior (2026)', 'Valor', '% do AFC 31/12/2025', '', '', 'Fonte']); r += 1
for nm, v, fu in [('CPFL Paulista — transferido ao intangível', 12737816, 'ITR 2T26 nota 11, p. 49 do PDF / 15 impressa'), ('CPFL Piratininga', 3686631, 'idem'),
                  ('CPFL RGE', 13175047, 'idem')]:
    put(ws, r, [nm, v, f'=B{r}/$B$5', '', '', fu], fmt={2: MIL, 3: PCT}, inp=(2,)); r += 1
put(ws, r, ['Total reclassificado (aplicação prospectiva, após 06/05/2026)', f'=SUM(B{r-3}:B{r-1})', f'=B{r}/$B$5', '', '', 'Total divulgado: 29.599.494'], fmt={2: MIL, 3: PCT}, bold=True); r += 1
put(ws, r, ['Conferência com o total divulgado', f'=B{r-1}-29599494', '', '', '', 'Esperado 0'], fmt={2: MIL})

# ======================= PT 550
ws = sheet('PT550_PartesRel', 'PT 550 — Partes relacionadas: saldos e transações divulgados (DF25 nota 31, p. 103–104/71–72)', [50, 14, 14, 14, 14, 14, 14, 14, 14])
head(ws, 4, ['Natureza / contraparte', 'Ativo 2025', 'Ativo 2024', 'Passivo 2025', 'Passivo 2024', 'Receita 2025', 'Receita 2024', 'Despesa 2025', 'Despesa 2024'])
pr = [('Outras operações financeiras — SGBP (debêntures CPFL Paulista)', 0, 0, 325563, 1507938, 0, 0, -97450, -177955),
      ('Energia e encargos — entidades sob controle comum', 0, 0, 0, 0, 2050, 1314, -424459, -410890),
      ('Energia e encargos — empreendimentos controlados em conjunto e outros', 0, 0, 76888, 104378, 678, 474, -897183, -965058),
      ('Intangível, imobilizado, serviços — entidades sob controle comum', 0, 15355, 557, 702, 0, 0, -25979, -96285),
      ('Intangível, imobilizado, serviços — controlados em conjunto e outros', 556, 250, 0, 0, 10813, 6614, 0, 0),
      ('Mútuo — entidades sob controle comum (SGEL)', 0, 0, 4412650, 0, 0, 0, -41667, 0),
      ('Mútuo — SGBP (liquidado antecipadamente em dez/2025)', 0, 0, 0, 2485694, 0, 0, -368077, -311573),
      ('Dividendos e JCP — controlados em conjunto', 11194, 19022, 0, 0, 0, 0, 0, 0),
      ('Outros — Instituto CPFL', 0, 0, 467, 0, 0, 0, -5299, -5115)]
for i, row in enumerate(pr, 5): put(ws, i, list(row), fmt={k: MIL for k in range(2, 10)}, inp=tuple(range(2, 10)))
e = 4 + len(pr)
put(ws, e + 1, ['Total'] + [f'=SUM({get_column_letter(c)}5:{get_column_letter(c)}{e})' for c in range(2, 10)], fmt={k: MIL for k in range(2, 10)}, bold=True)
put(ws, e + 3, ['Conciliação: mútuo passivo (nota 31) × BP (PC 19.416 + PNC 4.393.234)', '', '', '=D10+D11-(19416+4393234)', '=E10+E11-2485694', '', '', '', ''], fmt={4: MIL, 5: MIL})
put(ws, e + 4, ['Variação do mútuo passivo 2025 × 2024 (%)', '', '', '=(D10+D11)/(E10+E11)-1', '', '', '', '', ''], fmt={4: PCT})
put(ws, e + 5, ['Remuneração do pessoal-chave da administração', '', '', '', '', '', '', -80523, -80597], fmt={8: MIL, 9: MIL}, inp=(8, 9))
put(ws, e + 6, ['Despesas com partes relacionadas / despesas operacionais e financeiras (aprox.)', '', '', '', '', '', '', '=-H14/(20376889+4299812+5566883+985664+1722799+637243+4942969)', ''], fmt={8: PCT2})
os.makedirs(os.path.dirname(OUT), exist_ok=True); wb.save(OUT); print('OK', OUT)
