# -*- coding: utf-8 -*-
"""Bloco 01 - Controles de execução do dossiê CPFL 2025 (matrícula 231020069).
Gera CPFL_B01_Controles_Execucao.xlsx. Execução: python b01_controles.py <pasta_saida>
Todos os dados abaixo foram obtidos nesta sessão (22/10/2026) por leitura efetiva
das fontes indicadas; valores sem fonte aberta estão marcados como pendentes."""
import sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else "."
HOJE = "22/10/2026"
PREP = "Matrícula 231020069 (preparação assistida por IA, Claude)"
F = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
BRD = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
wb.remove(wb.active)

def sheet(name, title, header, rows, widths=None, note=None):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = Font(name=F, bold=True, size=12, color="1F3864")
    ws["A2"] = f"Preparador: {PREP} | Data: {HOJE} | Revisão: ____ (campo em branco — sem revisão de terceiros simulada)"
    ws["A2"].font = Font(name=F, italic=True, size=8, color="595959")
    r0 = 3
    if note:
        ws["A3"] = note
        ws["A3"].font = Font(name=F, size=9, color="7F6000")
        ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(header))
        ws.row_dimensions[3].height = 42
        r0 = 4
    for j, h in enumerate(header, 1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font = Font(name=F, bold=True, color="FFFFFF", size=9)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BRD
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=F, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BRD
    widths = widths or [18] * len(header)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    return ws, r0

# ---------------------------------------------------------------- 0 LEIA-ME
ws = wb.create_sheet("LEIA_ME")
txt = [
    "CONTROLES DE EXECUÇÃO — BLOCO 01 — Dossiê de auditoria didático CPFL Energia S.A. (CPFE3), exercício 2025 (consolidado)",
    "UnB · Auditoria 1 (CCA0024) · Turma 02 · 2026.2 · Matrícula 231020069 · Entrega 25/10/2026",
    "",
    "Finalidade: registrar o que foi efetivamente verificado em 22/10/2026 no início da execução (Bloco 01) e controlar a cobertura das exigências do enunciado.",
    "Natureza: trabalho ACADÊMICO. Nada aqui é opinião profissional, nem evidência obtida junto à companhia. Não houve contato com a CPFL, auditores ou terceiros.",
    "Legenda de estados: CONCLUÍDO = arquivo/verificação existe e foi conferido; EM ANDAMENTO = parcialmente produzido; PENDENTE = não iniciado; INSUMO = material preparatório (não é PT finalizado).",
    "Legenda de modalidades (enunciado, seção 2): PÚBLICA = Realizado com evidência pública; PLANEJADO = Planejado por falta de acesso; SIMULADO = Simulado para fins didáticos.",
    "Células em amarelo (aba Conferencia_DF25) são valores transcritos da fonte; as demais colunas numéricas são fórmulas.",
    "Abas: Inventario_Pasta · Fontes_Centrais · Divergencias · Quadro_Reaproveitamento · Matriz_Cobertura · Mapa_Normativo · Conferencia_DF25 · Base_Reproducao · Decisoes · Pendencias · Registro_Revisao",
]
for i, t in enumerate(txt, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(name=F, bold=(i <= 2), size=11 if i == 1 else 10)
    c.alignment = Alignment(wrap_text=True)
ws.column_dimensions["A"].width = 150

# ---------------------------------------------------------------- 1 Inventario
inv = [
 (1,"00_CONTROLE_PROJETO_Auditoria1_CPFL.md","1KgcxrqRlx6AJ9Xdu0dzhLx_mwOsmtZhY",27245,"A","Lido integralmente","Tamanho = registro etapa 5. Contém transcrição 129.669 (correto 129.659) — ver DV01."),
 (2,"CPFL_Etapa5_Estado_Continuidade.md","1ebK6uZDk6rgobxtKiUQLddxwKCr_AARc",3378,"A","Lido integralmente","Pendências P01–P11 incorporadas."),
 (3,"CPFL_Etapa5_Registro_Fontes_Continuidade.xlsx","10y3Sc-qBLR1JliGhuFcGX5vfFU4b8AkW",13777,"A","Lido (todas as abas)","IDs F01–F20 e I01–I15 preservados; URLs/datas vazias mantidas como não comprovadas."),
 (4,"5. CPFL_Etapa5_Relatorio_Metodologico.pdf","1tou1wO49tWoAKy1VGl4NTnsNQEWVMLj2",29818,"A","Não lido nesta sessão","Leitura prevista no Bloco 03 (PT 230)."),
 (5,"PROMPTS CPFL.pdf","1ktNnQ5Q60DCcfUCoMehSXG5mimi7QMJH",123025,"A","Lido integralmente (13 blocos)","Mesmo tamanho do 'CPFL_PROMPTS_GUIA_ADAPTADO_COMPLETO.pdf' do registro (123.025 bytes) — nome diferente; ver DV05."),
 (6,"4. CPFL_Etapa4_Indice_27_Papeis_Trabalho.xlsx","1uOFg2moAA2HfD809ZJKfLYcKiNX0nImR",10928,"A","Lido (todas as abas)","Controle mestre dos 27 PTs; todos 'não iniciados' como PT completo."),
 (7,"4. CPFL_Etapa4_Relatorio_Metodologico.pdf","1I8XMPA2MLzNlzVn7_5J71ozLUXvwmQ_1",72698,"A","Não lido nesta sessão","Leitura prevista no Bloco 03."),
 (8,"2. CPFL_PT320_Planilha_Materialidade_2025.xlsx","1Khcg57wpKS9j4yOQZ9otIQ-UW5gb6MNE",9167,"B","Não lido nesta sessão","Tamanho = registro etapa 5. Insumo do Bloco 04."),
 (9,"2. CPFL_PT320_Relatorio_Metodologico_Materialidade_2025.pdf","1mUwv5Udl3yfvk9nHmm_9ibwHzqOHaveP",75939,"B","Não lido nesta sessão","Insumo do Bloco 04."),
 (10,"3. CPFL_Etapa3_Matriz_Preliminar_Riscos_NBCTA315.xlsx","1ghrj6Pvng3ck82Rfk5MFkGbl9I22KI96",11269,"B","Não lido nesta sessão","Tamanho = registro etapa 5. Insumo do Bloco 04."),
 (11,"3. CPFL_Etapa3_Relatorio_Metodologico_NBCTA315.pdf","1XaYk7J1WnR4MV4Ia8IC6ykJxSb9UZGEL",74479,"B","Não lido nesta sessão","Insumo do Bloco 04."),
 (12,"CPFL_Base_Sintetica_500000.xlsx","1zfZSGJNJMHkwQvQpEB-bIEyfNyZ_Omrr",41337183,"C","Não baixável pelo conector (>10 MB)","Tamanho difere do registrado (35.830.106 bytes) — ver DV03. Base reproduzida pelo gerador (aba Base_Reproducao)."),
 (13,"CPFL_Base_Parametros.json","1k7BhaHvNrWdSoF-_f8P_yyl59jZXxnN4",947,"C","Lido; comparado","Idêntico aos parâmetros gerados na reexecução de 22/10/2026."),
 (14,"build_cpfl.py","1uH87nf9g4NVuE-ynswMy6afCG9SHu8Pw",6797,"C","Baixado e reexecutado","SHA-256 5c646bb1…b30b. Caminhos /mnt/data/ fixos (ajustar para reprodução — DV04)."),
 (15,"CPFL_LEIA_ME_Base.txt","1NlIglzobj3BRXuaRfOmf3XP4FBxxR3LA",438,"C","Não lido nesta sessão","Leitura no Bloco 07."),
 (16,"CPFL_Ata_Metodologica_Base_Sintetica.pdf","185gvmNVKsItnIaaLdpshcXzWqzuKLAiI",31228,"C","Não lido nesta sessão","Leitura no Bloco 06/07."),
 (17,"1. CPFL_Relatorio_Validacao_Integral_Definitivo0.pdf","1ZhCbkiN0PwQEldgQ9oeaTbKiT7uehdkX",32916,"C","Lido integralmente","Nome real tem sufixo '0'; registro etapa 5 cita sem o sufixo (DV05). Ressalva: 5 IDs técnicos com vínculo não pareado."),
 (18,"validacao_integral_cpfl.py","1FtbYGE3qfviGeq-tjlKLU9Y_fspQIcW4",7910,"C","Não lido nesta sessão","Leitura no Bloco 07."),
 (19,"CPFL_Gabarito_Erros.csv","13tkuD6W0ENOYF1wFSZUYrjyfzhS_HDIo",562182,"C","NÃO ABERTO (reservado)","Histórico; não usar como verdade final."),
 (20,"CPFL_Gabarito_Erros_CORRIGIDO.csv","1x0U3uyVsVLj-4kFB6qywwyGhJ-be5U44",1512694,"C","NÃO ABERTO (reservado)","Tamanho = registro etapa 5. Só será aberto após congelar os testes do PT 530."),
 (21,"EXEMPLO 1 DE DOSSIE.pdf","14-IeHTGqrrivjNc_Ykz9kktQteXCmvwK",8446021,"D","Não lido nesta sessão","Somente referência de forma (Bloco 03)."),
 (22,"EXEMPLO 2 DE DOSSIE.pdf","1V-iHB05PUPbmgyOrGGt9d29q_t-WMy5R",1263871,"D","Não lido nesta sessão","Somente referência de forma (Bloco 03)."),
 (23,"EXEMPLO 3 DE DOSSIE.pdf","1fG4UIODlvV2W9X3rfkWw4Lp2kTv4Q31s",6540508,"D","Não lido nesta sessão","Somente referência de forma (Bloco 03)."),
]
sheet("Inventario_Pasta","Inventário da pasta 'MANDAR PARA O CLAUDE PRO' (ID 1_06Wv8XzYPxV3vQ8TDlXiCpXa6rFI3WV) — 23 itens confirmados em 22/10/2026",
      ["Nº","Arquivo (nome exato no Drive)","ID Drive","Bytes (Drive)","Grupo","Situação de leitura","Observação / conferência"],
      [list(r) for r in inv],[5,48,36,13,7,26,70],
      note="Contagem: 23 arquivos diretamente na pasta (igual à informada no prompt inicial). Nenhum item ausente. 'Não lido nesta sessão' significa que a leitura está agendada para o bloco indicado — não significa arquivo ilegível.")

# ---------------------------------------------------------------- 2 Fontes centrais
fc = [
 ("FC01 (=F01)","DF25 CPFL.pdf — Demonstrações financeiras individuais e consolidadas 2025, relatório PwC","1JcqtNdI-70DCaukfmVni-rR-nQbzW25h","01_Demonstracoes_Contabeis_e_Auditoria",4076578,"0694f2b96a355d307eb446dd6201ec46ab79e64b221ae0411abc68bd488ba578","115 páginas; criado 04/03/2026; envelope DocuSign CD19B0C1-…; BP em imagem (p. física 34 = impressa 2)","Lido: texto integral + imagem da p. 34 a 300 dpi","https://publicidadelegal.valor.com.br/valor/2026/03/20/CPFLENERGIA1588267020032026.pdf (publicação legal — URL registrada no controle; não reaberta nesta sessão)"),
 ("FC02 (=F02)","DF24 CPFL.pdf — Demonstrações 2024","1bFNl9hR-0EcwDhAVuIx0ZjBSwON2cA4h","01",3529052,"pendente","—","Localizado; leitura no Bloco 02/09","pendente"),
 ("FC03 (=F03)","ITR 1TRI 2025.pdf","1ucrbjIVfI2Yw9MH817PKc9jODiXgh2rJ","01",1447067,"pendente","—","Localizado","pendente"),
 ("FC04 (=F04)","ITR 2TRI 2025.pdf","1WVl__lyr0WJYGawCUbE5iwJCeqruKRpi","01",8388608,"pendente","8.388.608 bytes = 8 MiB exatos — suspeita de truncamento mantida","Localizado; integridade a testar","pendente"),
 ("FC05 (=F05)","ITR 3TRI 2025.pdf","1ZL8tqXqSvTjh1I9BscuJr13AGZqv9F6k","01",1800414,"pendente","—","Localizado","pendente"),
 ("FC06 (=F06)","ITR 1TRI 2026.pdf","1HDduR-9R4KXLWylYt_U3Mjpfy3akms-6","01",4841995,"pendente","Período de 2026 — uso apenas em PT 560/570","Localizado","pendente"),
 ("FC07 (=F07)","ITR 2TRI 2026.pdf","1VW2PBbXed_SB5nzjefWYCzuIQwuej1Y_","01",4735816,"pendente","Período de 2026 — uso apenas em PT 560/570","Localizado","pendente"),
 ("FC08 (=F08)","REALESE DE RESULTADOS 4T25.pdf","1z2mFdomivT14QRo0Sefg0CESGC4wxbDC","01",4562291,"pendente","Nome do arquivo com grafia 'REALESE'","Localizado","pendente"),
 ("FC09 (=F18)","CONSOLIDAÇÃO DAS NORMAS AUD 1.pdf","1CbMTtLqFE4VuId0-Cl3r4wv-g-J4aXtJ","07_Normas_e_Legislacao",4775107,"f2edf5fa9503cbc90f0d900c23d7ba75198bd6a04e2dcc2b1362bc8c38f02775","756 páginas; contém EC, NBC TA 200–600 (versões vigentes para 2025); 2 erros de fluxo internos no PDF (p. 100 quase vazia)","Texto extraído; páginas iniciais e vigências indexadas (aba Mapa_Normativo)","Origem do download não registrada — conferir contra CFC (pendência P04)"),
 ("FC10","Enunciado P2 (texto colado pelo aluno no prompt inicial, 22/10/2026)","— (anexo da conversa)","—",None,"—","Arquivo 'Instruções P2.pdf' citado no guia NÃO foi localizado no Drive","Lido integralmente","Fonte: atividade da disciplina (texto fornecido pelo aluno)"),
 ("FC11","CFC — página 'NBC TA de Auditoria Independente'","—","Web",None,"—","Lista de normas e DOU","Acessada em 22/10/2026","https://cfc.org.br/tecnica/normas-brasileiras-de-contabilidade/nbc-ta-de-auditoria-independente/"),
 ("FC12","Notícia sobre NBC TA 240 (R2) e 570 (R1) — vigência","—","Web",None,"—","Publicação DOU 25/09/2026; aplicação a exercícios iniciados a partir de 01/01/2027","Acessada em 22/10/2026","https://www.lrmconsultoria.com.br/noticias/tecnicas/2026/09/30/cfc-atualiza-normas-de-auditoria-sobre-fraude-e-continuidade-operacional.html (fonte secundária; confirmar no texto oficial CFC/DOU)"),
 ("FC13","CVM Dados Abertos — diretório DFP","—","Web",None,"—","dfp_cia_aberta_2025.zip (12M) modificado em 18/10/2026","Diretório listado em 22/10/2026; arquivo NÃO baixado (bloqueio de rede do ambiente)","https://dados.cvm.gov.br/dados/CIA_ABERTA/DOC/DFP/DADOS/"),
]
sheet("Fontes_Centrais","Fontes centrais localizadas e conferidas no Bloco 01",
      ["ID","Documento","ID Drive","Pasta","Bytes","SHA-256 (calculado em 22/10/2026)","Características verificadas","Situação","URL / data de acesso"],
      [list(r) for r in fc],[11,42,30,22,11,40,44,34,55],
      note="IDs F01–F20 do registro da etapa 5 foram preservados (coluna ID mostra a correspondência). Novos IDs FC10–FC13 não colidem com os anteriores. Fontes web: data de acesso real = 22/10/2026.")

# ---------------------------------------------------------------- 3 Divergencias
dv = [
 ("DV01","00_CONTROLE seção 4 — 'Investimentos em instrumentos patrimoniais' 2025 = 129.669","DF25 CPFL.pdf, p. física 34 (impressa 2), BP consolidado: 129.659 (R$ mil)","R$ 10 mil: explica integralmente a diferença de soma do ANC (64.483.626 x 64.483.616)","Corrigir para 129.659 em todos os PTs; controle antigo preservado como histórico","Resolvida (Conferencia_DF25)"),
 ("DV02","PROMPTS CPFL.pdf — ativo total na 'página física 23' do DF25","DF25 CPFL.pdf: BP na p. física 34 (impressa 2)","Referência de página","Usar p. física 34 / impressa 2","Resolvida"),
 ("DV03","Registro etapa 5 — Base xlsx com 35.830.106 bytes e SHA f9c284c0…","Drive: 41.337.183 bytes (arquivo regravado após o hash?)","Integridade do arquivo entregue","Reexecução do gerador reproduz 35.830.104 bytes e parâmetros idênticos; propor que a base de entrega seja a reproduzida (D-B01-03)","Aberta — decisão"),
 ("DV04","build_cpfl.py grava em /mnt/data/ (caminho absoluto de outra sessão)","Exigência de reprodução com caminhos relativos (guia, Bloco 07)","Reprodutibilidade","Manter original intacto; criar versão de reprodução com caminhos relativos e mesma lógica","Aberta — Bloco 07"),
 ("DV05","Nomes no registro etapa 5: 'CPFL_Relatorio_Validacao_Integral_Definitivo.pdf' e 'CPFL_PROMPTS_GUIA_ADAPTADO_COMPLETO.pdf'","Drive: '1. CPFL_Relatorio_Validacao_Integral_Definitivo0.pdf' e 'PROMPTS CPFL.pdf' (mesmos tamanhos)","Rastreabilidade","Registrar ambos os nomes; usar nome do Drive","Resolvida"),
 ("DV06","Guia: 'Instruções P2.pdf' (4 páginas) e 'Consolidado Normas Aud1.pdf'","Drive: enunciado não localizado como PDF; normas = 'CONSOLIDAÇÃO DAS NORMAS AUD 1.pdf'","Fonte do enunciado","Usar o texto colado pelo aluno (FC10) como enunciado","Resolvida"),
 ("DV07","Guia: pasta de produção '09_Elaboracao_do_trabalho_CPFL' já pressuposta","Não existia no Drive em 22/10/2026","Organização","Criada nesta sessão com subpastas 00–07 e 99","Resolvida"),
 ("DV08","Guia: cronograma '21/10 B01–03; 20/10 B04–06; 21/10 B07–08; 22/10…'","Datas incoerentes (retroativas) e início real em 22/10","Prazo","Novo cronograma (aba Quadro_Reaproveitamento / ESTADO_DO_PROJETO)","Resolvida"),
 ("DV09","CFC lista NBC TA 240 (R2) e NBC TA 570 (R1), DOU 25/09/2026","Aplicáveis a exercícios iniciados a partir de 01/01/2027","Versão normativa","Para 2025 usar NBC TA 240 (R1) e NBC TA 570 (versão 2016), presentes no consolidado","Resolvida (Mapa_Normativo)"),
 ("DV11","00_CONTROLE seção 4 — 'IR/CS a compensar (circ.)' 2024 = 716.764","DF25 p. física 34: 715.764 (R$ mil)","R$ 1.000 mil no comparativo 2024 (circulante); R$ 1 mil remanescente = arredondamento da DF","Corrigir para 715.764 em PT 510/520","Resolvida (Conferencia_DF25)"),
 ("DV10","00_CONTROLE: ativo NC 2024 — soma das linhas","Soma = 60.577.646 vs total 60.577.648","R$ 2 mil — arredondamento da própria DF","Registrar como arredondamento de apresentação","Resolvida"),
]
sheet("Divergencias","Divergências identificadas (não corrigidas silenciosamente)",
      ["ID","Valor / trecho de origem","Fonte confrontada (documento, página)","Impacto","Decisão / tratamento","Estado"],
      [list(r) for r in dv],[7,48,48,30,48,20])

# ---------------------------------------------------------------- 4 Reaproveitamento
qr = [
 ("B01","Início, requisitos, mapa normativo","PT 230 (prep.)","Controles etapa 5; índice etapa 4","EXECUTAR + REAPROVEITAR","Este arquivo; inventário, cobertura, mapa normativo e conferência produzidos"),
 ("B02","Entidade e saldos","Suporte a todos","00_CONTROLE seção 3–4","VERIFICAR + COMPLETAR","Saldos do controle já checados no ativo; falta passivo, DRE, DFC, notas e CVM"),
 ("B03","Modelo e condições","EC, 200, 210, 220, 230","Índice etapa 4; exemplos (forma)","EXECUTAR","Nenhum PT redigido ainda"),
 ("B04","Materialidade, riscos, programas","240, 250, 300, 315, 320, 330 + 6 programas","Planilha PT 320; matriz 17 riscos","REAPROVEITAR (320, 315) + VERIFICAR + EXECUTAR","Não refazer cálculo; conferir bases com DF25 e redigir PTs"),
 ("B05","Grupo e prestadores","402, 600","DFs controladas pasta 06","EXECUTAR","NBC TA 600 (R2) aplicável (vigência 2024)"),
 ("B06","Projeto da base","Apoio 320/500/530","Conta aprovada 21/10; ata; parâmetros","REAPROVEITAR + VERIFICAR","Conta não reescolhida; documentar justificativa e dicionário"),
 ("B07","Geração e conciliação","500 (parcial)","build_cpfl.py; validação integral","REAPROVEITAR + COMPLETAR","Base reproduzida em 22/10; falta versão com caminhos relativos e CSV comprimido"),
 ("B08","Testes e amostragem","530, 450 (simulado)","Base; gabarito reservado","EXECUTAR","Gabarito só após congelar testes"),
 ("B09","Abertura, analíticos, estimativas, PR","510, 520, 540, 550","DF24/DF25","EXECUTAR","—"),
 ("B10","Itens específicos, confirmações, eventos, continuidade","501, 505, 560, 570","ITRs 2026; escrituras","EXECUTAR","NBC TA 570 versão 2016 (não a R1 de 2026)"),
 ("B11","Comunicações e conclusão","260, 265, 580; fechamento 450/500","Todos","EXECUTAR","—"),
 ("B12","Revisão técnica","Transversal","Todos","EXECUTAR","Revisão automatizada identificada como tal"),
 ("B13","Word/PDF/GitHub","Entrega","—","EXECUTAR","Base > 25 MB: verificar limite atual do GitHub"),
 ("Etapa 1","Base sintética + validação","530/500/450","xlsx, json, py, relatório","REAPROVEITAR (com ressalva DV03/DV04 e 5 anomalias)","Não regerar lógica; só reprodução"),
 ("Etapa 2","Materialidade","320","xlsx + pdf","REAPROVEITAR + VERIFICAR","Bases conferidas: LAIR 8.091.246 a confirmar na DRE (B02)"),
 ("Etapa 3","Matriz preliminar","315","xlsx + pdf","REAPROVEITAR + COMPLETAR","Converter em PT 315 com fontes/páginas"),
 ("Etapa 4","Índice 27 PTs","Todos","xlsx + pdf","REAPROVEITAR (controle mestre)","Status atualizado só com entregáveis reais"),
 ("Etapa 5","Registro de fontes","Todos","xlsx + md + pdf","REAPROVEITAR + COMPLETAR","URLs/datas preenchidas apenas quando acessadas"),
 ("Exemplos 1–3","Referência visual","—","pdf","NÃO APLICÁVEL como fonte de fatos","Apenas forma (decisão do aluno e do guia)"),
]
sheet("Quadro_Reaproveitamento","Quadro reaproveitar / verificar / completar / executar / não aplicável — por bloco e etapa",
      ["Bloco/etapa","Tema","PTs","Insumo existente","Ação","Justificativa / observação"],
      [list(r) for r in qr],[11,34,34,34,36,60])

# ---------------------------------------------------------------- 5 Matriz de cobertura
ob = [
 ("E-01","Dossiê da companhia atribuída (CPFL, CPFE3) cobrindo EC e NBC TA 200 a 600","Dossiê DOCX+PDF","B03–B13","Todos","27 PTs presentes com 9 campos","PENDENTE"),
 ("E-02","Usar DFs 2025, comparativos, notas e relatório do auditor","Fontes FC01/FC02","B02","Todos","Cada saldo com doc/página","EM ANDAMENTO (ativo conferido)"),
 ("E-03","Aplicar julgamento, ceticismo, risco, planejamento, evidência","PTs 200/300/315/330","B03–B04","200,300,315,330","Julgamentos fundamentados","PENDENTE"),
 ("E-04","Fontes CVM/B3/RI com URL, data, documento, página, exercício","Fontes_Usadas","Todos","230","Nenhuma fonte sem data real","EM ANDAMENTO"),
 ("E-05","Identificar alterações de nome empresarial e código de negociação","Memorando entidade","B02","EC/315","Fonte CVM/B3 verificada","PENDENTE (Draftll Participações — a reconfirmar)"),
 ("E-06","Não contatar empresa/terceiros","Declaração no dossiê","Todos","200/505","Nenhum envio","CUMPRIDO até aqui"),
 ("E-07","Classificar cada procedimento (Pública/Planejado/Simulado)","Coluna de modalidade","B03–B11","Todos","100% dos procedimentos classificados","PENDENTE"),
 ("E-08","Não apresentar confirmações/entrevistas/inspeções como realizadas","Revisão B12","B12","Todos","Varredura sem achados","PENDENTE"),
 ("E-09","Todos os PTs do roteiro; 'não se aplica' só com análise","27 PTs","B03–B11","Todos","Justificativa específica","PENDENTE"),
 ("E-10","Base simulada ≥ 500.000 linhas da conta de maior materialidade do ativo","Base reproduzida","B06–B07","500/530","500.000 linhas lidas do arquivo","EM ANDAMENTO (reproduzida; leitura de contagem no B07)"),
 ("E-11","Registros individuais compatíveis com a conta; sem repetição","Especificação","B06","530","Unidade de registro = bem da base de remuneração","INSUMO"),
 ("E-12","Justificar escolha da conta e materialidade","Memorando de conta","B06","320/500","AFC = 36,5% do ativo (29.623.619/81.099.930)","EM ANDAMENTO (saldo conferido na DF25 p. 34)"),
 ("E-13","Premissas e método de geração","Ata + memorando","B06","500","Documento revisado","INSUMO"),
 ("E-14","Dicionário de variáveis","Dicionário","B07","500","12 campos","INSUMO (aba Dicionario do gerador)"),
 ("E-15","Conciliar total da base com saldo contábil","Conciliação","B07","500","Diferença R$ 0,00 explicada","EM ANDAMENTO (soma reproduzida = alvo)"),
 ("E-16","Usar a base: amostras, exceções, avaliação","PT 530/450","B08","530/450","Código executado + resultados","PENDENTE"),
 ("E-17","GitHub: código, instruções, PTs da base","Repositório","B13","—","Link acessível","PENDENTE"),
 ("E-18","Erros inseridos documentados + registro comparável","Gabarito + comparação","B08","450/530","Comparação após congelamento","INSUMO (gabarito reservado)"),
 ("E-19","Dados expressamente sintéticos","Rótulos","B07","500","Campo origem_dados = SINTETICO","INSUMO"),
 ("E-20","9 campos mínimos em cada PT","Modelo de PT","B03","Todos","Checklist por PT","PENDENTE"),
 ("E-21","Campo de revisão sem assinatura simulada","Modelo de PT","B03","Todos","Campo em branco","APLICADO neste arquivo"),
 ("E-22","Dossiê em PDF com referências cruzadas","PDF","B13","—","PDF = DOCX final","PENDENTE"),
 ("E-23","Base ≥ 500k + dicionário no repositório (ou link)","Base","B13","—","Link testado","PENDENTE"),
 ("E-24","Código de geração e análise","Scripts","B07–B08","—","Reexecução limpa","EM ANDAMENTO (gerador reexecutado)"),
 ("E-25","Planilhas editáveis (materialidade, amostras, análises, conciliações)","XLSX","B04–B09","320/530/520","Fórmulas recalculadas","INSUMO (materialidade)"),
 ("E-26","Matriz de cobertura das normas","Esta aba + Mapa_Normativo","B01/B12","230","Norma → PT","EM ANDAMENTO"),
 ("E-27","Relação das fontes e links","Fontes_Usadas","Todos","230","URL + data","EM ANDAMENTO"),
 ("E-28","README com matrícula, organização e reprodução","README.md","B13","—","Passos testados","PENDENTE"),
 ("E-29","Link de download se base exceder limite do GitHub","README","B13","—","Link verificado","PENDENTE"),
 ("E-30","Programas: receitas, disponibilidades, estoques, imobilizado, passivos, estimativas","6 programas","B04","300/330","Aplicabilidade fundamentada","INSUMO (índice etapa 4)"),
 ("E-31","Conclusão integrada sem opinião profissional","Seção final","B11","—","Sem linguagem de opinião","PENDENTE"),
 ("E-32","Versões das normas aplicáveis a 2025 registradas","Mapa_Normativo","B01","Todos","Versão + vigência + página","EM ANDAMENTO (27 versões identificadas; itens por PT)"),
 ("E-33","Postar link até 25/10/2026 e verificar acesso do professor","Ação do aluno","B13","—","Teste em janela anônima","PENDENTE (ação do aluno)"),
]
sheet("Matriz_Cobertura","Matriz de cobertura das obrigações do enunciado (fonte: FC10)",
      ["ID","Obrigação (enunciado)","Saída","Bloco","PT","Critério de aceite","Estado em 22/10/2026"],
      [list(r) for r in ob],[7,60,28,10,18,36,36])

# ---------------------------------------------------------------- 6 Mapa normativo
mn = [
 ("EC","Estrutura Conceitual para Trabalhos de Asseguração","NBC TA Estrutura Conceitual","DOU 25/11/2015","Efeitos a partir de 01/01/2016 (texto no consolidado, p. 24)","4–24 (aprox.)","EC","Versão confirmada; itens a citar no B03"),
 ("200","Objetivos gerais do auditor","NBC TA 200 (R1)","DOU 05/09/2016","Períodos findos em/após 31/12/2016 (alteração 2016)","29","200","Versão confirmada"),
 ("210","Termos do trabalho","NBC TA 210 (R1)","DOU 05/09/2016","idem","57","210","Versão confirmada"),
 ("220","Gestão de qualidade","NBC TA 220 (R3)","DOU 09/12/2021","Item 10: períodos iniciados em/após 01/01/2023 (p. 93) — APLICÁVEL a 2025","89","220","Vigência lida no texto; p. 100 do PDF quase vazia — conferir"),
 ("230","Documentação","NBC TA 230 (R1)","DOU 05/09/2016","—","131","230","Versão confirmada"),
 ("240","Fraude","NBC TA 240 (R1) — NÃO a (R2)","DOU 05/09/2016","(R2) de DOU 25/09/2026 aplica-se a exercícios iniciados a partir de 01/01/2027 (FC12) → não aplicável a 2025","146","240","Atenção DV09"),
 ("250","Leis e regulamentos","NBC TA 250","DOU 14/02/2019","—","189","250","Versão confirmada"),
 ("260","Comunicação com a governança","NBC TA 260 (R2)","DOU 04/07/2016","—","203","260","Versão confirmada"),
 ("265","Deficiências de controle","NBC TA 265","Res. CFC 1.210/2009","—","225","265","Versão confirmada"),
 ("300","Planejamento","NBC TA 300 (R1)","DOU 05/09/2016","—","237","300","Versão confirmada"),
 ("315","Riscos","NBC TA 315 (R2)","DOU 02/09/2021","Item 10: períodos iniciados em/após 01/01/2022 (p. 256) — APLICÁVEL","252","315","Vigência lida no texto"),
 ("320","Materialidade","NBC TA 320 (R1)","DOU 05/09/2016","—","367","320","Versão confirmada"),
 ("330","Respostas aos riscos","NBC TA 330 (R1)","DOU 05/09/2016","—","379","330","Versão confirmada"),
 ("402","Organização prestadora de serviços","NBC TA 402","Res. CFC 1.215/2009","—","404","402","Versão confirmada"),
 ("450","Distorções","NBC TA 450 (R1)","DOU 05/09/2016","—","427","450","Versão confirmada"),
 ("500","Evidência","NBC TA 500 (R1)","DOU 05/09/2016","—","446","500","Versão confirmada"),
 ("501","Itens específicos","NBC TA 501","Res. CFC 1.218/2009","Exercícios iniciados em/após 01/01/2010 (p. 471)","471","501","Versão confirmada"),
 ("505","Confirmações externas","NBC TA 505","Res. CFC 1.219/2009","—","456","505","Versão confirmada"),
 ("510","Saldos iniciais","NBC TA 510 (R1)","DOU 05/09/2016","—","497","510","Versão confirmada"),
 ("520","Procedimentos analíticos","NBC TA 520","Res. CFC 1.221/2009","—","521","520","Versão confirmada"),
 ("530","Amostragem","NBC TA 530","Res. CFC 1.222/2009","—","530","530","Versão confirmada"),
 ("540","Estimativas","NBC TA 540 (R2)","DOU 23/10/2019","Item 10: períodos iniciados em/após 01/01/2020 (p. 552) — APLICÁVEL","548","540","Vigência lida no texto"),
 ("550","Partes relacionadas","NBC TA 550","Res. CFC 1.224/2009","—","600","550","Versão confirmada"),
 ("560","Eventos subsequentes","NBC TA 560 (R1)","DOU 05/09/2016","—","625","560","Versão confirmada"),
 ("570","Continuidade operacional","NBC TA 570 (versão 2016) — NÃO a (R1) de 2026","DOU 04/07/2016","Períodos findos em/após 31/12/2016 (p. 646). (R1) DOU 25/09/2026 só para exercícios iniciados a partir de 01/01/2027","639","570","Atenção DV09"),
 ("580","Representações formais","NBC TA 580 (R1)","DOU 05/09/2016","—","660","580","Versão confirmada"),
 ("600","Auditoria de grupos","NBC TA 600 (R2)","DOU 06/12/2023 (norma de 25/10/2023)","Item 12: períodos iniciados em/após 01/01/2024 (p. 688) — APLICÁVEL a 2025","684","600","Vigência lida no texto"),
]
sheet("Mapa_Normativo","Mapa normativo — versões aplicáveis ao exercício de 2025 (fontes FC09 e FC11; FC12 para 240/570)",
      ["Código","Tema","Versão aplicável a 2025","Publicação (CFC FC11)","Vigência / evidência (página do consolidado FC09)","Pág. inicial no consolidado","PT","Estado da verificação"],
      [list(r) for r in mn],[8,30,34,24,62,13,8,36],
      note="'Versão confirmada' = sigla e revisão coincidem entre a lista oficial do CFC (acesso 22/10/2026) e o consolidado do Drive. Os ITENS citados em cada PT serão conferidos no texto ao redigir o PT (estado por PT na Matriz_Cobertura). Comentários didáticos e modelos não são requisitos.")

# ---------------------------------------------------------------- 7 Conferencia DF25
ws, r0 = sheet("Conferencia_DF25","Conferência inicial — Balanço patrimonial consolidado, ativo (R$ mil) — DF25 CPFL.pdf, p. física 34 / impressa 2",
      ["Linha do BP (consolidado)","Nota","31/12/2025","31/12/2024","Fonte"],
      [],[52,8,16,16,40],
      note="Modalidade: Realizado com evidência pública (leitura visual da imagem a 300 dpi). Valores em amarelo = transcrição; totais calculados = fórmulas. Conferência CVM (BPA_con) PENDENTE: dfp_cia_aberta_2025.zip não baixável no ambiente (FC13).")
circ = [("Caixa e equivalentes de caixa",5,2229320,1973401),("Títulos e valores mobiliários",6,675762,1573654),("Consumidores, concessionárias e permissionárias",7,6244460,5883894),("Estoques",None,236566,221573),("Dividendo e juros sobre o capital próprio",None,14712,19115),("IR e CS a compensar",8,528584,715764),("Outros tributos a compensar",8,693224,653181),("PIS/COFINS a compensar sobre ICMS",8,1844885,1919127),("Derivativos",34,4513,915621),("Ativo financeiro setorial",9,1442244,393443),("Ativo contratual",15,872283,774368),("Outros ativos",12,1829761,1508406)]
nc = [("Consumidores, concessionárias e permissionárias",7,133102,186371),("Depósitos judiciais",22,805776,749936),("IR e CS a compensar",8,447583,503271),("Outros tributos a compensar",8,440936,479460),("PIS/COFINS a compensar sobre ICMS",8,2251657,3652513),("Ativo financeiro setorial",9,960063,337522),("Derivativos",34,591350,193542),("Créditos fiscais diferidos",10,202046,266798),("Ativo financeiro da concessão",11,29623619,25209768),("Investimentos em instrumentos patrimoniais",None,129659,144961),("Outros ativos",12,261064,331087),("Investimentos",13,380803,581364),("Imobilizado",14,9375879,9752282),("Intangível",16,6892839,7824433),("Ativo contratual",15,11987240,10364338)]
r = r0 + 1
def put(row, vals, bold=False, fill=None):
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = Font(name=F, size=9, bold=bold); c.border = BRD
        if j in (3, 4): c.number_format = '#,##0;(#,##0);"-"'
        if fill and j in (3, 4) and not (isinstance(v, str) and v.startswith("=")): c.fill = fill
ws.cell(row=r, column=1, value="CIRCULANTE").font = Font(name=F, bold=True, size=9); r += 1
c0 = r
for n, nt, a, b in circ:
    put(r, [n, nt, a, b, "DF25 p.34"], fill=INPUT_FILL); r += 1
c1 = r - 1
put(r, ["Total do circulante (calculado)", None, f"=SUM(C{c0}:C{c1})", f"=SUM(D{c0}:D{c1})", "fórmula"], bold=True); rc = r; r += 1
put(r, ["Total do circulante (divulgado)", None, 16616314, 16551546, "DF25 p.34"], bold=True, fill=INPUT_FILL); rcd = r; r += 1
put(r, ["Diferença", None, f"=C{rc}-C{rcd}", f"=D{rc}-D{rcd}", "2024: +1 = arredondamento da DF (DV11)"]); r += 2
ws.cell(row=r, column=1, value="NÃO CIRCULANTE").font = Font(name=F, bold=True, size=9); r += 1
n0 = r
for n, nt, a, b in nc:
    put(r, [n, nt, a, b, "DF25 p.34"], fill=INPUT_FILL); r += 1
n1 = r - 1
put(r, ["Total do não circulante (calculado)", None, f"=SUM(C{n0}:C{n1})", f"=SUM(D{n0}:D{n1})", "fórmula"], bold=True); rn = r; r += 1
put(r, ["Total do não circulante (divulgado)", None, 64483616, 60577648, "DF25 p.34"], bold=True, fill=INPUT_FILL); rnd = r; r += 1
put(r, ["Diferença", None, f"=C{rn}-C{rnd}", f"=D{rn}-D{rnd}", "2024: −2 = arredondamento da DF (DV10)"]); r += 2
put(r, ["TOTAL DO ATIVO (calculado)", None, f"=C{rc}+C{rn}", f"=D{rc}+D{rn}", "fórmula"], bold=True); rt = r; r += 1
put(r, ["TOTAL DO ATIVO (divulgado)", None, 81099930, 77129194, "DF25 p.34"], bold=True, fill=INPUT_FILL); rtd = r; r += 1
put(r, ["Diferença", None, f"=C{rt}-C{rtd}", f"=D{rt}-D{rtd}", ""]); r += 1
put(r, ["Valor de referência do controle do projeto", None, 81099930, None, "00_CONTROLE seção 4"]); rref = r; r += 1
put(r, ["Diferença vs. referência", None, f"=C{rtd}-C{rref}", None, "esperado 0"]); r += 2
afc = n0 + 8
put(r, ["Ativo financeiro da concessão / ativo total", None, f"=C{afc}/C{rtd}", None, "base da escolha da conta (E-12)"]); ws.cell(row=r, column=3).number_format = "0.0%"; r += 1
put(r, ["Ativo financeiro da concessão / ativo não circulante", None, f"=C{afc}/C{rnd}", None, ""]); ws.cell(row=r, column=3).number_format = "0.0%"; r += 2
notes = ["Divergência DV01: controle anterior registrava 129.669 em 'Investimentos em instrumentos patrimoniais' 2025; a imagem mostra 129.659. Com a correção, o ANC fecha sem diferença.",
         "Escala: 'Em milhares de reais' (cabeçalho do BP). Data-base 31/12/2025; comparativo 31/12/2024 apresentado no mesmo documento (não reapresentado).",
         "Conferência PDF × CVM (BPA_con, ESCALA_MOEDA=MIL, DT_FIM_EXERC=2025-12-31, ORDEM_EXERC=ÚLTIMO, CNPJ 02.429.144/0001-93): NÃO EXECUTADA — ver pendência PB01-01."]
for t in notes:
    c = ws.cell(row=r, column=1, value=t); c.font = Font(name=F, size=9, italic=True); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    c.alignment = Alignment(wrap_text=True); ws.row_dimensions[r].height = 28; r += 1

# ---------------------------------------------------------------- 8 Base reproducao
br = [
 ("Gerador","build_cpfl.py (Drive ID 1uH87nf9g4NVuE-ynswMy6afCG9SHu8Pw)","SHA-256 5c646bb1822149d8083df3c4211e43757c16468bd8603ecdd7c0b717aad0b30b; 6.797 bytes"),
 ("Alteração para execução","Somente 3 caminhos de saída (/mnt/data → saida/ e pasta reservada)","Lógica, semente e parâmetros inalterados"),
 ("Ambiente","Python 3.13.16; openpyxl (versão do ambiente); Linux","22/10/2026; duração 41,9 s"),
 ("Parâmetros gerados x JSON do Drive","raw_total_cents 4.139.537.816.245; precalibration 2.962.237.074.501; ajuste 124.825.499; final 2.962.361.900.000","IDÊNTICOS (comparação automática)"),
 ("Total contabilizado reproduzido","R$ 29.623.619.000,00","Igual ao saldo-alvo (DF25 p.34 × 1.000); diferença R$ 0,00"),
 ("Arquivo xlsx reproduzido","35.830.104 bytes","Registro etapa 5: 35.830.106 bytes; Drive atual: 41.337.183 bytes (DV03)"),
 ("Gabarito reproduzido","Gravado em pasta RESERVADO_NAO_ABRIR; conteúdo NÃO lido","SHA-256 1fd1bff4…c62b; 559.181 bytes (original no Drive: 562.182 bytes)"),
 ("Conclusão","Base reproduzível a partir do código e da semente","Leitura linha a linha do xlsx reproduzido (contagem, chaves) fica para o Bloco 07"),
]
sheet("Base_Reproducao","Reprodução da base sintética (SIMULADO para fins didáticos) — 22/10/2026",
      ["Item","Resultado","Observação"],[list(x) for x in br],[34,70,70])

# ---------------------------------------------------------------- 9 Decisoes
de = [
 ("D-B01-01","Perímetro: consolidado societário 2025 (com individual e componentes quando necessário)","Prompt inicial e controle (decisão já registrada)","Mantida","Aluno (21–22/10/2026)"),
 ("D-B01-02","Versões normativas: 240 (R1) e 570 (2016) para 2025, não as revisões de 2026","Vigência a partir de 01/01/2027 (FC12)","PROPOSTA — validar","Pendente validação do aluno"),
 ("D-B01-03","Base de entrega = base reproduzida pelo gerador original (mesma semente), com hash novo; arquivo de 41 MB do Drive preservado como histórico","DV03: arquivo do Drive não baixável e com tamanho diferente do registrado","PROPOSTA — validar","Pendente validação do aluno"),
 ("D-B01-04","Enunciado = texto colado pelo aluno (FC10)","PDF não localizado","Aplicada","Assistente (rotina)"),
 ("D-B01-05","Conta da base = ativo financeiro da concessão consolidado, R$ 29.623.619 mil","Aprovada em 21/10; saldo reconfirmado na DF25 p.34","Mantida","Aluno (21/10/2026)"),
]
sheet("Decisoes","Registro de decisões",["ID","Decisão","Fundamento","Estado","Responsável / data"],[list(x) for x in de],[11,70,55,22,30])

# ---------------------------------------------------------------- 10 Pendencias
pe = [
 ("PB01-01","Alta","Conferência PDF × CVM (BPA_con 2025): baixar dfp_cia_aberta_2025.zip","Ambiente bloqueia dados.cvm.gov.br; conector Drive limita 10 MB","Alternativas: (a) computador do aluno conectado; (b) aluno baixa o zip e salva no Drive (12 MB — ainda acima de 10 MB); (c) extrair apenas BPA_con via computador","B02"),
 ("PB01-02","Alta","DV03 — validar decisão D-B01-03 sobre a base","—","Aluno valida","B06/B07"),
 ("PB01-03","Média","Ler PDFs metodológicos das etapas 2–5 e exemplos (forma)","Agendado","—","B03/B04"),
 ("PB01-04","Média","Integridade ITR 2TRI 2025 (8 MiB exatos) e ratings .crdownload","Não testado","Testar ao usar","B10"),
 ("PB01-05","Média","Confirmar nome anterior (Draftll Participações S.A.) e eventual troca de ticker em fonte CVM/B3","Não reaberto","—","B02"),
 ("PB01-06","Média","Conferir no texto oficial CFC/DOU a vigência das NBC TA 240 (R2) e 570 (R1) (fonte usada é secundária)","Fonte secundária","—","B04/B10"),
 ("PB01-07","Baixa","Página 100 do consolidado quase vazia (NBC TA 220) — conferir se há texto ausente","Erro de fluxo no PDF","Comparar com texto CFC ao citar itens","B03"),
 ("P01–P11","—","Pendências da etapa 5 (preservadas)","—","P01 parcialmente resolvida (ativo conferido; DV01)","—"),
]
sheet("Pendencias","Pendências abertas após o Bloco 01",["ID","Prioridade","Pendência","Causa","Providência","Bloco"],[list(x) for x in pe],[10,10,60,40,55,8])

# ---------------------------------------------------------------- 11 Revisao
rv = [("RV-B01-01","Revisão automatizada (assistente de IA)","22/10/2026","Somas do BP recalculadas por fórmula; parâmetros da base comparados automaticamente; vigências lidas no texto","Sem achados além de DV01–DV11","Não substitui revisão humana"),
      ("RV-B01-02","Revisão do aluno","____","____","____","Campo em branco — preencher apenas quando efetivamente revisado")]
sheet("Registro_Revisao","Registro de revisão (sem revisão de terceiros simulada)",["ID","Tipo","Data","Escopo","Resultado","Observação"],[list(x) for x in rv],[11,34,12,60,34,40])

os.makedirs(OUT, exist_ok=True)
path = os.path.join(OUT, "CPFL_B01_Controles_Execucao.xlsx")
wb.save(path)
print(path)
