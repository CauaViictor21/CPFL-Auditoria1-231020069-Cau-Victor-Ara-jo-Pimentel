"""Converte um Markdown em DOCX (Arial, tabelas com borda) e PDF. Uso: python md2docpdf.py arquivo.md"""
import sys, os, subprocess
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
HERE = os.path.dirname(os.path.abspath(__file__))
md = sys.argv[1]; docx = os.path.splitext(md)[0] + ".docx"
subprocess.run(["pandoc", md, "-o", docx], check=True)
d = Document(docx)
for st in d.styles:
    try:
        if st.type == 1 and st.font is not None:
            st.font.name = "Arial"; rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
            if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), "Arial")
    except Exception: pass
for t in d.tables:
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x = OxmlElement(f"w:{e}"); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "4"); x.set(qn("w:color"), "A6A6A6"); b.append(x)
    t._tbl.tblPr.append(b)
    for row in t.rows:
        for c in row.cells:
            for par in c.paragraphs:
                for r in par.runs: r.font.size = Pt(8)
for s in d.sections:
    s.left_margin = s.right_margin = Cm(1.8); s.top_margin = s.bottom_margin = Cm(2)
d.save(docx)
subprocess.run([sys.executable, os.path.join(HERE, "lo_update_export.py"), docx], check=True)
