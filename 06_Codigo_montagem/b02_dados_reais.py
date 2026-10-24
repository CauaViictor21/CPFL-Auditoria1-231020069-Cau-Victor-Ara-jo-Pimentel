# -*- coding: utf-8 -*-
"""Bloco 02 - Dados reais CPFL Energia 2025/2024 (R$ mil) com fonte e página.
Execução: python b02_dados_reais.py <pasta_saida>
Transcrição visual de DF25 CPFL.pdf (SHA-256 0694f2b9...a578), feita em 22/10/2026.
Valores transcritos = células amarelas; totais e conciliações = fórmulas."""
import sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else "."
F = "Arial"; HOJE = "22/10/2026"
HEAD = PatternFill("solid", fgColor="1F3864"); INP = PatternFill("solid", fgColor="FFF2CC")
CHK = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="BFBFBF"); BRD = Border(left=thin, right=thin, top=thin, bottom=thin)
NUM = '#,##0;(#,##0);"-"'
DOC = "FC01 DF25 CPFL.pdf"

wb = Workbook(); wb.remove(wb.active)

def header(ws, title, cols, widths, note):
    ws["A1"] = title; ws["A1"].font = Font(name=F, bold=True, size=12, color="1F3864")
    ws["A2"] = f"Preparador: matrícula 231020069 (assistido por IA) | Data: {HOJE} | Modalidade: Realizado com evidência pública | Revisão: ____"
    ws["A2"].font = Font(name=F, italic=True, size=8, color="595959")
    ws["A3"] = note; ws["A3"].font = Font(name=F, size=9, color="7F6000"); ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(cols)); ws.row_dimensions[3].height = 40
    for j, h in enumerate(cols, 1):
        c = ws.cell(row=4, column=j, value=h); c.font = Font(name=F, bold=True, color="FFFFFF", size=9)
        c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BRD
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A5"

def cell(ws, r, c, v, bold=False, fill=None, num=False, italic=False):
    x = ws.cell(row=r, column=c, value=v); x.font = Font(name=F, size=9, bold=bold, italic=italic); x.border = BRD
    if num: x.number_format = NUM
    if fill: x.fill = fill
    return x

COLS = ["ID", "Linha (consolidado)", "Nota", "31/12/2025 | 2025", "31/12/2024 | 2024", "Documento", "Pág. física", "Pág. impressa", "Escala / perímetro"]
W = [9, 52, 6, 16, 16, 20, 9, 9, 22]
REG = {}  # nome lógico -> (aba, linha)

def block(ws, start, prefix, lines, page, sec_name):
    r = start
    cell(ws, r, 2, sec_name, bold=True); r += 1
    first = r
    for i, (name, nt, a, b) in enumerate(lines, 1):
        idn = f"{prefix}{i:02d}"
        cell(ws, r, 1, idn); cell(ws, r, 2, name); cell(ws, r, 3, nt)
        cell(ws, r, 4, a, fill=INP, num=True); cell(ws, r, 5, b, fill=INP, num=True)
        cell(ws, r, 6, DOC); cell(ws, r, 7, page); cell(ws, r, 8, page - 32); cell(ws, r, 9, "R$ mil | consolidado")
        REG[name + "|" + prefix] = (ws.title, r)
        r += 1
    return first, r - 1, r

def tot(ws, r, label, f25, f24, divulg=None, page=None, key=None):
    cell(ws, r, 2, label + " (calculado)", bold=True)
    cell(ws, r, 4, f25, bold=True, num=True); cell(ws, r, 5, f24, bold=True, num=True); cell(ws, r, 6, "fórmula")
    if key: REG[key] = (ws.title, r)
    r += 1
    if divulg:
        cell(ws, r, 2, label + " (divulgado)", bold=True)
        cell(ws, r, 4, divulg[0], bold=True, fill=INP, num=True); cell(ws, r, 5, divulg[1], bold=True, fill=INP, num=True)
        cell(ws, r, 6, DOC); cell(ws, r, 7, page); cell(ws, r, 8, page - 32)
        if key: REG[key + "_div"] = (ws.title, r)
        r += 1
        cell(ws, r, 2, "Diferença (calculado − divulgado)", italic=True)
        cell(ws, r, 4, f"=D{r-2}-D{r-1}", fill=CHK, num=True); cell(ws, r, 5, f"=E{r-2}-E{r-1}", fill=CHK, num=True)
        cell(ws, r, 6, "esperado 0 ou ±arredondamento (nota 2.4)")
        r += 1
    return r + 1

# ------------------------------------------------ BP ATIVO
ws = wb.create_sheet("BP_Ativo")
header(ws, "Balanço patrimonial consolidado — ATIVO (R$ mil)", COLS, W,
       "Fonte: DF25 CPFL.pdf, p. física 34 / impressa 2 (imagem lida a 300 dpi em 22/10/2026). Comparativo 2024 do mesmo documento. Nota 2.4 (p. física 46): arredondamento após totalização — somas podem diferir em unidades.")
circ = [("Caixa e equivalentes de caixa",5,2229320,1973401),("Títulos e valores mobiliários",6,675762,1573654),("Consumidores, concessionárias e permissionárias",7,6244460,5883894),("Estoques",None,236566,221573),("Dividendo e juros sobre o capital próprio",None,14712,19115),("Imposto de renda e contribuição social a compensar",8,528584,715764),("Outros tributos a compensar",8,693224,653181),("PIS/COFINS a compensar sobre ICMS",8,1844885,1919127),("Derivativos",34,4513,915621),("Ativo financeiro setorial",9,1442244,393443),("Ativo contratual",15,872283,774368),("Outros ativos",12,1829761,1508406)]
nc = [("Consumidores, concessionárias e permissionárias",7,133102,186371),("Depósitos judiciais",22,805776,749936),("Imposto de renda e contribuição social a compensar",8,447583,503271),("Outros tributos a compensar",8,440936,479460),("PIS/COFINS a compensar sobre ICMS",8,2251657,3652513),("Ativo financeiro setorial",9,960063,337522),("Derivativos",34,591350,193542),("Créditos fiscais diferidos",10,202046,266798),("Ativo financeiro da concessão",11,29623619,25209768),("Investimentos em instrumentos patrimoniais",None,129659,144961),("Outros ativos",12,261064,331087),("Investimentos",13,380803,581364),("Imobilizado",14,9375879,9752282),("Intangível",16,6892839,7824433),("Ativo contratual",15,11987240,10364338)]
a, b, r = block(ws, 5, "AC", circ, 34, "Circulante")
r = tot(ws, r, "Total do circulante", f"=SUM(D{a}:D{b})", f"=SUM(E{a}:E{b})", (16616314, 16551546), 34, "AC")
a2, b2, r = block(ws, r, "ANC", nc, 34, "Não circulante")
r = tot(ws, r, "Total do não circulante", f"=SUM(D{a2}:D{b2})", f"=SUM(E{a2}:E{b2})", (64483616, 60577648), 34, "ANC")
ac, anc = REG["AC_div"][1], REG["ANC_div"][1]
r = tot(ws, r, "TOTAL DO ATIVO", f"=D{ac}+D{anc}", f"=E{ac}+E{anc}", (81099930, 77129194), 34, "AT")
AFC_ROW = a2 + 8

# ------------------------------------------------ BP PASSIVO
wp = wb.create_sheet("BP_Passivo_PL")
header(wp, "Balanço patrimonial consolidado — PASSIVO E PATRIMÔNIO LÍQUIDO (R$ mil)", COLS, W,
       "Fonte: DF25 CPFL.pdf, p. física 35 / impressa 3 (imagem lida a 220 dpi em 22/10/2026).")
pc = [("Fornecedores",17,4098399,3890827),("Empréstimos e financiamentos",18,3160115,4587739),("Debêntures",19,1735590,815233),("Entidade de previdência privada",20,105186,336398),("Imposto de renda e contribuição social a recolher",21,748215,440904),("Outros impostos, taxas e contribuições a recolher",21,703760,784498),("Mútuo com partes relacionadas",31,19416,0),("Dividendo",None,1506041,1510207),("Derivativos",34,232540,8178),("Passivo financeiro setorial",9,1077501,927285),("Provisões para desmobilização e gastos ambientais",None,5153,12699),("PIS/COFINS devolução consumidores",8,1018818,1072326),("Outras contas a pagar",23,3465230,2984242)]
pnc = [("Fornecedores",17,239699,254364),("Empréstimos e financiamentos",18,6417553,7690254),("Debêntures",19,17742469,15695112),("Entidade de previdência privada",20,501738,451514),("Imposto de renda e contribuição social a recolher",21,254049,245142),("Outros impostos, taxas e contribuições a recolher",21,956254,997778),("Débitos fiscais diferidos",10,3001777,2665302),("Provisões para riscos fiscais, cíveis e trabalhistas",22,1740903,1561980),("Mútuo com partes relacionadas",31,4393234,2485694),("Derivativos",34,170854,272377),("Passivo financeiro setorial",9,913146,798912),("Provisões para desmobilização e gastos ambientais",None,170123,151778),("PIS/COFINS devolução consumidores",8,2384026,3864430),("Outras contas a pagar",23,833632,828992)]
pl = [("Capital social",24,9388071,9388071),("Reserva de capital",24,-1394956,-1394956),("Reserva legal",24,1877614,1877614),("Reserva estatutária - reforço de capital de giro",24,7510454,7033014),("Reserva de lucro a realizar",24,3822612,3075193),("Dividendo",24,2928598,1855190),("Resultado abrangente acumulado",24,-1642619,-1102363)]
a, b, r = block(wp, 5, "PC", pc, 35, "Circulante")
r = tot(wp, r, "Total do passivo circulante", f"=SUM(D{a}:D{b})", f"=SUM(E{a}:E{b})", (17875963, 17370535), 35, "PC")
a, b, r = block(wp, r, "PNC", pnc, 35, "Não circulante")
r = tot(wp, r, "Total do passivo não circulante", f"=SUM(D{a}:D{b})", f"=SUM(E{a}:E{b})", (39719457, 37963630), 35, "PNC")
a, b, r = block(wp, r, "PL", pl, 35, "Patrimônio líquido (nota 24)")
r = tot(wp, r, "PL atribuído aos controladores", f"=SUM(D{a}:D{b})", f"=SUM(E{a}:E{b})", (22489775, 20731762), 35, "PLC")
cell(wp, r, 1, "PL08"); cell(wp, r, 2, "PL atribuído aos acionistas não controladores"); cell(wp, r, 4, 1014735, fill=INP, num=True); cell(wp, r, 5, 1063267, fill=INP, num=True); cell(wp, r, 6, DOC); cell(wp, r, 7, 35); cell(wp, r, 8, 3)
nci = r; r += 1
plc = REG["PLC_div"][1]
r = tot(wp, r, "Total do patrimônio líquido", f"=D{plc}+D{nci}", f"=E{plc}+E{nci}", (23504510, 21795029), 35, "PLT")
pcr, pncr, pltr = REG["PC_div"][1], REG["PNC_div"][1], REG["PLT_div"][1]
r = tot(wp, r, "TOTAL DO PASSIVO E PL", f"=D{pcr}+D{pncr}+D{pltr}", f"=E{pcr}+E{pncr}+E{pltr}", (81099930, 77129194), 35, "PT")

# ------------------------------------------------ DRE
wd = wb.create_sheet("DRE")
header(wd, "Demonstração do resultado consolidada (R$ mil)", COLS, W,
       "Fonte: DF25 CPFL.pdf, p. física 36 / impressa 4. Despesas negativas. Subcomponentes (depreciação etc.) registrados para análise; totais recalculados por fórmula.")
dre = [("Receita operacional líquida",26,44367780,42628210),("Custo com energia elétrica",27,-20376889,-19519468),("Custo com operação",None,-4299812,-4371836),("Custo com serviço prestado a terceiros",28,-5566883,-5168133)]
a, b, r = block(wd, 5, "DR", dre, 36, "Receita e custos")
r = tot(wd, r, "Lucro operacional bruto", f"=SUM(D{a}:D{b})", f"=SUM(E{a}:E{b})", (14124196, 13568773), 36, "LB")
desp = [("Despesas com vendas",None,-985664,-981340),("Despesas gerais e administrativas",None,-1722799,-1588046),("Outras despesas operacionais",None,-637243,-505620)]
a, b, r = block(wd, r, "DD", desp, 36, "Despesas operacionais")
lb = REG["LB_div"][1]
r = tot(wd, r, "Resultado do serviço", f"=D{lb}+SUM(D{a}:D{b})", f"=E{lb}+SUM(E{a}:E{b})", (10778490, 10493767), 36, "RS")
fin = [("Resultado de participações societárias",13,275980,336462)]
a, b, r = block(wd, r, "DP", fin, 36, "Participações")
rs = REG["RS_div"][1]
r = tot(wd, r, "Resultado antes do resultado financeiro", f"=D{rs}+D{a}", f"=E{rs}+E{a}", (11054471, 10830229), 36, "RAF")
ff = [("Receitas financeiras",29,1979745,1601868),("Despesas financeiras",29,-4942969,-4343203)]
a, b, r = block(wd, r, "DF", ff, 36, "Resultado financeiro (nota 29)")
raf = REG["RAF_div"][1]
r = tot(wd, r, "Lucro antes dos tributos", f"=D{raf}+SUM(D{a}:D{b})", f"=E{raf}+SUM(E{a}:E{b})", (8091246, 8088893), 36, "LAIR")
tr = [("Contribuição social",None,-634771,-620678),("Imposto de renda",None,-1713923,-1706661)]
a, b, r = block(wd, r, "DT", tr, 36, "Tributos sobre o lucro")
lair = REG["LAIR_div"][1]
r = tot(wd, r, "Lucro líquido do exercício", f"=D{lair}+SUM(D{a}:D{b})", f"=E{lair}+SUM(E{a}:E{b})", (5742552, 5761554), 36, "LL")
det = [("  Depreciação e amortização (custo com operação)",None,-1930668,-1817542),("  Outros custos com operação",28,-2369144,-2554294),("  PDD (dentro de despesas com vendas)",None,-361253,-416808),("  Amortização de intangível da concessão (outras despesas)",None,-332792,-332791),("Lucro atribuído aos controladores",None,5484324,5457652),("Lucro atribuído aos não controladores",None,258229,303902)]
a, b, r = block(wd, r, "DX", det, 36, "Detalhes e atribuição")
ll = REG["LL_div"][1]
cell(wd, r, 2, "Controladores + não controladores − lucro líquido", italic=True); cell(wd, r, 4, f"=D{b-1}+D{b}-D{ll}", fill=CHK, num=True); cell(wd, r, 5, f"=E{b-1}+E{b}-E{ll}", fill=CHK, num=True); cell(wd, r, 6, "esperado 0/±1"); r += 1
cell(wd, r, 2, "Custo com operação − (depreciação + outros custos)", italic=True); cell(wd, r, 4, f"=D{REG['Custo com operação|DR'][1]}-D{a}-D{a+1}", fill=CHK, num=True); cell(wd, r, 5, f"=E{REG['Custo com operação|DR'][1]}-E{a}-E{a+1}", fill=CHK, num=True); r += 1
cell(wd, r, 1, "DX07"); cell(wd, r, 2, "Lucro por ação atribuído aos controladores (R$)"); cell(wd, r, 4, 4.76, fill=INP); cell(wd, r, 5, 4.74, fill=INP); cell(wd, r, 6, DOC); cell(wd, r, 7, 36); cell(wd, r, 8, 4)

# ------------------------------------------------ DFC
wf = wb.create_sheet("DFC")
header(wf, "Demonstração dos fluxos de caixa consolidada — método indireto (R$ mil)", COLS, W,
       "Fonte: DF25 CPFL.pdf, p. física 39 / impressa 7. Linhas detalhadas apenas onde servem a PTs (520, 570, 330); subtotais transcritos e conferidos.")
dfc = [("Lucro antes dos tributos",None,8091246,8088893),("Depreciação e amortização",None,2397217,2303124),("Provisão para créditos de liquidação duvidosa",None,361253,416808),("Encargos de dívida e atualizações monetárias e cambiais",None,2642184,2432471),("Equivalência patrimonial",None,-275980,-336462),("Caixa líquido gerado pelas operações",None,10806804,11850611),("Encargos de dívidas e debêntures pagos",None,-2084774,-2142754),("Imposto de renda e contribuição social pagos",None,-1559233,-2918915),("Caixa líquido gerado pelas atividades operacionais",None,7162797,6788942),("Adições de ativo contratual",None,-4963868,-4543589),("Aquisições de imobilizado",None,-369941,-509943),("Caixa líquido consumido pelas atividades de investimento",None,-4399571,-5531381),("Captação de empréstimos e debêntures",None,10266269,7761288),("Amortização de principal de empréstimos e debêntures",None,-10727722,-6542435),("Dividendo e juros sobre o capital próprio pagos",None,-3519769,-3479652),("Caixa líquido consumido pelas atividades de financiamento",None,-2507307,-3719346),("Aumento (redução) líquido de caixa e equivalentes",None,255919,-2461785),("Saldo inicial de caixa e equivalentes",None,1973401,4435186),("Saldo final de caixa e equivalentes",None,2229320,1973401)]
a, b, r = block(wf, 5, "FC", dfc, 39, "Linhas selecionadas")
def rr(i): return a + i - 1
cell(wf, r, 2, "Operações após encargos e IR (calculado) − divulgado", italic=True); cell(wf, r, 4, f"=D{rr(6)}+D{rr(7)}+D{rr(8)}-D{rr(9)}", fill=CHK, num=True); cell(wf, r, 5, f"=E{rr(6)}+E{rr(7)}+E{rr(8)}-E{rr(9)}", fill=CHK, num=True); r += 1
cell(wf, r, 2, "Oper. + invest. + financ. − variação líquida", italic=True); cell(wf, r, 4, f"=D{rr(9)}+D{rr(12)}+D{rr(16)}-D{rr(17)}", fill=CHK, num=True); cell(wf, r, 5, f"=E{rr(9)}+E{rr(12)}+E{rr(16)}-E{rr(17)}", fill=CHK, num=True); r += 1
cell(wf, r, 2, "Saldo inicial + variação − saldo final", italic=True); cell(wf, r, 4, f"=D{rr(18)}+D{rr(17)}-D{rr(19)}", fill=CHK, num=True); cell(wf, r, 5, f"=E{rr(18)}+E{rr(17)}-E{rr(19)}", fill=CHK, num=True); r += 1
DFC_LAIR, DFC_CAIXA, DFC_INI = rr(1), rr(19), rr(18)

# ------------------------------------------------ NOTA 11
wn = wb.create_sheet("Nota11_AFC")
header(wn, "Nota 11 — Ativo financeiro da concessão: movimentação (R$ mil, consolidado, não circulante)", COLS, W,
       "Fonte: DF25 CPFL.pdf, nota 11, p. física 68 / impressa 36 (tabela em imagem, lida a 200 dpi). Texto: baixas de 2025 = 113.091 (atualização) + 114.570 (ativo).")
mov = [("Saldo inicial",11,25209768,21617521),("Ajuste ao valor justo (VNR)",11,1285173,1104267),("Transferência - ativo contratual",11,3368915,2690645),("Transferência - intangível em serviço",11,-12576,-8185),("Baixas",11,-227661,-194481)]
a, b, r = block(wn, 5, "N11", mov, 68, "Movimentação")
r = tot(wn, r, "Saldo final", f"=SUM(D{a}:D{b})", f"=SUM(E{a}:E{b})", (29623619, 25209768), 68, "N11F")
cell(wn, r, 2, "Baixa da atualização + baixa do ativo (texto da nota)"); cell(wn, r, 4, "=-(113091+114570)", num=True); cell(wn, r, 5, "=-(96326+98156)", num=True); cell(wn, r, 6, DOC); cell(wn, r, 7, 68); r += 1
cell(wn, r, 2, "Diferença vs. linha Baixas", italic=True); cell(wn, r, 4, f"=D{r-1}-D{b}", fill=CHK, num=True); cell(wn, r, 5, f"=E{r-1}-E{b}", fill=CHK, num=True); r += 1
cell(wn, r, 2, "Ajuste VNR / lucro antes dos tributos (DRE)", italic=True); cell(wn, r, 4, f"=D{a+1}/DRE!D{REG['LAIR_div'][1]}"); wn.cell(row=r, column=4).number_format = "0.0%"; cell(wn, r, 6, "insumo PT 320/540"); r += 1

# ------------------------------------------------ CONCILIACOES
wc = wb.create_sheet("Conciliacoes")
header(wc, "Conciliações entre demonstrações e notas (todas por fórmula)", ["ID", "Teste", "2025", "2024", "Critério", "Resultado esperado", "", "", ""], [8, 70, 16, 16, 30, 30, 2, 2, 2],
       "Diferenças de ±1 a ±2 decorrem do arredondamento descrito na nota 2.4 da DF25 (p. física 46). Qualquer outra diferença deve ser investigada.")
at = REG["AT_div"][1]; pt = REG["PT_div"][1]
tests = [
 ("CC01", "Ativo total − (passivo + PL) [equação patrimonial]", f"=BP_Ativo!D{at}-BP_Passivo_PL!D{pt}", f"=BP_Ativo!E{at}-BP_Passivo_PL!E{pt}"),
 ("CC02", "Lucro antes dos tributos: DRE − DFC", f"=DRE!D{REG['LAIR_div'][1]}-DFC!D{DFC_LAIR}", f"=DRE!E{REG['LAIR_div'][1]}-DFC!E{DFC_LAIR}"),
 ("CC03", "Caixa final DFC − caixa BP", f"=DFC!D{DFC_CAIXA}-BP_Ativo!D6", f"=DFC!E{DFC_CAIXA}-BP_Ativo!E6"),
 ("CC04", "Caixa inicial DFC 2025 − caixa BP 2024", f"=DFC!D{DFC_INI}-BP_Ativo!E6", None),
 ("CC05", "Nota 11 saldo final − linha do BP (ativo financeiro da concessão)", f"=Nota11_AFC!D{REG['N11F_div'][1]}-BP_Ativo!D{AFC_ROW}", f"=Nota11_AFC!E{REG['N11F_div'][1]}-BP_Ativo!E{AFC_ROW}"),
 ("CC06", "Nota 11 saldo inicial 2025 − saldo final 2024", f"=Nota11_AFC!D6-Nota11_AFC!E{REG['N11F_div'][1]}", None),
 ("CC07", "PDD: DRE (vendas) − DFC (ajuste)", f"=-DRE!D{REG['  PDD (dentro de despesas com vendas)|DX'][1]}-DFC!D8", f"=-DRE!E{REG['  PDD (dentro de despesas com vendas)|DX'][1]}-DFC!E8"),
]
r = 5
for t in tests:
    cell(wc, r, 1, t[0]); cell(wc, r, 2, t[1]); cell(wc, r, 3, t[2], fill=CHK, num=True)
    if t[3]: cell(wc, r, 4, t[3], fill=CHK, num=True)
    cell(wc, r, 5, "|dif| ≤ 2"); cell(wc, r, 6, "0 ou arredondamento"); r += 1
cell(wc, r + 1, 2, "Indicadores-chave para PTs 320/520/570 (fórmulas)", bold=True); r += 2
ind = [
 ("IN01", "Ativo financeiro da concessão / ativo total", f"=BP_Ativo!D{AFC_ROW}/BP_Ativo!D{at}", f"=BP_Ativo!E{AFC_ROW}/BP_Ativo!E{at}", "0.0%"),
 ("IN02", "Liquidez corrente (AC/PC)", f"=BP_Ativo!D{REG['AC_div'][1]}/BP_Passivo_PL!D{REG['PC_div'][1]}", f"=BP_Ativo!E{REG['AC_div'][1]}/BP_Passivo_PL!E{REG['PC_div'][1]}", "0.00"),
 ("IN03", "Dívida bruta (empréstimos + debêntures, CP+LP)", "=BP_Passivo_PL!D7+BP_Passivo_PL!D8+BP_Passivo_PL!D{0}+BP_Passivo_PL!D{1}".format(REG['Empréstimos e financiamentos|PNC'][1], REG['Debêntures|PNC'][1]), "=BP_Passivo_PL!E7+BP_Passivo_PL!E8+BP_Passivo_PL!E{0}+BP_Passivo_PL!E{1}".format(REG['Empréstimos e financiamentos|PNC'][1], REG['Debêntures|PNC'][1]), NUM),
 ("IN04", "Margem líquida (LL/receita)", f"=DRE!D{REG['LL_div'][1]}/DRE!D6", f"=DRE!E{REG['LL_div'][1]}/DRE!E6", "0.0%"),
 ("IN05", "Variação da receita líquida 2025/2024", "=DRE!D6/DRE!E6-1", None, "0.0%"),
]
for t in ind:
    cell(wc, r, 1, t[0]); cell(wc, r, 2, t[1]); x = cell(wc, r, 3, t[2]); x.number_format = t[4]
    if t[3]: y = cell(wc, r, 4, t[3]); y.number_format = t[4]
    cell(wc, r, 5, "cálculo próprio — não é indicador divulgado"); r += 1

# ------------------------------------------------ GRUPO (nota 1)
wg = wb.create_sheet("Grupo_Nota1")
header(wg, "Estrutura do grupo — nota 1 (participações em 31/12/2025)", ["Segmento", "Entidade", "Tipo", "Participação", "Localização", "Término da concessão / obs.", "Documento", "Pág. física", "Pág. impressa"], [14, 46, 22, 16, 26, 26, 14, 9, 9],
       "Fonte: DF25 CPFL.pdf, nota 1, p. física 41–42 / impressa 9–10 (tabela em imagem). Classificação contábil (controlada / controlada em conjunto / coligada) a confirmar na nota 13 no Bloco 05.")
grp = [("Distribuição","Companhia Paulista de Força e Luz (CPFL Paulista)","S.A. aberta","Direta 100%","Interior de SP","Novembro de 2027"),
 ("Distribuição","Companhia Piratininga de Força e Luz (CPFL Piratininga)","S.A. aberta","Direta 100%","Interior e litoral de SP","Outubro de 2028"),
 ("Distribuição","RGE Sul Distribuidora de Energia S.A. (CPFL RGE)","S.A. aberta","Direta e indireta 100%","Interior do RS","Novembro de 2027"),
 ("Distribuição","Companhia Jaguari de Energia (CPFL Santa Cruz)","S.A. fechada","Direta 100%","SP, PR e MG","Julho de 2045"),
 ("Geração","CPFL Energias Renováveis S.A. (CPFL Renováveis)","S.A. aberta","Direta e indireta 100%","Vários estados","103 empreendimentos; 2.905,7 MW"),
 ("Geração","CPFL Geração de Energia S.A.","S.A. aberta","Direta 100%","SP","Incorporada pela CPFL Brasil em 02/02/2026 (nota 37.1)"),
 ("Geração","CERAN","S.A. fechada","Indireta 65%","RS","3 UHEs"),
 ("Geração","Foz do Chapecó Energia S.A.","S.A. fechada","Indireta 51% (c)","SC e RS","Controlada em conjunto via Chapecoense (nota 1c)"),
 ("Geração","Campos Novos Energia S.A. (ENERCAN)","S.A. fechada","Indireta 52,12%","SC",""),
 ("Geração","BAESA - Energética Barra Grande S.A.","S.A. fechada","Indireta 25,01%","SC e RS",""),
 ("Geração","Paulista Lajeado Energia S.A.","S.A. fechada","Indireta 56,93% (a)","TO",""),
 ("Transmissão","CPFL Transmissão (antiga CEEE-T)","S.A. aberta","Direta e indireta 100%","RS","Lote 3 Leilão 4/2025 arrematado em 31/10/2025"),
 ("Transmissão","CPFL Piracicaba, Morro Agudo, Maracanaú, Sul I, Sul II","Ltda.","Indireta 100%","SP, CE, SC, RS",""),
 ("Transmissão","TESB","S.A. fechada","Indireta 98,01%","RS",""),
 ("Transmissão","TPAE","S.A. fechada","Indireta 9,65%","RS",""),
 ("Transmissão","ETAU","S.A. fechada","Indireta 10%","RS",""),
 ("Comercialização","CPFL Comercialização Brasil S.A. (CPFL Brasil)","S.A. fechada","Direta 100%","—","Sucessora da CPFL Geração (2026)"),
 ("Comercialização","CPFL Meridional, CPFL Planalto, CPFL Brasil Varejista","Ltda.","100% (direta/indireta)","—","")]
r = 5
for g in grp:
    for j, v in enumerate(g, 1): cell(wg, r, j, v)
    cell(wg, r, 7, DOC); cell(wg, r, 8, 41); cell(wg, r, 9, 9); r += 1

# ------------------------------------------------ AUDITOR
wa = wb.create_sheet("Relatorio_Auditor")
header(wa, "Relatório do auditor independente sobre as DF 2025 (fonte pública — não é procedimento do aluno)", ["Item", "Conteúdo", "Documento", "Pág. física", "Pág. impressa", "", "", "", ""], [30, 90, 16, 9, 9, 2, 2, 2, 2],
       "Leitura do relatório real da PwC. Os procedimentos descritos nele foram executados pela PwC, não por este trabalho acadêmico.")
aud = [("Firma","PricewaterhouseCoopers Auditores Independentes Ltda. (CRC 2SP027613/F-1), Campinas/SP",26),
 ("Data do relatório","Campinas, 4 de março de 2026",32),
 ("Opinião","Sem modificação — demonstrações individuais e consolidadas apresentam adequadamente, em todos os aspectos relevantes, conforme BR GAAP e IFRS",26),
 ("Base da opinião","Normas brasileiras e internacionais de auditoria; independência conforme Código de Ética do CFC (entidades de interesse público)",26),
 ("PAA 1","Mensuração do ativo contratual das transmissoras e receita de remuneração/atualização (notas 3.2 e 15)",27),
 ("PAA 2","Receita de energia fornecida e não faturada (notas 3.11 e 7): R$ 1.763.858 mil em 31/12/2025; PwC envolveu especialistas de TI e recalculou",28),
 ("PAA 3","Mensuração e classificação do ativo financeiro da concessão (R$ 29.623.619 mil), ativo contratual (R$ 4.760.953 mil) e infraestrutura de distribuição em serviço (R$ 3.935.652 mil) (notas 11, 15, 16)",29),
 ("Outros assuntos","Demonstrações do Valor Adicionado (DVA) examinadas como informação suplementar (CPC 09)",30),
 ("Ênfase / ressalva","Nenhum parágrafo de ênfase ou ressalva identificado na leitura do texto das p. 26–32",26),
 ("Aprovação das DF","Autorização pela administração em 23/02/2026; aprovação pelo Conselho de Administração em 04/03/2026 (nota 2.1)",45)]
r = 5
for it, ct, pg in aud:
    cell(wa, r, 1, it, bold=True); cell(wa, r, 2, ct); wa.cell(row=r, column=2).alignment = Alignment(wrap_text=True)
    cell(wa, r, 3, DOC); cell(wa, r, 4, pg); cell(wa, r, 5, pg - 32 if pg > 32 else "relatório"); r += 1

# ------------------------------------------------ FONTES USADAS
wfu = wb.create_sheet("Fontes_Usadas")
header(wfu, "FONTES_USADAS — Bloco 02", ["ID", "Documento", "Versão / exercício", "Local", "URL", "Data de acesso", "Páginas usadas", "Uso", "Observação"], [8, 40, 18, 26, 44, 12, 22, 30, 40],
       "Fonte = documento efetivamente aberto nesta sessão. URL de portal só como localizador. IDs F01–F20 (etapa 5) preservados; FC = fontes centrais do Bloco 01.")
fu = [("FC01=F01","DF25 CPFL.pdf — DF individuais e consolidadas 2025 + relatório PwC","Exercício 2025 (comparativo 2024); 115 p.; SHA-256 0694f2b9…a578","Drive 01_Demonstracoes (ID 1JcqtNdI-70DCaukfmVni-rR-nQbzW25h)","https://publicidadelegal.valor.com.br/valor/2026/03/20/CPFLENERGIA1588267020032026.pdf (publicação legal registrada no controle; não reaberta)","22/10/2026","26–36, 39, 41–42, 45–46, 68, 113–115","Saldos, notas 1, 2, 11, 37; relatório do auditor","Data de acesso = abertura do arquivo do Drive; origem web original não comprovada nesta sessão"),
 ("FC14","FR_CPFL Energia 2025_V6_06-04-26_PT.pdf — Formulário de Referência 2025, versão 6","FR 2025 v6","Drive 03 (ID 12xVd4uA7Fj9cwYYlC3hYPGFw0PSJDb-d)","https://www.rad.cvm.gov.br/ENETWEB/frmConsultaExternaCVM.aspx (portal)","22/10/2026","Item 1.1, p. 1 (numeração do FR)","Histórico: 'anteriormente denominada Draft II'; NYSE/ADS deslistado em 2020; controle State Grid","ARQUIVO TRUNCADO: 16.777.216 bytes; texto termina no item 2.1 (DV15)"),
 ("FC10","Enunciado P2 (texto do aluno)","2026.2","Conversa","—","22/10/2026","—","Escopo","—")]
r = 5
for f in fu:
    for j, v in enumerate(f, 1): cell(wfu, r, j, v); wfu.cell(row=r, column=j).alignment = Alignment(wrap_text=True, vertical="top")
    r += 1

# ------------------------------------------------ DIVERGENCIAS B02
wv = wb.create_sheet("Divergencias_B02")
header(wv, "Divergências identificadas no Bloco 02", ["ID", "Origem (controle antigo)", "Fonte primária", "Impacto", "Tratamento", "Estado", "", "", ""], [7, 40, 40, 36, 40, 14, 2, 2, 2], "Continuação de DV01–DV11 (Bloco 01).")
dv = [("DV12","00_CONTROLE: despesas com vendas 2025 = (986.664)","DF25 p. física 36: (985.664)","R$ 1.000 mil; não altera LAIR (controle tinha LAIR correto)","Usar 985.664","Resolvida"),
 ("DV13","00_CONTROLE: resultado do serviço 2025 = 10.776.490","DF25 p. física 36: 10.778.490","R$ 2.000 mil em subtotal; LAIR e LL corretos no controle","Usar 10.778.490","Resolvida"),
 ("DV14","00_CONTROLE seção 3: nome anterior 'Draftll Participações S.A.'","FR 2025 v6, item 1.1, p. 1: 'anteriormente denominada Draft II'","Identificação societária (exigência do enunciado)","Usar 'Draft II' como consta no FR; confirmar grafia completa no cadastro CVM quando acessível (PB02-02)","Parcial"),
 ("DV15","FR 2025 e FR 2026 no Drive com 16.777.216 bytes (16 MiB exatos)","Texto do FR 2025 termina no item 2.1 (índice prevê 13 seções, ~296 p.)","FR incompleto: itens 4 (riscos), 9 (auditores), 11 (partes relacionadas), 15 (controle) indisponíveis","Usar apenas o trecho legível; baixar versão íntegra (ENET/RI) quando houver rota (PB02-03)","Aberta")]
r = 5
for d in dv:
    for j, v in enumerate(d, 1): cell(wv, r, j, v); wv.cell(row=r, column=j).alignment = Alignment(wrap_text=True, vertical="top")
    r += 1

os.makedirs(OUT, exist_ok=True)
p = os.path.join(OUT, "CPFL_B02_Dados_Reais_2025.xlsx"); wb.save(p); print(p)
