# -*- coding: utf-8 -*-
"""B12 — consistência de números-chave no dossiê e nas planilhas (variantes próximas = possível erro)."""
import re, sys, glob, json
from openpyxl import load_workbook
md = open(sys.argv[1], encoding='utf-8').read()
CH = {'materialidade global': '404.562', 'execução': '303.421', 'trivial': '20.228', 'AFC': '29.623.619', 'ativo total': '81.099.930',
      'LAIR': '8.091.246', 'receita líquida': '44.367.780', 'provisões': '1.740.903', 'perdas possíveis': '10.748.481', 'mútuo SGEL': '4.412.650',
      'IRRF Jantus': '139.291', 'reclassificação 2026': '29.599.494', 'distorção sintética': '117.484', 'Paulista AFC': '12.617.237', 'RGE AFC': '12.734.234',
      'Piratininga AFC': '3.836.600', 'Santa Cruz AFC': '435.548', 'CCL': '1.259.649', 'caixa+TVM': '2.905.082', 'passivo 12m': '12.916.273'}
nums = re.findall(r'(?<![\d.,])\d{1,3}(?:\.\d{3})+(?:,\d+)?', md)
out = {}
for nome, v in CH.items():
    base = v.replace('.', ''); n = len(base)
    var = set()
    for x in nums:
        xb = x.split(',')[0].replace('.', '')
        if len(xb) == n and xb != base and xb[:max(3, n-2)] == base[:max(3, n-2)]:
            var.add(x)
    out[nome] = dict(valor=v, ocorrencias=md.count(v), variantes=sorted(var))
for k, r in out.items(): print(f"{k:22s} {r['valor']:>12s} x{r['ocorrencias']:<3d} variantes: {r['variantes']}")
json.dump(out, open('/root/cpfl/revisao_b12/check_numeros.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
# planilhas: erros de fórmula
err = {}
for f in sorted(glob.glob('/root/cpfl/producao/*/*.xlsx')):
    wb = load_workbook(f, data_only=True); e = 0
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('#') and c.value.rstrip('!?/0A').upper() in ('#REF', '#VALUE', '#NAME', '#DIV', '#N/', '#NUM', '#NULL') or (isinstance(c.value, str) and re.match(r'#(REF|VALUE|NAME|DIV/0|N/A|NUM|NULL)', c.value)):
                    e += 1
    err[f.split('/producao/')[1]] = e
print(json.dumps(err, ensure_ascii=False, indent=0))
