# -*- coding: utf-8 -*-
"""Bloco 04 - Planilha de planejamento: materialidade (fórmulas), matriz de riscos NBC TA 315 (R2),
respostas NBC TA 330, programas por ciclo, fraude (240), leis (250) e cronograma (300).
Uso: python b04_planejamento.py <pasta_saida>"""
import sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from b04_dados import RISCOS, RESPOSTAS, PROGRAMAS, FRAUDE, LEIS, CRONOGRAMA

OUT = sys.argv[1] if len(sys.argv) > 1 else "."
F = "Arial"; HEAD = PatternFill("solid", fgColor="1F3864"); INP = PatternFill("solid", fgColor="FFF2CC"); CHK = PatternFill("solid", fgColor="E2EFDA")
th = Side(style="thin", color="BFBFBF"); BRD = Border(left=th, right=th, top=th, bottom=th)
wb = Workbook(); wb.remove(wb.active)

def sheet(name, title, cols, widths, note):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = Font(name=F, bold=True, size=12, color="1F3864")
    ws["A2"] = "Preparador: matrícula 231020069 (assistido por IA) | Data: 22/10/2026 | Revisão: ____ (sem revisão de terceiros simulada)"
    ws["A2"].font = Font(name=F, italic=True, size=8, color="595959")
    ws["A3"] = note; ws["A3"].font = Font(name=F, size=9, color="7F6000"); ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(cols)); ws.row_dimensions[3].height = 48
    for j, h in enumerate(cols, 1):
        c = ws.cell(row=4, column=j, value=h); c.font = Font(name=F, bold=True, color="FFFFFF", size=9); c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BRD
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A5"; return ws

def put(ws, r, vals, fills=None, fmt=None, bold=False):
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=j, value=v); c.font = Font(name=F, size=9, bold=bold); c.border = BRD
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if fills and j in fills: c.fill = fills[j]
        if fmt and j in fmt: c.number_format = fmt[j]

# ------------------------------------------------------------ MATERIALIDADE
ws = sheet("Materialidade", "PT 320 — Materialidade (R$ mil) — fórmulas", ["Item", "Valor", "Unidade", "Fonte / natureza", "Observação"], [44, 18, 12, 46, 60],
           "Bases reconferidas na DF25 em 22/10/2026 (B02). Percentuais = julgamento didático (NBC TA 320 item A8 dá 5% do lucro antes do imposto como EXEMPLO, não como regra). Células amarelas são entradas; as demais, fórmulas.")
N = "#,##0.00"; P = "0.00%"
rows = [
 ("Lucro antes dos tributos (LAIR)", 8091246, "R$ mil", "DF25 p. 36/4 (DRE) e p. 39/7 (DFC)", "Base principal: companhia lucrativa e estável (LAIR 2024: 8.088.893)"),
 ("Receita operacional líquida", 44367780, "R$ mil", "DF25 p. 36/4", "Base alternativa"),
 ("Ativo total", 81099930, "R$ mil", "DF25 p. 34/2", "Base alternativa (intensiva em capital)"),
 ("Patrimônio líquido total", 23504510, "R$ mil", "DF25 p. 35/3", "Base alternativa (referência)"),
 ("Ajuste ao valor justo VNR (nota 11)", 1285173, "R$ mil", "DF25 p. 68/36", "Componente não caixa do resultado"),
 ("% sobre LAIR", 0.05, "%", "Julgamento (exemplo em 320 A8)", ""),
 ("% sobre receita", 0.005, "%", "Julgamento", ""),
 ("% sobre ativo", 0.005, "%", "Julgamento", ""),
 ("% sobre PL", 0.01, "%", "Julgamento", ""),
 ("% execução / global", 0.75, "%", "Julgamento (aprovado na etapa 2)", "Mantido: sem histórico de distorções conhecido; relatório sem modificação. Sensibilidade abaixo"),
 ("% trivial / global", 0.05, "%", "Julgamento (limite operacional NBC TA 450)", ""),
]
r = 5
for it in rows:
    put(ws, r, list(it), fills={2: INP}, fmt={2: N if it[2] == "R$ mil" else P}); r += 1
r += 1
calc = [
 ("Materialidade pelo LAIR", "=B5*B10"), ("Materialidade pela receita", "=B6*B11"), ("Materialidade pelo ativo", "=B7*B12"), ("Materialidade pelo PL", "=B8*B13"),
 ("MATERIALIDADE GLOBAL SELECIONADA (LAIR)", "=B17"), ("Materialidade para execução", "=B21*B14"), ("Limite de distorções claramente triviais", "=B21*B15"),
 ("LAIR sem ajuste VNR (sensibilidade mecânica)", "=B5-B9"), ("Materialidade sensibilidade VNR", "=B24*B10"), ("Diferença vs. global", "=B21-B25"),
 ("Execução a 60% (sensibilidade)", "=B21*0.6"), ("Execução a 50% (sensibilidade)", "=B21*0.5"),
 ("Estoques / materialidade global", "=236566/B21"), ("Caixa e equivalentes / materialidade global", "=2229320/B21"),
]
for lab, f in calc:
    put(ws, r, [lab, f, "R$ mil" if "/" not in lab else "x", "fórmula", ""], fills={2: CHK}, fmt={2: N if "/" not in lab else "0.00"}, bold=lab.startswith("MATERIALIDADE")); r += 1
r += 1
for t in ["Materialidade específica (item 10): não se fixa valor quantitativo inferior para nenhuma classe. Fundamento: as divulgações sensíveis (remuneração da administração, partes relacionadas — PT 550) são avaliadas qualitativamente (NBC TA 450), e as estimativas de maior risco já recebem resposta substantiva extensiva (PT 330). Reavaliar se o PT 550 identificar transações com o controlador sensíveis aos usuários.",
          "Gatilhos de revisão (item 12–13): (i) distorções acumuladas acima de 50% da execução; (ii) evento subsequente que altere o LAIR (ex.: renovação das concessões afetando o VNR); (iii) reapresentação de 2024 identificada no PT 510.",
          "Documentação (item 14): valores e fatores registrados nesta aba e no PT 320."]:
    c = ws.cell(row=r, column=1, value=t); c.font = Font(name=F, size=9, italic=True); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    c.alignment = Alignment(wrap_text=True); ws.row_dimensions[r].height = 52; r += 1

# ------------------------------------------------------------ MATRIZ 315
ws = sheet("Matriz_Riscos_315", "PT 315 — Matriz de riscos por conta e afirmação (NBC TA 315 R2)",
           ["ID", "Conta / processo", "Saldo 2025 (R$ mil)", "Fonte", "Afirmações", "Risco", "Fatores de risco inerente*", "Prob. (1-3)", "Magn. (1-3)", "Pontuação", "Nível (espectro)", "Risco significativo?", "Fundamento", "Saldo / MG", "Risco de controle"],
           [6, 30, 14, 26, 22, 40, 12, 8, 8, 9, 11, 13, 42, 9, 26],
           "*C = complexidade; S = subjetividade; M = mudança; I = incerteza; V = suscetibilidade a viés/fraude (315 item 12(f)). Pontuação e nível = proxy didática do espectro de risco inerente (315 item 31), não fórmula da norma. Risco de controle: sem teste de efetividade, o risco de distorção relevante = risco inerente (315 item 34). Revisão da matriz preliminar da etapa 3 (17 riscos) + R18 e R19.")
r = 5
for (i, c, s, fo, af, ri, fa, p, m, nv, sig, fu) in RISCOS:
    put(ws, r, [i, c, s, fo, af, ri, fa, p, m, f"=H{r}*I{r}", nv, sig, fu, (f"=C{r}/Materialidade!$B$21" if isinstance(s, (int, float)) else "—"), "Não testado → RDR = risco inerente (315 item 34)"],
        fills={3: INP, 8: INP, 9: INP, 10: CHK, 14: CHK}, fmt={3: "#,##0", 14: "0.0"}); r += 1

# ------------------------------------------------------------ RESPOSTAS 330
ws = sheet("Respostas_330", "PT 330 — Respostas aos riscos avaliados (natureza, época, extensão, modalidade)",
           ["Risco", "Procedimento", "Natureza", "Época", "Extensão", "Modalidade", "PT"], [7, 60, 22, 14, 30, 34, 10],
           "Respostas gerais (330 item 5): ceticismo reforçado em estimativas, imprevisibilidade na seleção (amostragem aleatória com semente registrada), reconferência de toda fonte secundária. Riscos significativos (R01, R02, R03, R16): procedimentos substantivos específicos (330 item 21). Época 'interina' é planejada; nada foi executado na companhia.")
r = 5
for rid, lst in RESPOSTAS.items():
    for pr in lst: put(ws, r, [rid] + list(pr)); r += 1

# ------------------------------------------------------------ PROGRAMAS
ws = sheet("Programas_Ciclos", "Programas de auditoria por ciclo (PT 300/330)", ["Código", "Ciclo", "Riscos", "Saldo-base", "Aplicabilidade e extensão", "Procedimento", "Nº"], [10, 18, 16, 28, 40, 72, 5],
           "Legenda: [Pub] realizado com evidência pública; [P] planejado por falta de acesso; [S] simulado. O estado de execução de cada procedimento é atualizado nos blocos seguintes.")
r = 5
for cod, ci, rs, sb, ap, procs in PROGRAMAS:
    for k, pr in enumerate(procs, 1): put(ws, r, [cod, ci, rs, sb, ap, pr, k]); r += 1

# ------------------------------------------------------------ FRAUDE
ws = sheet("Fraude_240", "PT 240 — Fatores de risco de fraude (NBC TA 240 R1)", ["Categoria", "Fator", "Fonte", "Natureza"], [24, 70, 40, 14],
           "Classificação dos fatores conforme a NBC TA 240. Fator de risco não é alegação de fraude.")
r = 5
for f in FRAUDE: put(ws, r, list(f)); r += 1

# ------------------------------------------------------------ LEIS
ws = sheet("Leis_250", "PT 250 — Mapa de leis e regulamentos", ["Categoria (250 itens 14–15)", "Lei / regulamento", "Fonte", "Efeito nas DF", "Procedimento e modalidade"], [22, 52, 34, 30, 44],
           "Mapeamento para fins de auditoria; não é parecer jurídico.")
r = 5
for l in LEIS: put(ws, r, list(l)); r += 1

# ------------------------------------------------------------ CRONOGRAMA
ws = sheet("Cronograma_300", "PT 300 — Cronograma", ["Fase", "Data", "Bloco", "PTs"], [48, 16, 10, 50], "Cronograma do trabalho acadêmico (prazo do enunciado: 25/10/2026).")
r = 5
for c in CRONOGRAMA: put(ws, r, list(c)); r += 1

os.makedirs(OUT, exist_ok=True); p = os.path.join(OUT, "CPFL_B04_Planejamento_Riscos_Materialidade.xlsx"); wb.save(p); print(p)
